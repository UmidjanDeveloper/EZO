"""Official-layout Word documents for drafts from `writing.service`.

Layout follows common Uzbek office practice: A4, margins 30/15/20/20 mm, Times New Roman
14 pt (12 pt for CV and academic work), justified body with a 1.25 cm first-line indent,
letterhead centred at the top, outgoing number and date, the recipient block on the
right, the title centred, and a signature block whose signature and seal places are
always placeholders. Drafts of official papers carry a removable "QORALAMA" line in the
page header and "Draft" in the document properties.
"""

from __future__ import annotations

import io
import re
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ezo_api.organization.formatting import apply as apply_formatting
from ezo_api.organization.formatting import formatting_of
from ezo_api.organization.templates import open_template
from ezo_api.writing.catalog import DocType
from ezo_api.writing.service import Draft, OrgFacts

DRAFT_LABEL = {"uz": "QORALAMA", "uz-cyrl": "ҚОРАЛАМА", "ru": "ПРОЕКТ", "en": "DRAFT"}
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_NUMBERED = re.compile(r"^\s*(\d{1,2})[.)]\s+(.*)$")


def _placeholder(doc: DocType, draft: Draft, key: str) -> str:
    value = draft.fields.get(key)
    if value:
        return value
    label = next((f.label for f in doc.fields if f.key == key), key)
    return f"[{label}]"


def _runs(paragraph: Any, text: str, *, bold: bool = False) -> None:
    position = 0
    for match in _BOLD.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()]).bold = bold or None
        paragraph.add_run(match.group(1)).bold = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:]).bold = bold or None


def _rule(paragraph: Any) -> None:
    """A thin line under the letterhead (paragraph bottom border)."""
    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "8"), ("space", "1"), ("color", "000000")):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    ppr.append(borders)


def _body(document: Any, text: str, *, indent: bool, indent_cm: float = 1.25) -> None:
    for block in [b for b in re.split(r"\n\s*\n", text.strip()) if b.strip()]:
        lines = block.splitlines()
        if all(re.match(r"^\s*[-•]\s+", ln) for ln in lines):
            for line in lines:
                _runs(document.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-•]\s+", "", line))
            continue
        if all(_NUMBERED.match(ln) for ln in lines):
            for line in lines:
                match = _NUMBERED.match(line)
                assert match is not None
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.first_line_indent = Cm(indent_cm) if indent else None
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                _runs(paragraph, f"{match.group(1)}. {match.group(2)}")
            continue
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent and indent_cm:
            paragraph.paragraph_format.first_line_indent = Cm(indent_cm)
        _runs(paragraph, " ".join(ln.strip() for ln in lines))


def _signature(document: Any, left: str, right: str = "[imzo]   [F.I.Sh.]") -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    left_cell, right_cell = table.rows[0].cells
    left_cell.paragraphs[0].add_run(left).bold = True
    right_cell.paragraphs[0].add_run(right)
    right_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT


def render_docx(
    doc: DocType,
    draft: Draft,
    *,
    org: OrgFacts | None = None,
    language: str = "uz",
    draft_label: bool = True,
) -> bytes:
    small = doc.layout in ("cv", "academic")
    template = org.template if org else None
    if template is not None:
        # The organization's approved blank: its letterhead, styles and margins win.
        document = open_template(template)
        section = document.sections[0]
    else:
        document = Document()
        fmt = formatting_of(org.formatting if org else None)
        apply_formatting(document, fmt, small=small)
        normal = document.styles["Normal"]
        normal.paragraph_format.space_after = Pt(0 if not small else 4)
        section = document.sections[0]
    props = document.core_properties
    props.title = (draft.title or doc.names["uz"])[:200]
    props.author = (org.name if org and org.name else "")[:200]
    props.comments = "Prepared with EZO. Review all facts before use."
    if doc.draft and draft_label:
        props.category = "Draft"
        header = section.header
        line = header.paragraphs[0] if template is None else header.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        label = line.add_run(DRAFT_LABEL.get(language, "DRAFT"))
        label.font.size = Pt(9)
        label.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    official = doc.layout in ("letter", "act", "minutes")
    if template is None and official and org and (org.letterhead or org.name):
        if org.logo:
            picture = document.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.add_run().add_picture(io.BytesIO(org.logo), height=Cm(2))
        lines = [ln for ln in (org.letterhead or "").splitlines() if ln.strip()] or [org.name]
        if not org.letterhead and org.department:
            lines.append(org.department)
        for index, text in enumerate(lines):
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(text.strip())
            run.bold = index == 0
        contacts = " · ".join(v for v in (org.address, org.phone, org.email, org.website) if v)
        if contacts and not org.letterhead:
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(contacts).font.size = Pt(10)
        _rule(document.paragraphs[-1])
        document.add_paragraph()

    keys = {f.key for f in doc.fields}
    if doc.layout == "letter":
        if "number" in keys:
            document.add_paragraph(
                f"№ {_placeholder(doc, draft, 'number')}        {_placeholder(doc, draft, 'date')}"
            )
            document.add_paragraph()
        if "recipient" in keys:
            block = document.add_paragraph()
            block.paragraph_format.left_indent = Cm(8.5)
            _runs(block, _placeholder(doc, draft, "recipient"), bold=True)
            if "applicant" in keys:
                applicant = document.add_paragraph()
                applicant.paragraph_format.left_indent = Cm(8.5)
                applicant.add_run(_placeholder(doc, draft, "applicant"))
            if "sender" in keys:
                sender = document.add_paragraph()
                sender.paragraph_format.left_indent = Cm(8.5)
                sender.add_run(_placeholder(doc, draft, "sender"))
            document.add_paragraph()
        title = draft.title or doc.names["uz"]
        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.add_run(
            title.upper()
            if doc.key in ("ariza", "tushuntirish_xati", "bildirishnoma", "malumotnoma")
            else title
        ).bold = True
        if "subject" in keys and draft.fields.get("subject"):
            subject = document.add_paragraph()
            subject.add_run(f"Mavzu: {draft.fields['subject']}").italic = True
        document.add_paragraph()
    elif doc.layout in ("act", "minutes"):
        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.add_run(
            f"{(draft.title or doc.names['uz']).upper()} № {_placeholder(doc, draft, 'number')}"
        ).bold = True
        place = document.add_paragraph(
            f"{_placeholder(doc, draft, 'place')}        {_placeholder(doc, draft, 'date')}"
        )
        place.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for key, label in (
            ("chair", "Rais"),
            ("secretary", "Kotib"),
            ("attendees", "Qatnashdilar"),
            ("members", "Komissiya a'zolari"),
            ("basis", "Asos"),
        ):
            if key in keys:
                entry = document.add_paragraph()
                entry.add_run(f"{label}: ").bold = True
                entry.add_run(_placeholder(doc, draft, key))
        document.add_paragraph()
    else:
        heading = document.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run(draft.title or doc.names["uz"])
        run.bold = True
        run.font.size = Pt(16 if not small else 14)
        document.add_paragraph()

    indent = doc.layout not in ("cv", "speech")
    for section_part in draft.sections:
        if section_part.heading:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.add_run(section_part.heading).bold = True
        _body(
            document,
            section_part.text,
            indent=indent,
            indent_cm=formatting_of(org.formatting if org else None).indent_cm,
        )

    document.add_paragraph()
    if doc.layout == "letter":
        if "signatory" in keys:
            _signature(document, _placeholder(doc, draft, "signatory"), "[imzo]")
        else:
            person = next((k for k in ("applicant", "sender", "name") if k in keys), None)
            if person:
                name = draft.fields.get(person) or "[F.I.Sh.]"
                _signature(document, "", f"[imzo]   {name}")
        if "date" in keys and "number" not in keys:
            document.add_paragraph(_placeholder(doc, draft, "date"))
    elif doc.layout == "minutes":
        _signature(document, "Rais", f"[imzo]   {draft.fields.get('chair') or '[F.I.Sh.]'}")
        _signature(document, "Kotib", f"[imzo]   {draft.fields.get('secretary') or '[F.I.Sh.]'}")
    elif doc.layout == "act":
        if "signatory" in keys:
            _signature(document, _placeholder(doc, draft, "signatory"))
        else:
            _signature(document, "[Lavozimi]")
        document.add_paragraph("[M.O.]")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
