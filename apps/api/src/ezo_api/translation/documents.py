"""Translation outputs for files: Word script conversion in place, bilingual Word
documents, and source/translation paragraph pairs of Word files."""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ezo_api.documents.docx import fidelity
from ezo_api.translation.translit import Script, detect_script, spans

_XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


def _own_texts(paragraph: Any) -> list[Any]:
    """w:t elements of this paragraph (not of paragraphs nested in its text boxes)."""
    tag = qn("w:p")
    return [
        t
        for t in paragraph.iter(qn("w:t"))
        if next((a for a in t.iterancestors() if a.tag == tag), None) is paragraph
    ]


def transliterate_docx(data: bytes, to: Script) -> tuple[bytes, int]:
    """Convert the Uzbek script of a Word file in place, run by run (formatting, tables,
    headers and footers kept). Returns the new file and the number of paragraphs changed."""
    document = fidelity.open_docx(data)
    changed = 0
    for _, root in fidelity.document_parts(document):
        for paragraph in root.iter(qn("w:p")):
            nodes = _own_texts(paragraph)
            texts = [node.text or "" for node in nodes]
            full = "".join(texts)
            if not full.strip():
                continue
            bounds: list[int] = []
            position = 0
            for text in texts:
                bounds.append(position)
                position += len(text)
            outputs = [""] * len(nodes)
            owner = 0
            for span in spans(full, to):
                while owner + 1 < len(nodes) and span.start >= bounds[owner + 1]:
                    owner += 1
                outputs[owner] += span.out
            if outputs != texts:
                changed += 1
                for node, text in zip(nodes, outputs, strict=True):
                    node.text = text
                    if text != text.strip():
                        node.set(_XML_SPACE, "preserve")
    props = document.core_properties
    if props.title and detect_script(props.title) not in (None, to):
        props.title = "".join(s.out for s in spans(props.title, to))[:255]
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue(), changed


def docx_pairs(before: bytes, after: bytes) -> list[tuple[str, str]]:
    """(original, translated) paragraph texts of a Word file edited in place."""
    old = {u.index: u.text for u in fidelity.read_units(fidelity.open_docx(before))}
    new = {u.index: u.text for u in fidelity.read_units(fidelity.open_docx(after))}
    return [(text, new.get(index, "")) for index, text in old.items() if text.strip()]


def bilingual_docx(
    pairs: list[tuple[str, str]], *, title: str, source_label: str, target_label: str
) -> bytes:
    """A two-column Word document: the original on the left, the translation on the right,
    one row per paragraph."""
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(1.5))
    document.core_properties.title = title[:255]
    document.core_properties.comments = (
        "Bilingual translation prepared with EZO. Review before use."
    )
    heading = document.add_paragraph()
    heading.add_run(title).bold = True
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, (source_label, target_label), strict=True):
        run = cell.paragraphs[0].add_run(label)
        run.bold = True
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    for source, target in pairs:
        row = table.add_row().cells
        row[0].paragraphs[0].add_run(source)
        row[1].paragraphs[0].add_run(target)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
