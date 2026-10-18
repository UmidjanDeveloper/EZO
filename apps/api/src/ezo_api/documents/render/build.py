"""DeckSpec → laid-out slides → QA → PPTX, PDF, HTML and PNG previews.

Automatic fix before export (reported, never silent): more than 6 bullets on a
slide — the extra ones move to the speaker notes. Text is shrunk to fit down to
a minimum size; anything that still does not fit is a QA error, not hidden.
An export is only reported as ready after the PPTX was re-opened and checked.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field

from pptx import Presentation

from ezo_api.documents.render.html import render_deck_html
from ezo_api.documents.render.layouts import layout_deck
from ezo_api.documents.render.pdf import render_deck_pdf
from ezo_api.documents.render.png import render_png
from ezo_api.documents.render.pptx_renderer import render_pptx
from ezo_api.documents.render.primitives import SlidePlan
from ezo_api.documents.render.qa import MAX_BULLETS, QAReport, run_qa
from ezo_api.documents.render.spec import DeckSpec
from ezo_api.documents.render.theme import Theme, choose_theme


class DeckBuildError(Exception):
    pass


@dataclass(slots=True)
class DeckBuild:
    deck: DeckSpec
    theme: Theme
    plans: list[SlidePlan]
    qa: QAReport
    fixes: list[str] = field(default_factory=list)
    pptx: bytes = b""
    pdf: bytes = b""
    html: bytes = b""
    previews: list[bytes] = field(default_factory=list)


def _move_extra_bullets(deck: DeckSpec, fixes: list[str]) -> None:
    for index, slide in enumerate(deck.slides, start=1):
        if len(slide.bullets) > MAX_BULLETS:
            extra = slide.bullets[MAX_BULLETS:]
            slide.bullets = slide.bullets[:MAX_BULLETS]
            slide.speaker_notes = (
                slide.speaker_notes + "\n\n" + "\n".join(f"• {b}" for b in extra)
            ).strip()
            fixes.append(
                f"Slide {index}: {len(extra)} extra bullet points moved to the speaker notes."
            )


def verify_pptx(data: bytes, expected_slides: int) -> None:
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:  # the file must open with a real PPTX reader
        raise DeckBuildError(
            f"The PPTX file could not be re-opened: {type(exc).__name__}."
        ) from exc
    count = len(presentation.slides)
    if count != expected_slides:
        raise DeckBuildError(f"The PPTX has {count} slides, expected {expected_slides}.")
    for number, slide in enumerate(presentation.slides, start=1):
        if len(slide.shapes) == 0:
            raise DeckBuildError(f"Slide {number} of the PPTX is empty.")


def check_deck(deck: DeckSpec, theme: Theme | None = None) -> QAReport:
    """Lay out a copy of the deck and run QA without rendering any file (fast)."""
    trial = deck.model_copy(deep=True)
    _move_extra_bullets(trial, [])
    theme = theme or choose_theme(trial.theme, trial.audience, trial.goal)
    return run_qa(trial, layout_deck(trial, theme))


def build_deck(
    deck: DeckSpec,
    theme: Theme | None = None,
    *,
    previews: bool = True,
    images: dict[int, bytes] | None = None,
) -> DeckBuild:
    """`images` maps a slide number (1-based) to a generated image for that slide."""
    theme = theme or choose_theme(deck.theme, deck.audience, deck.goal)
    fixes: list[str] = []
    _move_extra_bullets(deck, fixes)
    plans = layout_deck(deck, theme, images)
    report = run_qa(deck, plans)
    build = DeckBuild(deck=deck, theme=theme, plans=plans, qa=report, fixes=fixes)
    build.pptx = render_pptx(plans, deck.title, deck.language)
    verify_pptx(build.pptx, len(plans))
    build.pdf = render_deck_pdf(plans, deck.title)
    build.html = render_deck_html(plans, deck.title, deck.language)
    if previews:
        build.previews = [render_png(plan, 1280) for plan in plans]
    return build
