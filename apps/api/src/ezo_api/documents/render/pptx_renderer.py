"""Presentations as PPTX (python-pptx).

Everything stays editable in PowerPoint, Keynote and LibreOffice: real text boxes
with real bullets (no pre-broken lines), native charts and tables, speaker
notes. Text uses Arial, which is metric-compatible with the Liberation Sans used
for fitting, so the sizes EZO chose fit the same way.
"""

from __future__ import annotations

import io
import math

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_AXIS_CROSSES, XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from ezo_api.documents.fonts import OFFICE_FONT
from ezo_api.documents.render.primitives import (
    ChartBox,
    Ellipse,
    ImageBox,
    Line,
    Rect,
    SlidePlan,
    TableBox,
    TextBox,
    Wedge,
)
from ezo_api.documents.render.text import BULLET_INDENT, PARAGRAPH_GAP

EMU_PER_PT = 12700
_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
_CHARTS = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


def emu(points: float) -> Emu:
    return Emu(round(points * EMU_PER_PT))


def rgb(color: str) -> RGBColor:
    return RGBColor.from_string(color.lstrip("#").upper())


def _no_shadow(shape) -> None:  # type: ignore[no-untyped-def]
    """No theme effects: drop the shape's theme style reference (it adds a shadow in some apps)."""
    shape.shadow.inherit = False
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def _fill(shape, fill: str | None, line: str | None, line_width: float) -> None:  # type: ignore[no-untyped-def]
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    if line and line_width > 0:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    _no_shadow(shape)


def _rect(slide, item: Rect) -> None:  # type: ignore[no-untyped-def]
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if item.radius > 0 else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, emu(item.x), emu(item.y), emu(item.w), emu(item.h))
    if item.radius > 0:
        shape.adjustments[0] = min(0.5, item.radius / max(min(item.w, item.h), 1))
    _fill(shape, item.fill, item.line, item.line_width)


def _ellipse(slide, item: Ellipse) -> None:  # type: ignore[no-untyped-def]
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, emu(item.x), emu(item.y), emu(item.w), emu(item.h)
    )
    _fill(shape, item.fill, item.line, item.line_width)


def _line(slide, item: Line) -> None:  # type: ignore[no-untyped-def]
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, emu(item.x1), emu(item.y1), emu(item.x2), emu(item.y2)
    )
    connector.line.color.rgb = rgb(item.color)
    connector.line.width = Pt(item.width)
    style = connector._element.find(qn("p:style"))
    if style is not None:
        connector._element.remove(style)
    if item.arrow:
        ln = connector.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})


def _bullet(paragraph, size: float) -> None:  # type: ignore[no-untyped-def]
    p_pr = paragraph._p.get_or_add_pPr()
    indent = round(size * BULLET_INDENT * EMU_PER_PT)
    p_pr.set("marL", str(indent))
    p_pr.set("indent", str(-indent))
    font = etree.SubElement(p_pr, qn("a:buFont"))
    font.set("typeface", OFFICE_FONT)
    char = etree.SubElement(p_pr, qn("a:buChar"))
    char.set("char", "•")


def _text(slide, item: TextBox) -> None:  # type: ignore[no-untyped-def]
    box = slide.shapes.add_textbox(emu(item.x), emu(item.y), emu(item.w), emu(item.h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = _ANCHOR[item.valign]
    size = item.fitted_size or item.size
    for index, para in enumerate(item.paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = _ALIGN[item.align]
        paragraph.line_spacing = item.line_spacing
        if index < len(item.paragraphs) - 1:
            paragraph.space_after = Pt(size * PARAGRAPH_GAP)
        run = paragraph.add_run()
        run.text = para.text
        font = run.font
        font.name = OFFICE_FONT
        font.size = Pt(round(size * para.scale, 1))
        font.bold = para.bold or item.bold
        font.color.rgb = rgb(para.color or item.color)
        if para.bullet:
            _bullet(paragraph, size * para.scale)


def _chart(slide, item: ChartBox) -> None:  # type: ignore[no-untyped-def]
    data = CategoryChartData()
    data.categories = item.categories
    for name, values in item.series:
        clean = [None if math.isnan(v) else v for v in values[: len(item.categories)]]
        data.add_series(name or " ", clean, number_format="#,##0.##")
    frame = slide.shapes.add_chart(
        _CHARTS[item.kind], emu(item.x), emu(item.y), emu(item.w), emu(item.h), data
    )
    chart = frame.chart
    chart.has_title = False  # the slide title already says what the chart is
    chart.font.name = OFFICE_FONT
    chart.font.size = Pt(12)
    chart.font.color.rgb = rgb(item.muted)
    pie = item.kind in ("pie", "doughnut")
    chart.has_legend = pie or len(item.series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.RIGHT if pie else XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)
    plot = chart.plots[0]
    if pie:
        for index, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = rgb(item.colors[index % len(item.colors)])
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.show_percentage = True
        labels.show_value = False
        labels.number_format = "0%"
        labels.number_format_is_linked = False
        labels.font.size = Pt(12)
        labels.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        return
    for index, series in enumerate(plot.series):
        color = rgb(item.colors[index % len(item.colors)])
        if item.kind == "line":
            series.format.line.color.rgb = color
            series.format.line.width = Pt(3)
            series.smooth = False
            series.marker.format.fill.solid()
            series.marker.format.fill.fore_color.rgb = color
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color
    if item.kind in ("column", "bar"):
        plot.gap_width = 60
        if len(item.categories) * len(item.series) <= 16:
            plot.has_data_labels = True
            plot.data_labels.font.size = Pt(11)
            plot.data_labels.font.color.rgb = rgb(item.ink)
            plot.data_labels.number_format = "#,##0.##"
            plot.data_labels.number_format_is_linked = False
    value_axis = chart.value_axis
    value_axis.has_major_gridlines = True
    value_axis.major_gridlines.format.line.color.rgb = rgb(item.grid)
    value_axis.format.line.fill.background()
    value_axis.tick_labels.font.size = Pt(11)
    chart.category_axis.tick_labels.font.size = Pt(12)
    if item.kind == "bar":
        # PowerPoint draws bars bottom-up: keep the written order, value axis at the bottom.
        chart.category_axis.reverse_order = True
        chart.category_axis.crosses = XL_AXIS_CROSSES.MAXIMUM
    chart.category_axis.format.line.color.rgb = rgb(item.grid)


def _table(slide, item: TableBox) -> None:  # type: ignore[no-untyped-def]
    rows = [item.header, *item.rows]
    columns = max(len(item.header), 1)
    frame = slide.shapes.add_table(
        len(rows), columns, emu(item.x), emu(item.y), emu(item.w), emu(item.h)
    )
    table = frame.table
    table.first_row = True
    table.horz_banding = False
    for c in range(columns):
        table.columns[c].width = emu(item.w / columns)
    for r, row in enumerate(rows):
        table.rows[r].height = emu(item.h / len(rows))
        for c in range(columns):
            cell = table.cell(r, c)
            cell.text = row[c] if c < len(row) else ""
            cell.margin_left = cell.margin_right = emu(8)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(
                item.header_fill if r == 0 else (item.stripe if r % 2 == 0 else "#FFFFFF")
            )
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = OFFICE_FONT
                    run.font.size = Pt(item.size)
                    run.font.bold = r == 0
                    run.font.color.rgb = rgb(item.header_ink if r == 0 else item.ink)


def _image(slide, item: ImageBox) -> None:  # type: ignore[no-untyped-def]
    picture = slide.shapes.add_picture(
        io.BytesIO(item.data), emu(item.x), emu(item.y), emu(item.w), emu(item.h)
    )
    if item.alt:
        picture._element.nvPicPr.cNvPr.set("descr", item.alt[:500])


def render_pptx(plans: list[SlidePlan], title: str, language: str) -> bytes:
    presentation = Presentation()
    first = plans[0]
    presentation.slide_width = emu(first.width)
    presentation.slide_height = emu(first.height)
    presentation.core_properties.title = title
    presentation.core_properties.author = "EZO"
    presentation.core_properties.language = language
    blank = presentation.slide_layouts[6]
    for plan in plans:
        slide = presentation.slides.add_slide(blank)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = rgb(plan.background)
        for item in plan.items:
            if isinstance(item, Rect):
                _rect(slide, item)
            elif isinstance(item, Ellipse):
                _ellipse(slide, item)
            elif isinstance(item, Line):
                _line(slide, item)
            elif isinstance(item, TextBox):
                _text(slide, item)
            elif isinstance(item, ChartBox):
                _chart(slide, item)
            elif isinstance(item, TableBox):
                _table(slide, item)
            elif isinstance(item, ImageBox):
                _image(slide, item)
            elif isinstance(item, Wedge):  # only produced by chart expansion, never on PPTX slides
                continue
        if plan.notes:
            slide.notes_slide.notes_text_frame.text = plan.notes
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
