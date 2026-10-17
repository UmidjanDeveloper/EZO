"""Tables → XLSX (openpyxl) or CSV. Numbers stay numbers, headers are styled and
frozen, columns are sized to their content."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ezo_api.documents.fonts import OFFICE_FONT

_NUMBER = re.compile(r"^-?\d{1,15}([.,]\d+)?$")
_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
# Cells starting with these are treated as formulas by Excel; keep text as text.
_FORMULA_START = ("=", "+", "-", "@")


@dataclass(slots=True)
class Sheet:
    name: str
    columns: list[str]
    rows: list[list[Any]] = field(default_factory=list)


def coerce(value: str | float | int | None) -> str | float | int | None:
    """Turn numeric-looking text into numbers (accepts '1 200,5' style); leave the rest."""
    if not isinstance(value, str):
        return value
    text = value.strip().replace(" ", " ")
    compact = text.replace(" ", "")
    if _NUMBER.match(compact):
        number = float(compact.replace(",", "."))
        return (
            int(number)
            if number.is_integer() and "." not in compact and "," not in compact
            else number
        )
    return text


def safe_text(value: str) -> str:
    """Neutralise spreadsheet formula injection in text cells."""
    return (
        "'" + value
        if value.startswith(_FORMULA_START) and not _NUMBER.match(value.replace(" ", ""))
        else value
    )


def sheet_title(name: str, used: set[str]) -> str:
    title = _INVALID_SHEET_CHARS.sub(" ", name).strip()[:31] or "Sheet"
    base, counter = title, 2
    while title.lower() in used:
        suffix = f" ({counter})"
        title = base[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(title.lower())
    return title


def render_xlsx(sheets: list[Sheet], title: str = "") -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = title
    workbook.properties.creator = "EZO"
    header_fill = PatternFill("solid", fgColor="0B7A72")
    used: set[str] = set()
    for sheet in sheets or [Sheet("Sheet", [])]:
        ws = workbook.create_sheet(sheet_title(sheet.name, used))
        ws.append(sheet.columns)
        for cell in ws[1]:
            cell.font = Font(name=OFFICE_FONT, bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in sheet.rows:
            values = [coerce(v) for v in row]
            ws.append([safe_text(v) if isinstance(v, str) else v for v in values])
        ws.freeze_panes = "A2"
        for index, name in enumerate(sheet.columns, start=1):
            longest = max(
                [len(str(name))]
                + [
                    len(str(r[index - 1]))
                    for r in sheet.rows
                    if index - 1 < len(r) and r[index - 1] is not None
                ]
            )
            ws.column_dimensions[get_column_letter(index)].width = min(max(10, longest + 2), 60)
        if sheet.columns and sheet.rows:
            ws.auto_filter.ref = ws.dimensions
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_csv(sheet: Sheet) -> bytes:
    """UTF-8 with BOM so Excel opens Uzbek/Russian text correctly."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(sheet.columns)
    for row in sheet.rows:
        writer.writerow(
            [safe_text(str(v)) if isinstance(v, str) else ("" if v is None else v) for v in row]
        )
    return buffer.getvalue().encode("utf-8-sig")
