"""Read XLSX/CSV and describe the data (types, missing values, statistics, top values)
without pandas. The summary is what the model sees; it never invents numbers."""

from __future__ import annotations

import csv
import io
import statistics
from collections import Counter
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook

from ezo_api.documents.spreadsheet.writer import Sheet, coerce
from ezo_api.knowledge.parsers.text import decode_text

MAX_ROWS = 50_000


@dataclass(slots=True)
class ColumnSummary:
    name: str
    kind: str  # number | text | empty | mixed
    filled: int
    missing: int
    distinct: int
    stats: dict[str, float] = field(default_factory=dict)
    top: list[tuple[str, int]] = field(default_factory=list)


@dataclass(slots=True)
class SheetSummary:
    name: str
    rows: int
    columns: list[ColumnSummary]
    truncated: bool = False


def read_sheets(data: bytes, filename: str) -> list[Sheet]:
    if filename.lower().endswith(".csv"):
        text = decode_text(data)
        reader = csv.reader(io.StringIO(text))
        records = list(reader)
        if not records:
            return [Sheet("CSV", [])]
        header = [c.strip() for c in records[0]]
        return [Sheet("CSV", header, [list(r) for r in records[1 : MAX_ROWS + 1]])]
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    for ws in workbook.worksheets:
        values = ws.iter_rows(values_only=True)
        first = next(values, None)
        if first is None:
            continue
        columns = [
            str(c).strip() if c is not None else f"Column {i + 1}" for i, c in enumerate(first)
        ]
        rows: list[list[Any]] = []
        for row in values:
            if len(rows) >= MAX_ROWS:
                break
            if any(cell is not None and str(cell).strip() for cell in row):
                rows.append(list(row))
        sheets.append(Sheet(ws.title, columns, rows))
    workbook.close()
    return sheets


def summarize_sheet(sheet: Sheet) -> SheetSummary:
    columns = []
    for index, name in enumerate(sheet.columns):
        values = [coerce(row[index]) if index < len(row) else None for row in sheet.rows]
        present = [v for v in values if v not in (None, "")]
        numbers = [
            float(v) for v in present if isinstance(v, int | float) and not isinstance(v, bool)
        ]
        texts = [str(v) for v in present if not isinstance(v, int | float) or isinstance(v, bool)]
        if not present:
            kind = "empty"
        elif len(numbers) == len(present):
            kind = "number"
        elif not numbers:
            kind = "text"
        else:
            kind = "mixed"
        summary = ColumnSummary(
            name=name,
            kind=kind,
            filled=len(present),
            missing=len(values) - len(present),
            distinct=len({str(v) for v in present}),
        )
        if numbers:
            summary.stats = {
                "sum": round(sum(numbers), 4),
                "mean": round(statistics.fmean(numbers), 4),
                "min": min(numbers),
                "max": max(numbers),
                "median": round(statistics.median(numbers), 4),
            }
        if texts:
            summary.top = Counter(texts).most_common(5)
        columns.append(summary)
    return SheetSummary(sheet.name, len(sheet.rows), columns, truncated=len(sheet.rows) >= MAX_ROWS)


def describe(summaries: list[SheetSummary]) -> str:
    """A compact, factual description for the model (and the user's report)."""
    lines = []
    for sheet in summaries:
        lines.append(
            f"Sheet “{sheet.name}”: {sheet.rows} rows, {len(sheet.columns)} columns"
            + (" (first 50,000 rows only)" if sheet.truncated else "")
        )
        for column in sheet.columns:
            detail = (
                f"- {column.name}: {column.kind}, {column.filled} filled, "
                f"{column.missing} empty, {column.distinct} distinct"
            )
            if column.stats:
                s = column.stats
                detail += (
                    f"; sum {s['sum']:g}, mean {s['mean']:g}, median {s['median']:g}, "
                    f"min {s['min']:g}, max {s['max']:g}"
                )
            if column.top:
                detail += "; most common: " + ", ".join(
                    f"{value} ({count})" for value, count in column.top[:3]
                )
            lines.append(detail)
    return "\n".join(lines)
