"""Deterministic spreadsheet operations (no model): sort, filter, remove duplicates,
remove empty rows, totals and group-by summaries.

Requests are understood in Uzbek, Russian and English by finding the sheet's real
column names in the text, so EZO never acts on a column that does not exist. All
numbers in the result are computed by EZO and written as values; the reply says so.

Formulas: workbooks are read with the results Excel saved (`data_only`). Formula
cells whose result was never saved are reported, never "calculated" by EZO. Macros
are never executed (openpyxl does not run them; .xlsm files are not accepted).
"""

from __future__ import annotations

import io
import re
from collections import OrderedDict
from dataclasses import dataclass, field
from typing import Any, Literal

from openpyxl import load_workbook

from ezo_api.core.text import normalize, quoted
from ezo_api.documents.spreadsheet.writer import Sheet, coerce

Op = Literal["sort", "filter", "dedupe", "drop_empty", "totals", "group"]
Agg = Literal["sum", "count", "mean", "min", "max"]


class SheetOpError(Exception):
    def __init__(self, message: str) -> None:
        self.message = message
        super().__init__(message)


@dataclass(slots=True)
class SheetStep:
    op: Op
    column: str | None = None
    descending: bool = False
    comparison: str | None = None  # = != > < >= <= contains
    value: str | None = None
    columns: list[str] = field(default_factory=list)
    agg: Agg = "sum"


@dataclass(slots=True)
class OpsResult:
    sheet: Sheet
    report: list[str]
    warnings: list[str] = field(default_factory=list)


# --------------------------------------------------------------------------- formulas


def formula_report(data: bytes, filename: str) -> dict[str, Any]:
    """Formula cells and how many have no saved result (they read as empty)."""
    if filename.lower().endswith(".csv"):
        return {"formulas": 0, "without_saved_result": 0}
    formulas = workbook_formulas = load_workbook(io.BytesIO(data), read_only=True)
    values = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    total = missing = 0
    examples: list[str] = []
    try:
        for ws_f, ws_v in zip(formulas.worksheets, values.worksheets, strict=True):
            for row_f, row_v in zip(ws_f.iter_rows(), ws_v.iter_rows(), strict=False):
                for cell_f, cell_v in zip(row_f, row_v, strict=False):
                    if isinstance(cell_f.value, str) and cell_f.value.startswith("="):
                        total += 1
                        if cell_v.value is None:
                            missing += 1
                            if len(examples) < 5:
                                examples.append(f"{ws_f.title}!{cell_f.coordinate}")
    finally:
        workbook_formulas.close()
        values.close()
    return {"formulas": total, "without_saved_result": missing, "examples": examples}


# --------------------------------------------------------------------------- parsing

_SORT = re.compile(r"\b(sarala|saralab|tartibla|tartiblab|sort|отсортир|сортир|упорядоч)")
_DESC = re.compile(
    r"\b(kamayish|kattadan kichik|descending|desc\b|по убыванию|убыван|eng kattasi|teskari)"
)
_DEDUPE = re.compile(
    r"\b(dublikat\w*|takror\w*\s+(qator\w*\s+)?(o'chir|olib tashla)|"
    r"remove duplicates|duplicates|дубл\w*|повтор\w*)"
)
_EMPTY = re.compile(r"\b(bo'sh qator\w*|empty rows|blank rows|пуст\w* строк\w*)")
_TOTALS = re.compile(r"\b(jami\w*|yig'indi\w*|totals?\b|итог\w*|сумм\w* по столбц\w*)")
_GROUP = re.compile(r"\b(guruhla\w*|group(ed)? by|сгруппир\w*|группир\w*|по группам)")
_FILTER = re.compile(
    r"\b(faqat|qoldir|filtr\w*|filter|only rows|keep rows|where|только|оставь|фильтр\w*)"
)
_AGG: list[tuple[Agg, re.Pattern[str]]] = [
    ("mean", re.compile(r"\b(o'rtacha|average|mean|средн\w*)")),
    ("count", re.compile(r"\b(soni|nechta|count|количеств\w*|число)")),
    ("max", re.compile(r"\b(eng katta|maximum|max\b|максимум\w*)")),
    ("min", re.compile(r"\b(eng kichik|minimum|min\b|минимум\w*)")),
    ("sum", re.compile(r"\b(yig'indi\w*|jami\w*|sum\b|total|сумм\w*)")),
]
_COMPARE = re.compile(r"(>=|<=|!=|>|<|=)\s*([\d\s.,]+)")
_WORD_COMPARE: list[tuple[str, re.Pattern[str]]] = [
    (">", re.compile(r"(?:dan|than|больше)\s*(?:katta|ko'p|more|greater)?\s*$")),
]


def mentioned_columns(text: str, columns: list[str]) -> list[str]:
    """Columns named in the text, in the order they appear (longest names first win)."""
    norm = normalize(text)
    found: list[tuple[int, str]] = []
    taken: list[tuple[int, int]] = []
    for column in sorted(columns, key=len, reverse=True):
        name = normalize(str(column)).strip()
        if len(name) < 2:
            continue
        for match in re.finditer(rf"(?<!\w){re.escape(name)}(?!\w)", norm):
            span = (match.start(), match.end())
            if not any(a < span[1] and span[0] < b for a, b in taken):
                taken.append(span)
                found.append((match.start(), column))
                break
    return [column for _, column in sorted(found)]


def parse_sheet_request(text: str, columns: list[str]) -> list[SheetStep]:
    norm = normalize(text)
    named = mentioned_columns(text, columns)
    steps: list[SheetStep] = []
    if _EMPTY.search(norm):
        steps.append(SheetStep("drop_empty"))
    if _DEDUPE.search(norm):
        steps.append(SheetStep("dedupe", columns=named if named and "bo'yicha" in norm else []))
    if _FILTER.search(norm) and named:
        values = quoted(text)
        compare = _COMPARE.search(norm)
        if values:
            steps.append(SheetStep("filter", column=named[0], comparison="=", value=values[0]))
        elif compare:
            steps.append(
                SheetStep(
                    "filter",
                    column=named[0],
                    comparison=compare.group(1),
                    value=compare.group(2).strip(),
                )
            )
    if _GROUP.search(norm) and named:
        agg: Agg = next((a for a, p in _AGG if p.search(norm)), "sum")
        steps.append(SheetStep("group", column=named[0], columns=named[1:], agg=agg))
    elif _SORT.search(norm) and named:
        steps.append(SheetStep("sort", column=named[0], descending=bool(_DESC.search(norm))))
    if _TOTALS.search(norm) and not any(s.op == "group" for s in steps):
        steps.append(SheetStep("totals", columns=named))
    return steps


# --------------------------------------------------------------------------- running


def _number(value: Any) -> float | None:
    value = coerce(value) if isinstance(value, str) else value
    if isinstance(value, bool) or value is None:
        return None
    return float(value) if isinstance(value, int | float) else None


def _index(sheet: Sheet, column: str | None) -> int:
    if column is None or column not in sheet.columns:
        raise SheetOpError(f"The column “{column}” does not exist in “{sheet.name}”.")
    return sheet.columns.index(column)


def _cell(row: list[Any], index: int) -> Any:
    return row[index] if index < len(row) else None


def _fmt(number: float) -> int | float:
    return int(number) if float(number).is_integer() else round(number, 6)


def run_steps(sheet: Sheet, steps: list[SheetStep]) -> OpsResult:
    if not steps:
        raise SheetOpError("No spreadsheet operation was recognised.")
    rows = [list(r) for r in sheet.rows]
    columns = list(sheet.columns)
    report: list[str] = []
    warnings: list[str] = []
    current = Sheet(sheet.name, columns, rows)
    for step in steps:
        if step.op == "drop_empty":
            before = len(current.rows)
            current.rows = [
                r for r in current.rows if any(v is not None and str(v).strip() for v in r)
            ]
            report.append(f"Removed {before - len(current.rows)} empty row(s).")
        elif step.op == "dedupe":
            keys = [_index(current, c) for c in step.columns] if step.columns else None
            seen: set[tuple[str, ...]] = set()
            kept: list[list[Any]] = []
            for row in current.rows:
                row_key = tuple(
                    str(_cell(row, i)).strip().lower()
                    for i in (keys if keys is not None else range(len(current.columns)))
                )
                if row_key not in seen:
                    seen.add(row_key)
                    kept.append(row)
            removed = len(current.rows) - len(kept)
            current.rows = kept
            suffix = f" (by {', '.join(step.columns)})" if step.columns else ""
            report.append(
                f"Removed {removed} duplicate row(s){suffix}; the first occurrence was kept."
            )
        elif step.op == "filter":
            index = _index(current, step.column)
            before = len(current.rows)
            current.rows = [r for r in current.rows if _matches(_cell(r, index), step)]
            report.append(
                f"Kept {len(current.rows)} of {before} row(s) where {step.column} "
                f"{step.comparison} {step.value}."
            )
        elif step.op == "sort":
            index = _index(current, step.column)
            maybe = [_number(_cell(r, index)) for r in current.rows]
            numeric = sum(n is not None for n in maybe) >= max(1, len(maybe) // 2)

            def sort_key(
                row: list[Any], i: int = index, numeric: bool = numeric
            ) -> tuple[int, Any]:
                value = _cell(row, i)
                if numeric:
                    n = _number(value)
                    return (0, n) if n is not None else (1, 0)
                text = str(value).strip().lower() if value is not None else ""
                return (0, text) if text else (1, "")

            blanks = [r for r in current.rows if sort_key(r)[0] == 1]
            filled = sorted(
                (r for r in current.rows if sort_key(r)[0] == 0),
                key=sort_key,
                reverse=step.descending,
            )
            current.rows = filled + blanks
            report.append(
                f"Sorted by {step.column} ({'descending' if step.descending else 'ascending'}"
                f"{', as numbers' if numeric else ', alphabetically'}); empty cells last."
            )
        elif step.op == "totals":
            targets = step.columns or [
                c
                for i, c in enumerate(current.columns)
                if any(_number(_cell(r, i)) is not None for r in current.rows)
            ]
            total_row: list[Any] = [""] * len(current.columns)
            parts: list[str] = []
            for column in targets:
                col = _index(current, column)
                values = [_number(_cell(r, col)) for r in current.rows]
                present = [v for v in values if v is not None]
                if not present:
                    warnings.append(f"{column} has no numbers, so it has no total.")
                    continue
                skipped = sum(
                    1
                    for v, r in zip(values, current.rows, strict=True)
                    if v is None and _cell(r, col) not in (None, "")
                )
                if skipped:
                    warnings.append(f"{column}: {skipped} non-numeric cell(s) were not added.")
                total_row[col] = _fmt(sum(present))
                parts.append(f"{column} = {total_row[col]}")
            if parts:
                label = next((n for n in range(len(current.columns)) if total_row[n] == ""), None)
                if label is not None:
                    total_row[label] = "Jami / Total"
                current.rows.append(total_row)
                report.append("Added a total row (computed by EZO): " + "; ".join(parts) + ".")
        elif step.op == "group":
            by = _index(current, step.column)
            targets = step.columns or [
                c
                for i, c in enumerate(current.columns)
                if i != by and any(_number(_cell(r, i)) is not None for r in current.rows)
            ]
            indexes = [_index(current, c) for c in targets]
            groups: OrderedDict[str, list[list[Any]]] = OrderedDict()
            for row in current.rows:
                group = str(_cell(row, by) if _cell(row, by) is not None else "").strip()
                groups.setdefault(group or "(empty)", []).append(row)
            new_rows: list[list[Any]] = []
            for group, members in groups.items():
                out: list[Any] = [group]
                for col in indexes:
                    found = [n for n in (_number(_cell(r, col)) for r in members) if n is not None]
                    out.append(_aggregate(found, step.agg, len(members)))
                new_rows.append(out)
            current = Sheet(
                f"{sheet.name} by {step.column}"[:31],
                [str(step.column)] + [f"{c} ({step.agg})" for c in targets],
                new_rows,
            )
            report.append(
                f"Grouped {sum(len(m) for m in groups.values())} row(s) by {step.column} into "
                f"{len(groups)} group(s); {step.agg} of {', '.join(targets) or 'rows'} "
                "(computed by EZO)."
            )
    return OpsResult(current, report, warnings)


def _aggregate(numbers: list[float], agg: Agg, members: int) -> int | float | str:
    if agg == "count":
        return members
    if not numbers:
        return ""
    if agg == "sum":
        return _fmt(sum(numbers))
    if agg == "mean":
        return round(sum(numbers) / len(numbers), 4)
    return _fmt(max(numbers) if agg == "max" else min(numbers))


def _matches(value: Any, step: SheetStep) -> bool:
    op, wanted = step.comparison or "=", step.value or ""
    if op in (">", "<", ">=", "<=") or (_number(wanted) is not None and op in ("=", "!=")):
        left, right = _number(value), _number(wanted)
        if left is None or right is None:
            return op == "!="
        return {
            ">": left > right,
            "<": left < right,
            ">=": left >= right,
            "<=": left <= right,
            "=": left == right,
            "!=": left != right,
        }[op]
    text = normalize(str(value or "")).strip()
    target = normalize(wanted).strip()
    if op == "!=":
        return text != target
    if op == "contains":
        return target in text
    return text == target
