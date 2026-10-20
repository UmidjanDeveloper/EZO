"""Approved Word templates: the organization's own blank (letterhead, header, footer,
styles, margins), into which EZO writes the document body.

Where the body goes: at a paragraph containing the marker {{matn}} ({{text}} or {{текст}}
also work). The marker paragraph and everything after it in the body is replaced by the
document; everything before it (a letterhead in the body) is kept. Without a marker the
whole body is kept and the document follows it. Headers, footers, styles and page
settings always stay as in the template.
"""

from __future__ import annotations

import io
import re
from typing import Any

from docx import Document
from docx.oxml.ns import qn

from ezo_api.documents.isolation import ISOLATION_MESSAGE, native_fixtures_allowed
from ezo_api.documents.office_convert import OfficeConvertError, external_links

MARKER = re.compile(r"\{\{\s*(matn|text|текст|матн)\s*\}\}", re.IGNORECASE)
MAX_TEMPLATE_BYTES = 10 * 1024 * 1024


class TemplateError(Exception):
    def __init__(self, message: str) -> None:
        self.message = message
        super().__init__(message)


def check_template(data: bytes) -> dict[str, Any]:
    """Validate an uploaded template (Word only, no external links). Returns facts about it."""
    if not native_fixtures_allowed():
        raise TemplateError(ISOLATION_MESSAGE)
    if len(data) > MAX_TEMPLATE_BYTES:
        raise TemplateError("The template is larger than 10 MB.")
    try:
        linked = external_links(data)
    except OfficeConvertError as exc:
        raise TemplateError(exc.message) from exc
    if linked:
        raise TemplateError(
            "The template links to files outside itself ("
            + ", ".join(sorted(set(linked)))
            + "); embed them or remove the links."
        )
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several types for damaged files
        raise TemplateError("The template could not be opened as a Word document.") from exc
    has_marker = any(MARKER.search(p.text) for p in document.paragraphs)
    header_text = " ".join(
        p.text for s in document.sections for p in s.header.paragraphs if p.text.strip()
    )
    return {"has_marker": has_marker, "header_text": header_text[:300]}


def open_template(data: bytes) -> Any:
    """The template as a python-docx Document, with the body from the marker on removed."""
    if not native_fixtures_allowed():
        raise TemplateError(ISOLATION_MESSAGE)
    document = Document(io.BytesIO(data))
    body = document.element.body
    children = [child for child in body.iterchildren() if child.tag != qn("w:sectPr")]
    marker_index = next(
        (
            index
            for index, child in enumerate(children)
            if child.tag == qn("w:p")
            and MARKER.search("".join(t.text or "" for t in child.iter(qn("w:t"))))
        ),
        None,
    )
    if marker_index is not None:
        for child in children[marker_index:]:
            body.remove(child)
    return document
