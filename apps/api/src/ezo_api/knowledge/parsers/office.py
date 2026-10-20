"""Word (DOCX) and Excel (XLSX) parsers."""

from __future__ import annotations

import io
import zipfile
from typing import Any

import docx
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from ezo_api.documents.ooxml_limits import OfficeInputLimit, validate_ooxml
from ezo_api.knowledge.parsers.base import MAX_TABLE_ROWS, ParsedDocument, ParseError, Section


def _iter_blocks(document: DocxDocument) -> Any:
    """Paragraphs and tables in document order."""
    body = document.element.body
    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def parse_docx(data: bytes, filename: str) -> ParsedDocument:
    _preflight(data)
    try:
        document = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ParseError("The Word document could not be read.") from exc

    sections: list[Section] = []
    current: list[str] = []
    heading: str | None = None

    def flush() -> None:
        if current:
            sections.append(Section(text="\n".join(current), heading=heading))
            current.clear()

    for block in _iter_blocks(document):
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if not text:
                continue
            style = (block.style.name if block.style is not None else "") or ""
            if style.lower().startswith(("heading", "title", "заголовок", "sarlavha")):
                flush()
                heading = text
            current.append(text)
        else:
            for row in block.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    current.append(" | ".join(cells))
    flush()
    title = (document.core_properties.title or "").strip() or None
    return ParsedDocument(sections=sections, title=title, metadata={"format": "docx"})


def parse_xlsx(data: bytes, filename: str) -> ParsedDocument:
    _preflight(data)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ParseError("The Excel file could not be read.") from exc

    sections: list[Section] = []
    truncated = False
    sheet_names: list[str] = []
    try:
        for sheet in workbook.worksheets:
            sheet_names.append(sheet.title)
            lines: list[str] = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_TABLE_ROWS:
                    truncated = True
                    break
                cells = ["" if value is None else str(value) for value in row]
                if any(cell.strip() for cell in cells):
                    lines.append(" | ".join(cells).rstrip(" |"))
            if lines:
                sections.append(Section(text="\n".join(lines), heading=f"Sheet: {sheet.title}"))
    finally:
        workbook.close()
    return ParsedDocument(
        sections=sections,
        metadata={"format": "xlsx", "sheets": sheet_names, "truncated": truncated},
    )


def _preflight(data: bytes) -> None:
    try:
        validate_ooxml(data)
    except OfficeInputLimit as exc:
        raise ParseError(str(exc), code=exc.code) from exc
