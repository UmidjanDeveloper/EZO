"""Run under the test image’s tini ENTRYPOINT, not `--entrypoint python` (which skips init).

CI mounts this file into the immutable image. Real conversions plus deliberately
orphaned native children make a missing PID-1 reaper observable on every run.
"""

import io
import os
import signal
import sys
import time
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from ezo_api.documents.isolation import trusted_native_regression_fixtures
from ezo_api.documents.office_convert import OfficeConverter
from ezo_api.documents.process_limits import run_native

with trusted_native_regression_fixtures():
    assert Path("/proc/1/comm").read_text().strip() == "tini"
    assert os.getuid() == 10001
    converter = OfficeConverter.discover()
    assert converter is not None
    word = Document()
    word.add_paragraph("O‘zbekiston — Ўзбекистон — Learning")
    a = io.BytesIO()
    word.save(a)
    book = Workbook()
    book.active.append(["Jami", 42])
    b = io.BytesIO()
    book.save(b)

    def zombies():
        result = []
        for file in Path("/proc").glob("[0-9]*/stat"):
            try:
                if file.read_text().rsplit(")", 1)[1].split()[0] == "Z":
                    result.append(file.parent.name)
            except FileNotFoundError:
                pass
        return result

    assert not zombies()
    for index in range(6):
        assert converter.to_pdf_sync(a.getvalue(), "docx").pages == 1
        assert converter.to_pdf_sync(b.getvalue(), "xlsx").pages == 1
        # Same lifecycle as a native leader exiting before its descendant; no setsid.
        run_native(
            [
                sys.executable,
                "-c",
                "import os,time; pid=os.fork(); time.sleep(30) if pid == 0 else None",
            ],
            timeout=3,
            file_limit=4096,
            stdout_limit=4096,
        )
        deadline = time.monotonic() + 3
        stable = 0
        observed = []
        while time.monotonic() < deadline:
            time.sleep(0.05)
            observed = zombies()
            stable = stable + 1 if not observed else 0
            if stable >= 5:
                break
        assert stable >= 5, f"iteration {index}: zombies accumulated: {observed}"
    print("12 conversions + 6 orphan lifecycles: no zombies", flush=True)
    # Host can use --shutdown-probe and docker stop to verify forwarding and clean exit.
    if "--shutdown-probe" in sys.argv:
        signal.signal(signal.SIGTERM, lambda *_: sys.exit(0))
        print("shutdown probe ready", flush=True)
        while True:
            signal.pause()
