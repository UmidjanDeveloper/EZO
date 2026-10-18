"""A Word document with the formatting EZO must keep when it edits text in place."""

from __future__ import annotations

import io

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from PIL import Image


def _png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (60, 30), (20, 90, 200)).save(buffer, format="PNG")
    return buffer.getvalue()


def _hyperlink(paragraph, text: str, url: str) -> None:  # type: ignore[no-untyped-def]
    rel = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    rpr.append(style)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def rich_docx() -> bytes:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(13)
    section = doc.sections[0]
    section.header.paragraphs[0].text = "MCHJ «Ezgu Kelajak»"
    footer = section.footer.paragraphs[0]
    footer.add_run("Sahifa ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    doc.add_heading("Ijara shartnomasi", level=1)
    p = doc.add_paragraph("Ijara narxi oyiga ")
    bold = p.add_run("4 500 000 so'm")
    bold.bold = True
    p.add_run(" bo'lib, to'lov har oyning 5-sanasigacha amalga oshiriladi.")
    p2 = doc.add_paragraph("Batafsil ma'lumot ")
    _hyperlink(p2, "kompaniya saytida", "https://example.uz/ijara")
    p2.add_run(" keltirilgan.")
    p3 = doc.add_paragraph()
    red = p3.add_run("Muhim: ")
    red.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    red.italic = True
    p3.add_run("shartnoma ikki nusxada tuziladi.")
    doc.add_paragraph("Ijarachi binoni toza saqlaydi.", style="List Number")
    doc.add_paragraph("Ijarachi kommunal to'lovlarni o'z vaqtida to'laydi.", style="List Number")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Tomon"
    table.cell(0, 1).text = "Imzo"
    table.cell(1, 0).text = "Ijarachi"
    table.cell(1, 1).text = "________"
    doc.add_picture(io.BytesIO(_png()))
    doc.add_paragraph("Shartnoma imzolangan kundan kuchga kiradi.")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
