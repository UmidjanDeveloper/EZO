"""Builders for real test files (PDF, DOCX, XLSX, images) generated on the fly."""

from __future__ import annotations

import io

from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

PNG_1PX = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00"
    b"\x1f\x15\xc4\x89\x00\x00\x00\rIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N\x00\x00"
    b"\x00\x00IEND\xaeB`\x82"
)


def make_pdf(pages: list[str], title: str | None = None) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    if title:
        pdf.setTitle(title)
    for page in pages:
        y = 800
        for line in page.split("\n"):
            pdf.drawString(50, y, line)
            y -= 16
        pdf.showPage()
    pdf.save()
    return buffer.getvalue()


def make_docx(heading: str, paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = Document()
    document.core_properties.title = heading
    document.add_heading(heading, level=1)
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx(rows: list[list[object]], sheet: str = "Sales") -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = sheet
    for row in rows:
        worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
