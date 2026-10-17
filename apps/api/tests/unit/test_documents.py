import csv
import io

import pytest
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from ezo_api.documents.compare import compare_texts
from ezo_api.documents.docx.writer import render_docx
from ezo_api.documents.markdown_io import parse_inline, parse_markdown, to_markdown, to_text
from ezo_api.documents.model import Heading, ListBlock, Paragraph, Quote, Table, word_count
from ezo_api.documents.pdf.writer import render_pdf
from ezo_api.documents.spreadsheet.analysis import describe, read_sheets, summarize_sheet
from ezo_api.documents.spreadsheet.writer import Sheet, coerce, render_csv, render_xlsx
from tests.office import office_to_pdf, requires_office

MARKDOWN = """# Oʻzbekiston raqamli loyihasi

Bu **muhim** hujjat. *Muddat*: 2026-yil, [manba](https://gov.uz).

## Maqsadlar

- Hujjatlarni raqamlashtirish
- Xizmatlarni tezlashtirish

1. Tahlil
2. Joriy etish

| Koʻrsatkich | Qiymat |
|---|---|
| Xizmatlar | 120 |
| Foydalanuvchilar | 1 500 |

> Ўзбекистон — келажаги буюк давлат.
"""


def test_markdown_parsing_structure() -> None:
    spec = parse_markdown(MARKDOWN)
    assert spec.title == "Oʻzbekiston raqamli loyihasi"
    kinds = [type(b) for b in spec.blocks]
    assert kinds == [Paragraph, Heading, ListBlock, ListBlock, Table, Quote]
    lists = [b for b in spec.blocks if isinstance(b, ListBlock)]
    assert [b.ordered for b in lists] == [False, True]
    table = next(b for b in spec.blocks if isinstance(b, Table))
    assert table.header == ("Koʻrsatkich", "Qiymat") and table.rows[1] == (
        "Foydalanuvchilar",
        "1 500",
    )
    assert word_count(spec) > 20


def test_inline_runs_and_safe_links() -> None:
    runs = parse_inline("**qalin** va *kursiv* `kod` [ok](https://a.uz) [bad](javascript:alert(1))")
    assert (runs[0].text, runs[0].bold) == ("qalin", True)
    assert any(r.italic and r.text == "kursiv" for r in runs)
    assert any(r.code and r.text == "kod" for r in runs)
    assert next(r for r in runs if r.text == "ok").link == "https://a.uz"
    assert all(r.link is None for r in runs if r.text == "bad")


def test_markdown_round_trip_and_text() -> None:
    spec = parse_markdown(MARKDOWN)
    again = parse_markdown(to_markdown(spec))
    assert again.title == spec.title and len(again.blocks) == len(spec.blocks)
    text = to_text(spec)
    assert "• Hujjatlarni raqamlashtirish" in text and "1. Tahlil" in text


def test_docx_has_real_structure() -> None:
    data = render_docx(parse_markdown(MARKDOWN))
    document = Document(io.BytesIO(data))
    styles = [(p.style.name, p.text) for p in document.paragraphs if p.text]
    assert ("Title", "Oʻzbekiston raqamli loyihasi") in styles
    assert ("Heading 1", "Maqsadlar") in styles
    assert ("List Bullet", "Hujjatlarni raqamlashtirish") in styles
    assert ("List Number", "Tahlil") in styles
    table = document.tables[0]
    assert table.rows[0].cells[0].text == "Koʻrsatkich" and table.rows[2].cells[1].text == "1 500"
    assert document.core_properties.title == "Oʻzbekiston raqamli loyihasi"


def test_pdf_embeds_unicode_text() -> None:
    data = render_pdf(parse_markdown(MARKDOWN))
    assert data.startswith(b"%PDF-")
    text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(data)).pages)
    assert "Maqsadlar" in text and "Ўзбекистон" in text and "1 500" in text


def test_xlsx_numbers_headers_and_formula_safety() -> None:
    sheet = Sheet(
        "Sotuvlar: 2026/Q1",
        ["Mahsulot", "Soni", "Narx"],
        [["Choy", "12", "8 500,5"], ['=HYPERLINK("x")', 3, None]],
    )
    workbook = load_workbook(io.BytesIO(render_xlsx([sheet], title="Hisobot")))
    ws = workbook.worksheets[0]
    assert ws.title == "Sotuvlar  2026 Q1"  # invalid sheet characters removed
    assert [c.value for c in ws[1]] == ["Mahsulot", "Soni", "Narx"] and ws["A1"].font.bold
    assert ws["B2"].value == 12 and ws["C2"].value == 8500.5
    assert ws["A3"].value == '\'=HYPERLINK("x")'  # never becomes a formula
    assert ws.freeze_panes == "A2"


def test_csv_has_bom_for_excel() -> None:
    data = render_csv(Sheet("x", ["Ism", "Shahar"], [["Aziz", "Toshkent"], ["@SUM", "Тошкент"]]))
    assert data.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows == [["Ism", "Shahar"], ["Aziz", "Toshkent"], ["'@SUM", "Тошкент"]]


def test_coerce() -> None:
    assert coerce("1 200") == 1200 and coerce("3,5") == 3.5 and coerce("-7") == -7
    assert coerce("12-uy") == "12-uy" and coerce(None) is None


def test_spreadsheet_analysis() -> None:
    data = render_xlsx(
        [
            Sheet(
                "Data",
                ["Viloyat", "Aholi", "Izoh"],
                [["Toshkent", 3000, ""], ["Samarqand", 4000, "x"], ["Toshkent", None, ""]],
            )
        ]
    )
    summary = summarize_sheet(read_sheets(data, "data.xlsx")[0])
    columns = {c.name: c for c in summary.columns}
    assert summary.rows == 3
    assert (
        columns["Aholi"].kind == "number"
        and columns["Aholi"].stats["sum"] == 7000
        and columns["Aholi"].missing == 1
    )
    assert columns["Viloyat"].top[0] == ("Toshkent", 2)
    text = describe([summary])
    assert "Aholi: number" in text and "sum 7000" in text

    csv_sheet = read_sheets("a;b\n1;2\n".encode(), "x.csv")[0]
    assert csv_sheet.columns == ["a;b"]  # comma CSV reader; semicolons stay text


def test_compare_texts() -> None:
    result = compare_texts(
        "Birinchi band.\n\nIkkinchi band.\n\nUchinchi band.",
        "Birinchi band.\n\nIkkinchi band yangilandi.\n\nToʻrtinchi band.\n\nBeshinchi.",
    )
    assert result.changed and not result.identical
    assert "changed" in result.summary()
    assert compare_texts("a\n\nb", "a\n\nb").identical


@requires_office
@pytest.mark.parametrize("kind", ["docx", "xlsx"])
def test_libreoffice_opens_generated_files(kind: str) -> None:
    if kind == "docx":
        pdf = office_to_pdf(render_docx(parse_markdown(MARKDOWN)), ".docx")
    else:
        pdf = office_to_pdf(render_xlsx([Sheet("A", ["x"], [[1]])]), ".xlsx")
    assert len(PdfReader(io.BytesIO(pdf)).pages) >= 1
