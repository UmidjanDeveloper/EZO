import io
import math

from pptx import Presentation
from pptx.oxml.ns import qn
from pypdf import PdfReader

from ezo_api.documents.render.build import build_deck
from ezo_api.documents.render.primitives import TextBox, TextPara
from ezo_api.documents.render.spec import DeckSpec, SlideSpec
from ezo_api.documents.render.text import fit
from ezo_api.documents.render.theme import THEMES, choose_theme, contrast
from tests.office import office_to_pdf, requires_office
from tests.sample_deck import SAMPLE_DECK, sample_deck


def deck(*slides: dict, **extra: object) -> DeckSpec:  # type: ignore[type-arg]
    return DeckSpec.model_validate({"title": "Test", "slides": list(slides), **extra})


def test_spec_normalises_model_output() -> None:
    slide = SlideSpec.model_validate(
        {
            "layout": "fancy_unknown",
            "title": "  Ko‘p   bo‘sh joy  ",
            "bullets": [f"band {i}" for i in range(12)],
            "chart": {
                "type": "column",
                "categories": ["a", "b"],
                "series": [{"name": "x", "values": ["1 200", "3,5"]}],
            },
            "edges": [{"from": "a", "to": "b"}, ["b", "c"], "bad"],
            "sources": ["1", "x", 3],
        }
    )
    assert slide.layout == "bullets" and slide.title == "Ko‘p bo‘sh joy"
    assert len(slide.bullets) == 8
    assert slide.chart is not None and slide.chart.series[0].values == [1200.0, 3.5]
    assert slide.edges == [("a", "b"), ("b", "c")]
    assert slide.sources == [1, 3]
    bad = SlideSpec.model_validate({"chart": {"series": [{"values": ["n/a"]}]}})
    assert bad.chart is not None and math.isnan(bad.chart.series[0].values[0])


def test_themes_are_readable_and_chosen_by_audience() -> None:
    for theme in THEMES.values():
        assert contrast(theme.ink, theme.background) >= 7, theme.name
        assert contrast(theme.muted, theme.background) >= 3.5, theme.name
    assert choose_theme("Tuman hokimiga taqdimot").name == "civic"
    assert choose_theme("investor pitch deck").name == "pitch"
    assert choose_theme("dars uchun").name == "paper"
    assert choose_theme("").name == "ezo"
    assert choose_theme("midnight").name == "midnight"


def test_text_fitting_shrinks_then_flags_overflow() -> None:
    short = fit(
        TextBox(0, 0, 400, 100, [TextPara("Qisqa matn")], size=30, color="#000", min_size=12)
    )
    assert short.fitted_size == 30 and not short.overflow
    long = fit(
        TextBox(0, 0, 400, 100, [TextPara("so‘z " * 60)], size=30, color="#000", min_size=12)
    )
    assert 12 <= long.fitted_size < 30
    impossible = fit(
        TextBox(0, 0, 200, 40, [TextPara("so‘z " * 200)], size=30, color="#000", min_size=12)
    )
    assert impossible.overflow and impossible.fitted_size == 12


def test_sample_deck_builds_every_format() -> None:
    build = build_deck(sample_deck())
    n = len(SAMPLE_DECK["slides"])  # type: ignore[arg-type]
    assert len(build.plans) == n and build.qa.passed, build.qa.to_dict()
    assert {i.code for i in build.qa.issues} == {"image_pending"}

    presentation = Presentation(io.BytesIO(build.pptx))
    assert len(presentation.slides) == n
    assert (
        presentation.slide_width == 960 * 12700
        and presentation.core_properties.title == SAMPLE_DECK["title"]
    )
    shapes = [shape for slide in presentation.slides for shape in slide.shapes]
    assert sum(1 for s in shapes if s.has_chart) == 5
    assert sum(1 for s in shapes if s.has_table) == 1
    # editable text with real bullets, no theme shadows
    assert any(
        s.has_text_frame and s.text_frame._txBody.find(".//" + qn("a:buChar")) is not None
        for s in shapes
    )
    assert all(s._element.find(qn("p:style")) is None for s in shapes if s.shape_type == 1)
    bar = next(s.chart for s in shapes if s.has_chart and s.chart.chart_type == 57)  # BAR_CLUSTERED
    assert bar.category_axis.reverse_order and not bar.has_title

    pdf = PdfReader(io.BytesIO(build.pdf))
    assert len(pdf.pages) == n
    assert "Xatirchi tumani" in pdf.pages[0].extract_text()

    html = build.html.decode()
    assert html.count('<section class="slide">') == n
    assert "Content-Security-Policy" in html and "http" not in html.split("<body>")[1].split(
        "<script>"
    )[0].replace("http://www.w3.org/2000/svg", "")

    assert len(build.previews) == n and build.previews[0][:8] == b"\x89PNG\r\n\x1a\n"


def test_qa_detects_real_problems() -> None:
    build = build_deck(
        deck(
            {"layout": "bullets"},
            {"layout": "bullets", "title": "Bir xil", "bullets": ["x"]},
            {"layout": "bullets", "title": "Bir xil"},
            {
                "layout": "chart",
                "title": "Grafik",
                "chart": {
                    "categories": ["a", "b", "c"],
                    "series": [{"name": "s", "values": [1, 2]}],
                },
            },
            {"layout": "chart", "title": "Bo‘sh grafik"},
            {"layout": "table", "title": "Jadval"},
        ),
        previews=False,
    )
    codes = {(i.slide, i.code) for i in build.qa.issues}
    assert (1, "empty_slide") in codes
    assert (3, "duplicate_title") in codes and (3, "title_only") in codes
    assert (4, "chart_data") in codes
    assert (5, "chart_missing") in codes and (6, "table_missing") in codes
    assert not build.qa.passed


def test_automatic_fixes_are_reported() -> None:
    many = [f"Muhim fikr raqami {i}, batafsil izoh bilan birga yozilgan" for i in range(8)]
    build = build_deck(
        deck({"layout": "bullets", "title": "Ko‘p fikr", "bullets": many}), previews=False
    )
    assert any("speaker notes" in fix for fix in build.fixes)
    assert len(build.deck.slides[0].bullets) == 6
    assert "Muhim fikr raqami 7" in build.deck.slides[0].speaker_notes


def test_overflow_is_reported_not_hidden() -> None:
    from ezo_api.documents.render.qa import run_qa

    build = build_deck(deck({"layout": "bullets", "title": "T", "bullets": ["a"]}), previews=False)
    body = next(i for i in build.plans[0].items if isinstance(i, TextBox) and i.role == "body")
    body.paragraphs = [TextPara("so‘z " * 1000)]
    fit(body)
    report = run_qa(build.deck, build.plans)
    assert not report.passed and any(i.code == "text_overflow" for i in report.issues)


def test_text_is_escaped_in_web_presentation() -> None:
    build = build_deck(
        deck({"layout": "bullets", "title": "<script>alert(1)</script>", "bullets": ["a & b"]}),
        previews=False,
    )
    html = build.html.decode()
    assert (
        "<script>alert(1)</script>" not in html and "&lt;script&gt;" in html and "a &amp; b" in html
    )


def test_four_by_three_and_dark_theme() -> None:
    build = build_deck(
        deck(
            {"layout": "kpi", "title": "KPI", "metrics": [{"value": "1", "label": "a"}]},
            aspect_ratio="4:3",
            theme="pitch",
        ),
        previews=False,
    )
    assert build.plans[0].width == 720 and build.theme.dark


@requires_office
def test_libreoffice_opens_the_presentation() -> None:
    pdf = office_to_pdf(build_deck(sample_deck(), previews=False).pptx, ".pptx")
    assert len(PdfReader(io.BytesIO(pdf)).pages) == len(SAMPLE_DECK["slides"])  # type: ignore[arg-type]
