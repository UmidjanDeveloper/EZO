"""Organization profiles without a model: formatting, approved templates with the
{{matn}} marker, file checks (logo, templates, examples) and letterheads."""

from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from docx.shared import Cm, Pt

from ezo_api.documents.docx.writer import render_docx as render_markdown_docx
from ezo_api.documents.markdown_io import parse_markdown
from ezo_api.organization.files import (
    OrgFileError,
    check_example_file,
    check_logo,
    check_template_file,
)
from ezo_api.organization.formatting import Formatting, formatting_of
from ezo_api.organization.letterhead import from_facts, mentions_organization, wants_letterhead
from ezo_api.organization.templates import check_template, open_template
from ezo_api.writing.catalog import BY_KEY
from ezo_api.writing.render import render_docx
from ezo_api.writing.service import Draft, OrgFacts, Section
from tests.fakes import png_bytes


def template_docx(*, marker: bool = True) -> bytes:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "TOSHKENT SHAHAR HOKIMLIGI — rasmiy blank"
    document.sections[0].left_margin = Cm(2.5)
    document.add_paragraph("Toshkent shahar hokimligi")
    document.add_paragraph("Manzil: Mustaqillik maydoni 5")
    if marker:
        document.add_paragraph("{{matn}}")
        document.add_paragraph("Namuna matn — bu yerga xat yoziladi")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_formatting_defaults_and_limits() -> None:
    assert formatting_of({}).font_name == "Times New Roman" and formatting_of(None).font_size == 14
    assert formatting_of({"font_size": 99}) == Formatting()  # invalid values fall back
    custom = formatting_of({"font_name": "Arial", "font_size": 12, "margins_cm": {"left": 2}})
    assert (custom.font_name, custom.font_size, custom.margins_cm.left) == ("Arial", 12, 2)


def test_template_marker_keeps_letterhead_and_drops_sample_text() -> None:
    document = open_template(template_docx())
    assert [p.text for p in document.paragraphs] == [
        "Toshkent shahar hokimligi",
        "Manzil: Mustaqillik maydoni 5",
    ]
    whole = open_template(template_docx(marker=False))
    assert len(whole.paragraphs) == 2  # no marker: the body stays, the text follows it
    assert check_template(template_docx())["has_marker"] is True
    assert "rasmiy blank" in check_template(template_docx())["header_text"]


def test_official_letter_in_the_approved_template() -> None:
    org = OrgFacts(
        name="Toshkent shahar hokimligi", template=template_docx(), template_title="Blank"
    )
    draft = Draft(
        title="Xizmat xati",
        fields={"recipient": "Vazirlikka", "subject": "Hisobot"},
        sections=[Section(text="Hisobot ilova qilinmoqda.")],
    )
    data = render_docx(BY_KEY["xizmat_xati"], draft, org=org, language="uz")
    document = Document(io.BytesIO(data))
    texts = [p.text for p in document.paragraphs]
    assert texts[:2] == ["Toshkent shahar hokimligi", "Manzil: Mustaqillik maydoni 5"]
    assert "Hisobot ilova qilinmoqda." in texts and "{{matn}}" not in " ".join(texts)
    assert not any("Namuna matn" in t for t in texts)
    header = [p.text for p in document.sections[0].header.paragraphs]
    assert header[0] == "TOSHKENT SHAHAR HOKIMLIGI — rasmiy blank" and "QORALAMA" in header
    assert abs(document.sections[0].left_margin - Cm(2.5)) < Cm(0.01)  # the template's margins


def test_formatting_and_logo_without_a_template() -> None:
    org = OrgFacts(
        name="Hokimlik",
        logo=check_logo("logo.png", png_bytes(300, 300)).data,
        formatting={"font_name": "Arial", "font_size": 12, "margins_cm": {"left": 2.0}},
    )
    draft = Draft(title="Xizmat xati", sections=[Section(text="Matn.")])
    document = Document(io.BytesIO(render_docx(BY_KEY["xizmat_xati"], draft, org=org)))
    normal = document.styles["Normal"]
    assert (normal.font.name, normal.font.size) == ("Arial", Pt(12))
    assert abs(document.sections[0].left_margin - Cm(2.0)) < Cm(0.01)
    assert len(document.inline_shapes) == 1  # the logo
    assert document.paragraphs[1].text == "Hokimlik"


def test_file_checks() -> None:
    logo = check_logo("photo.jpg", png_bytes(1200, 400))
    assert logo.mime_type == "image/png" and logo.filename == "logo.png"
    with pytest.raises(OrgFileError):
        check_logo("logo.png", b"not an image")
    template = check_template_file("Blank.docx", template_docx())
    assert template.mime_type.endswith("wordprocessingml.document")
    with pytest.raises(OrgFileError) as error:
        check_template_file("notes.txt", b"hello")
    assert error.value.code == "wrong_type"
    example = check_example_file("namuna.txt", b"Hurmatli rahbar! Ma'lum qilamizki ...")
    assert example.text_excerpt and example.text_excerpt.startswith("Hurmatli rahbar")
    with pytest.raises(OrgFileError):
        check_example_file("logo.png", png_bytes(10, 10))


def test_template_with_an_external_link_is_refused() -> None:
    source = template_docx()
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source)) as original, zipfile.ZipFile(buffer, "w") as copy:
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdX" Type="http://schemas.openxmlformats.org/'
                    b'officeDocument/2006/relationships/image" Target="http://x.example/a.png"'
                    b' TargetMode="External"/></Relationships>',
                )
            copy.writestr(item, data)
    with pytest.raises(OrgFileError) as error:
        check_template_file("Blank.docx", buffer.getvalue())
    assert error.value.code == "bad_template" and "image" in error.value.message


def test_letterhead_for_other_documents() -> None:
    org = OrgFacts(
        name="Hokimlik",
        department="Yoshlar bo'limi",
        phone="71 200 00 00",
        formatting={"font_size": 13},
    )
    letterhead = from_facts(org)
    assert (
        letterhead.lines == ["Hokimlik", "Yoshlar bo'limi"]
        and letterhead.contacts == "71 200 00 00"
    )
    data = render_markdown_docx(parse_markdown("# Reja\n\nMatn."), letterhead)
    document = Document(io.BytesIO(data))
    assert [p.text for p in document.paragraphs[:3]] == [
        "Hokimlik",
        "Yoshlar bo'limi",
        "71 200 00 00",
    ]
    assert document.styles["Normal"].font.size == Pt(13)
    assert wants_letterhead("Tashkilotimiz blankasida reja tayyorla")
    assert wants_letterhead("Подготовь на бланке")
    assert mentions_organization("Tashkilotimiz faoliyati haqida taqdimot")
    assert not mentions_organization("Quyosh tizimi haqida taqdimot")
