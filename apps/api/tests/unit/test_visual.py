"""Visual building blocks without a database: image safety, request rules, providers,
fallback, composition and images on slides."""

from __future__ import annotations

import base64
import io
from types import SimpleNamespace
from typing import Any

import httpx
import openai
import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ezo_api.documents.render.build import build_deck
from ezo_api.documents.render.spec import DeckSpec, Node, SlideSpec
from ezo_api.visual.base import ImageError, ImageRequest
from ezo_api.visual.compose import (
    DiagramSpec,
    LogoSpec,
    SocialSpec,
    compose_diagram,
    compose_logo,
    compose_social,
)
from ezo_api.visual.gemini_images import GeminiImageProvider
from ezo_api.visual.images import ImageRejectedError, cover_fit, normalize_image
from ezo_api.visual.kinds import detect_request, wants_edit, wants_no_images
from ezo_api.visual.openai_images import OpenAIImageProvider, size_for
from ezo_api.visual.service import AllImageProvidersFailedError, VisualService
from tests.fakes import FakeImages, png_bytes


def _size(data: bytes) -> tuple[int, int]:
    return Image.open(io.BytesIO(data)).size


# --- image safety ----------------------------------------------------------------------------


def test_normalize_crops_reencodes_and_strips_metadata() -> None:
    buffer = io.BytesIO()
    exif = Image.Exif()
    exif[0x010E] = "secret location note"  # ImageDescription
    Image.new("RGB", (1200, 600), "red").save(buffer, format="JPEG", exif=exif)
    result = normalize_image(buffer.getvalue(), aspect=1.0)
    assert (result.width, result.height) == (600, 600)
    image = Image.open(io.BytesIO(result.data))
    assert image.format == "PNG"
    assert not image.getexif() and b"secret location" not in result.data


def test_normalize_rejects_non_images_and_bombs() -> None:
    with pytest.raises(ImageRejectedError):
        normalize_image(b"<svg xmlns='http://www.w3.org/2000/svg'></svg>")
    with pytest.raises(ImageRejectedError):
        normalize_image(png_bytes(10, 10)[:40])  # truncated
    huge = io.BytesIO()
    Image.new("1", (8000, 6000)).save(huge, format="PNG")  # 48 MP, tiny file
    with pytest.raises(ImageRejectedError, match="dimensions"):
        normalize_image(huge.getvalue())


def test_large_images_are_downscaled_and_cover_fit_matches_the_box() -> None:
    result = normalize_image(png_bytes(4000, 1000))
    assert (result.width, result.height) == (2048, 512)
    fitted = cover_fit(png_bytes(1000, 1000), 400, 200)
    width, height = _size(fitted)
    assert abs(width / height - 2.0) < 0.01
    assert Image.open(io.BytesIO(fitted)).format == "JPEG"


# --- request rules ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("text", "kind", "aspect", "count"),
    [
        ("Kosmosdagi mushuk rasmini chiz", "image", "1:1", 1),
        ("Nur Academy uchun logo yarat, 3 ta variant", "logo", "1:1", 3),
        ("EZO arxitekturasi diagrammasini chiz", "diagram", "16:9", 1),
        ("Tadbir uchun Instagram story tayyorla", "social", "1:1", 1),
        ("Sayt uchun hero rasm, 16:9", "hero", "16:9", 1),
        ("Mahsulot rasmi kerak, vertikal", "product", "3:4", 1),
        ("Make 9 variants of a cat picture", "image", "1:1", 1),  # out of range: one
    ],
)
def test_detect_request(text: str, kind: str, aspect: str, count: int) -> None:
    request = detect_request(text, max_count=4)
    assert (request.kind, request.aspect_ratio, request.count) == (kind, aspect, count)


def test_social_formats_follow_the_platform() -> None:
    assert detect_request("Instagram story uchun banner", 4).social_format == "story"
    assert detect_request("LinkedIn post", 4).social_format == "landscape"
    assert detect_request("Instagram post", 4).social_format == "square"
    assert detect_request("Konsert afishasi", 4).social_format == "portrait"


def test_edit_and_no_image_rules() -> None:
    assert wants_edit("Fonini to'q ko'k qil")
    assert wants_edit("Bu rasmni yorqinroq qil")
    assert wants_edit("Make the background darker")
    assert not wants_edit("Kosmosdagi mushuk rasmini chiz")
    assert not wants_edit("Logo yarat, fonni ko'k qil")  # a new logo, described fully
    assert wants_no_images("Taqdimot qil, rasmsiz")
    assert not wants_no_images("Taqdimot qil")


# --- providers -------------------------------------------------------------------------------


class _Images:
    def __init__(self, response: Any = None, error: Exception | None = None) -> None:
        self.response, self.error, self.calls = response, error, []

    async def generate(self, **kwargs: Any) -> Any:
        self.calls.append(kwargs)
        if self.error:
            raise self.error
        return self.response

    async def edit(self, **kwargs: Any) -> Any:
        return await self.generate(**kwargs)


def _openai(images: _Images) -> OpenAIImageProvider:
    return OpenAIImageProvider("k", "gpt-image-test", 30, client=SimpleNamespace(images=images))


async def test_openai_adapter_decodes_and_maps_sizes() -> None:
    encoded = base64.b64encode(png_bytes(8, 8)).decode()
    images = _Images(SimpleNamespace(data=[SimpleNamespace(b64_json=encoded, revised_prompt="r")]))
    result = await _openai(images).generate(ImageRequest("a cat", "16:9", "low"))
    assert _size(result.data) == (8, 8) and result.revised_prompt == "r"
    assert images.calls[0] == {
        "model": "gpt-image-test",
        "prompt": "a cat",
        "size": "1536x1024",
        "quality": "low",
        "n": 1,
    }
    assert size_for("9:16") == "1024x1536" and size_for("4:3") == "1536x1024"
    assert size_for("1:1") == "1024x1024"
    edited = await _openai(images).edit(b"png", "image/png", ImageRequest("bluer"))
    assert edited.provider == "openai" and images.calls[1]["image"][0] == "image.png"


async def test_openai_moderation_is_a_refusal() -> None:
    response = httpx.Response(400, request=httpx.Request("POST", "https://api.test/v1/images"))
    error = openai.BadRequestError(
        "blocked", response=response, body={"code": "moderation_blocked", "message": "blocked"}
    )
    with pytest.raises(ImageError) as caught:
        await _openai(_Images(error=error)).generate(ImageRequest("x"))
    assert caught.value.kind == "refused"
    with pytest.raises(ImageError) as empty:
        await _openai(_Images(SimpleNamespace(data=[]))).generate(ImageRequest("x"))
    assert empty.value.kind == "failed"


class _GeminiModels:
    def __init__(self, response: Any) -> None:
        self.response, self.calls = response, []

    async def generate_content(self, **kwargs: Any) -> Any:
        self.calls.append(kwargs)
        return self.response


def _gemini(response: Any) -> tuple[GeminiImageProvider, _GeminiModels]:
    models = _GeminiModels(response)
    client = SimpleNamespace(aio=SimpleNamespace(models=models))
    return GeminiImageProvider("k", "gemini-image-test", 30, client=client), models


async def test_gemini_adapter_reads_inline_image_and_blocks() -> None:
    part = SimpleNamespace(
        inline_data=SimpleNamespace(data=png_bytes(4, 4), mime_type="image/png"), thought=False
    )
    candidate = SimpleNamespace(finish_reason="STOP", content=SimpleNamespace(parts=[part]))
    provider, models = _gemini(SimpleNamespace(prompt_feedback=None, candidates=[candidate]))
    result = await provider.generate(ImageRequest("a lake", "16:9"))
    assert _size(result.data) == (4, 4) and result.model == "gemini-image-test"
    config = models.calls[0]["config"]
    assert config.image_config.aspect_ratio == "16:9"
    edit = await provider.edit(png_bytes(4, 4), "image/png", ImageRequest("night"))
    assert edit.provider == "gemini" and len(models.calls[1]["contents"]) == 2

    blocked, _ = _gemini(
        SimpleNamespace(prompt_feedback=SimpleNamespace(block_reason="SAFETY"), candidates=[])
    )
    with pytest.raises(ImageError) as caught:
        await blocked.generate(ImageRequest("x"))
    assert caught.value.kind == "refused"
    no_image, _ = _gemini(
        SimpleNamespace(
            prompt_feedback=None,
            candidates=[SimpleNamespace(finish_reason="IMAGE_SAFETY", content=None)],
        )
    )
    with pytest.raises(ImageError) as safety:
        await no_image.generate(ImageRequest("x"))
    assert safety.value.kind == "refused"


# --- fallback --------------------------------------------------------------------------------


async def test_service_falls_back_but_never_after_a_refusal() -> None:
    down = FakeImages(name="down", fail_with="unavailable")
    broken = FakeImages(name="broken", data=b"not an image")
    good = FakeImages(name="good")
    service = VisualService([down, broken, good], [good], quality="low")
    outcome = await service.generate("a cat", "16:9")
    assert outcome.provider == "good"
    assert (outcome.image.width, outcome.image.height) == (2048, 1152)  # cropped to 16:9
    assert [a.kind for a in outcome.attempts] == ["unavailable", "failed", None]
    assert good.requests[0].quality == "low"

    refusing = FakeImages(name="strict", fail_with="refused")
    other = FakeImages(name="other")
    with pytest.raises(AllImageProvidersFailedError) as caught:
        await VisualService([refusing, other]).generate("x")
    assert caught.value.refused and other.requests == []
    with pytest.raises(AllImageProvidersFailedError, match="No image provider"):
        await VisualService().generate("x")
    assert not VisualService().can_generate


# --- composition and slides ------------------------------------------------------------------


def test_diagram_has_exact_labels_and_svg() -> None:
    spec = DiagramSpec(
        title="EZO oqimi",
        kind="flow",
        nodes=[Node(id="a", label="Foydalanuvchi"), Node(id="b", label="EzoCore")],
        edges=[("a", "b")],
    )
    composed = compose_diagram(spec, "uz")
    assert (composed.width, composed.height) == (1920, 1080) and composed.text_fits
    assert composed.svg is not None
    assert b"Foydalanuvchi" in composed.svg and b"EzoCore" in composed.svg


def test_social_and_logo_layouts() -> None:
    story = compose_social(
        SocialSpec(headline="Yangi kurs", subline="Ro'yxatdan o'ting", format="story"),
        png_bytes(512, 900),
    )
    assert (story.width, story.height) == (1080, 1920)
    placeholder = compose_social(SocialSpec(headline="Salom", format="landscape"), None)
    assert (placeholder.width, placeholder.height) == (1600, 900)
    logo = compose_logo(LogoSpec(brand="Nur Academy", tagline="Learn smarter"), None)
    assert (logo.width, logo.height) == (1024, 1024) and logo.text_fits


def _deck() -> DeckSpec:
    return DeckSpec(
        title="Deck",
        slides=[
            SlideSpec(layout="cover", title="Cover"),
            SlideSpec(
                layout="hero_image", title="Hero", body="One sentence.", image_prompt="a lake"
            ),
            SlideSpec(layout="closing", title="Thanks"),
        ],
    )


def test_slide_images_reach_every_format() -> None:
    build = build_deck(_deck(), images={2: png_bytes(600, 800, "#ff0000")})
    presentation = Presentation(io.BytesIO(build.pptx))
    pictures = [s for s in presentation.slides[1].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 1
    assert b"data:image/jpeg;base64," in build.html
    assert b"/Subtype /Image" in build.pdf
    preview = Image.open(io.BytesIO(build.previews[1])).convert("RGB")
    red, green, blue = preview.getpixel((960, 300))  # inside the image panel
    assert red > 200 and green < 40 and blue < 40
    assert not build.qa.issues


def test_hero_slide_without_image_is_flagged_as_placeholder() -> None:
    build = build_deck(_deck(), previews=False)
    assert [i.code for i in build.qa.issues] == ["image_placeholder"]
    # An image for a slide that is not a hero slide is ignored.
    other = build_deck(_deck(), previews=False, images={1: png_bytes(10, 10)})
    assert len(Presentation(io.BytesIO(other.pptx)).slides[0].shapes) > 0
    assert [i.code for i in other.qa.issues] == ["image_placeholder"]


def test_spec_accepts_only_uuid_image_ids() -> None:
    slide = SlideSpec(layout="hero_image", image_artifact_id="../../etc/passwd")
    assert slide.image_artifact_id == ""
    good = "123e4567-e89b-12d3-a456-426614174000"
    assert SlideSpec(image_artifact_id=good).image_artifact_id == good


def test_providers_come_from_settings_in_the_configured_order(tmp_path: Any) -> None:
    from ezo_api.visual.service import build_visual_service
    from tests.conftest import make_settings

    url = f"sqlite+aiosqlite:///{tmp_path}/x.db"
    keys = {"OPENAI_API_KEY": "sk-test", "GEMINI_API_KEY": "g-test"}
    auto = build_visual_service(make_settings(url, **keys, EZO_OPENAI_IMAGE_MODEL="img-x"))
    assert [(p.name, p.model) for p in auto.generators] == [
        ("openai", "img-x"),
        ("gemini", "gemini-2.5-flash-image"),
    ]
    ordered = build_visual_service(make_settings(url, **keys, EZO_IMAGE_PROVIDERS="gemini"))
    assert [p.name for p in ordered.generators] == ["gemini"]
    assert not build_visual_service(
        make_settings(url, **keys, EZO_IMAGE_PROVIDERS="none")
    ).can_generate
    assert not build_visual_service(make_settings(url)).can_generate  # no keys, no images
    limited = build_visual_service(make_settings(url, EZO_IMAGE_MAX_PER_DECK="0"))
    assert limited.max_per_deck == 0
