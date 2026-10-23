"""Executor protocol v1 additions: OCR options and Office → PDF conversion.

Container isolation itself is tested against real Docker in
tests/integration/test_isolated_executor_container.py (opt-in, CI).
"""

from __future__ import annotations

import base64
import hashlib
import io
import json
import os
import shutil
import subprocess
import sys
from collections.abc import Iterator

import pytest
from pydantic import ValidationError

from ezo_api.documents import isolation
from ezo_api.documents.executor import child
from ezo_api.documents.executor.controller import IsolatedExecutor
from ezo_api.documents.executor.protocol import (
    ConversionRequest,
    ConversionResult,
    ExecutorError,
    ExecutorRequest,
    ExtractionRequest,
    ExtractionResult,
    OcrSummary,
)
from ezo_api.documents.isolation import (
    IsolationUnavailable,
    native_execution_allowed,
    require_isolation,
    trusted_native_regression_fixtures,
    verified_executor_workload,
)
from tests.office_fixtures import official_letter_docx, report_xlsx, with_external_image

SHA = "a" * 64
needs_office = pytest.mark.skipif(
    shutil.which("soffice") is None and not os.environ.get("EZO_REQUIRE_OFFICE"),
    reason="LibreOffice is not installed",
)


@pytest.fixture
def production_isolation() -> Iterator[None]:
    """Leave the session-wide regression-fixture exception, as the API runs."""
    token = isolation._FIXTURES.set(False)
    try:
        yield
    finally:
        isolation._FIXTURES.reset(token)


def test_request_union_is_data_only_and_closed() -> None:
    extract = ExecutorRequest.validate_json(
        json.dumps({"operation": "extract", "format": "pdf", "data_base64": ""})
    )
    assert isinstance(extract, ExtractionRequest) and extract.ocr_languages == ()
    convert = ExecutorRequest.validate_json(
        json.dumps({"operation": "convert_pdf", "format": "pptx", "data_base64": ""})
    )
    assert isinstance(convert, ConversionRequest)
    for bad in (
        {"format": "pdf", "data_base64": ""},  # the operation is always explicit
        {"operation": "run", "format": "pdf", "data_base64": ""},
        {"operation": "convert_pdf", "format": "odt", "data_base64": ""},  # ODF links unchecked
        {"operation": "extract", "format": "pdf", "data_base64": "", "ocr_languages": ["deu"]},
        {"operation": "extract", "format": "pdf", "data_base64": "", "command": "sh"},
        {"operation": "convert_pdf", "format": "docx", "data_base64": "", "binary": "/bin/sh"},
    ):
        with pytest.raises(ValidationError):
            ExecutorRequest.validate_json(json.dumps(bad))


def test_completed_ocr_cannot_still_need_ocr() -> None:
    with pytest.raises(ValidationError):
        ExtractionResult(
            source_sha256=SHA,
            text="x",
            sections=[],
            needs_ocr=True,
            ocr=OcrSummary(status="completed"),
        )
    result = ExtractionResult(
        source_sha256=SHA, text="", sections=[], needs_ocr=True, ocr=OcrSummary(status="deadline")
    )
    assert result.ocr is not None and result.ocr.status == "deadline"


def test_native_binaries_stay_denied_outside_a_verified_workload(production_isolation) -> None:
    assert not native_execution_allowed()
    with pytest.raises(IsolationUnavailable):
        require_isolation()
    # This test process is not UID 10001 in a capability-free, cgroup-limited,
    # network-less container: the workload context must refuse to open.
    with (
        pytest.raises((RuntimeError, OSError, KeyError, ValueError)),
        verified_executor_workload(),
    ):
        pytest.fail("entered the executor context outside the container boundary")
    assert not native_execution_allowed()


def test_child_fails_closed_outside_the_container_with_fixed_error() -> None:
    request = ExtractionRequest(format="pdf", data_base64=base64.b64encode(b"%PDF-").decode())
    completed = subprocess.run(
        [sys.executable, "-m", "ezo_api.documents.executor.child"],
        input=request.model_dump_json().encode(),
        capture_output=True,
        timeout=60,
        check=False,
    )
    assert completed.returncode == 1 and completed.stdout == b""
    assert json.loads(completed.stderr) == {"code": "invalid_document_or_resource_limit"}


class Recording(IsolatedExecutor):
    def __init__(self, output: bytes) -> None:
        super().__init__("sha256:" + "0" * 64)
        self.output = output
        self.requests: list[ExtractionRequest | ConversionRequest] = []

    def _run(self, request, cancelled):  # type: ignore[no-untyped-def]
        self.requests.append(request)
        return self.output


def test_controller_sends_ocr_languages_and_validates_result() -> None:
    data = b"%PDF-scan"
    result = ExtractionResult(
        source_sha256=hashlib.sha256(data).hexdigest(), text="Salom", sections=[]
    )
    executor = Recording(result.model_dump_json().encode())
    assert executor.extract(data, "pdf", ocr_languages=("uzb", "rus")).text == "Salom"
    sent = executor.requests[0]
    assert isinstance(sent, ExtractionRequest) and sent.ocr_languages == ("uzb", "rus")
    with pytest.raises(ExecutorError, match="source_mismatch"):
        Recording(result.model_dump_json().encode()).extract(b"%PDF-other", "pdf")


def test_controller_verifies_converted_pdf_identity() -> None:
    data, pdf = b"PK-docx", b"%PDF-1.7 converted"
    good = ConversionResult(
        source_sha256=hashlib.sha256(data).hexdigest(),
        pdf_base64=base64.b64encode(pdf).decode(),
        pdf_sha256=hashlib.sha256(pdf).hexdigest(),
        pages=1,
        engine="LibreOffice",
    )
    out, meta = Recording(good.model_dump_json().encode()).convert_to_pdf(data, "docx")
    assert out == pdf and meta.pages == 1
    tampered = good.model_copy(update={"pdf_sha256": "b" * 64})
    with pytest.raises(ExecutorError, match="invalid_result"):
        Recording(tampered.model_dump_json().encode()).convert_to_pdf(data, "docx")
    not_pdf = good.model_copy(
        update={
            "pdf_base64": base64.b64encode(b"MZ").decode(),
            "pdf_sha256": hashlib.sha256(b"MZ").hexdigest(),
        }
    )
    with pytest.raises(ExecutorError, match="invalid_result"):
        Recording(not_pdf.model_dump_json().encode()).convert_to_pdf(data, "docx")
    with pytest.raises(ExecutorError, match="source_mismatch"):
        Recording(good.model_dump_json().encode()).convert_to_pdf(b"PK-other", "docx")


def test_ocr_reports_unavailable_without_isolation(production_isolation) -> None:
    sections, summary = child._ocr(b"%PDF-", "pdf", ("eng",))
    assert sections == [] and summary.status == "unavailable"


def test_images_require_ocr_languages() -> None:
    request = ExtractionRequest(format="png", data_base64="")
    with pytest.raises(ValueError, match="ocr_required"):
        child.extract(request, b"\x89PNG\r\n\x1a\n")


@needs_office
@pytest.mark.parametrize("kind", ["docx", "xlsx", "pptx"])
def test_conversion_produces_a_real_pdf(kind: str) -> None:
    from pptx import Presentation

    from ezo_api.documents.pdf import engine

    if kind == "docx":
        data = official_letter_docx()
    elif kind == "xlsx":
        data = report_xlsx()
    else:
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Fotosintez"
        output = io.BytesIO()
        deck.save(output)
        data = output.getvalue()
    request = ConversionRequest(operation="convert_pdf", format=kind, data_base64="")  # type: ignore[arg-type]
    with trusted_native_regression_fixtures():
        result = child.convert(request, data)
    pdf = base64.b64decode(result.pdf_base64)
    assert pdf.startswith(b"%PDF-") and result.pages >= 1
    assert result.pdf_sha256 == hashlib.sha256(pdf).hexdigest()
    assert result.source_sha256 == hashlib.sha256(data).hexdigest()
    assert engine.inspect_pdf(pdf).page_count == result.pages
    if kind == "pptx":
        from pypdf import PdfReader

        assert "Fotosintez" in PdfReader(io.BytesIO(pdf)).pages[0].extract_text()


@needs_office
def test_conversion_refuses_external_links() -> None:
    from ezo_api.documents.office_convert import OfficeConvertError

    request = ConversionRequest(operation="convert_pdf", format="docx", data_base64="")
    with trusted_native_regression_fixtures(), pytest.raises(OfficeConvertError):
        child.convert(request, with_external_image(official_letter_docx()))


def test_conversion_is_denied_without_isolation(production_isolation) -> None:
    request = ConversionRequest(operation="convert_pdf", format="docx", data_base64="")
    with pytest.raises(RuntimeError, match="converter_unavailable"):
        child.convert(request, official_letter_docx())
