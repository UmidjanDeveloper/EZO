"""Professional writing: document types, options, style checks, validation and layout."""

from __future__ import annotations

import io

import pytest
from docx import Document
from pypdf import PdfReader

from ezo_api.core.text import normalize
from ezo_api.writing.catalog import BY_KEY, CATALOG, detect, wants_template, wants_writing
from ezo_api.writing.render import render_docx
from ezo_api.writing.service import Draft, OrgFacts, Section, blank, to_markdown, validate
from ezo_api.writing.style import lint, options_from
from tests.office import office_to_pdf, requires_office

ORG = OrgFacts(
    name="Xatirchi tumani hokimligi",
    department="Yoshlar ishlari bo'limi",
    address="Yangirabod sh., Mustaqillik ko'chasi 1",
    phone="+998 79 000 00 00",
)


@pytest.mark.parametrize(
    ("request_text", "key"),
    [
        ("Hokimga ariza yozib ber", "ariza"),
        ("Shu faktlar asosida tushuntirish xati yoz", "tushuntirish_xati"),
        ("Hokimlik uchun xizmat xati tayyorla", "xizmat_xati"),
        ("Напиши служебное письмо в хокимият", "xizmat_xati"),
        ("Javob xati tayyorla", "javob_xati"),
        ("Yig'ilish bayonnomasi shablonini ber", "bayonnoma"),
        ("Dalolatnoma tuz", "dalolatnoma"),
        ("Buyruq loyihasini tayyorla", "buyruq"),
        ("Qaror loyihasi yoz", "qaror"),
        ("Oylik hisobot yoz", "hisobot"),
        ("Ma'ruza tayyorla", "maruza"),
        ("Press-reliz yozib ber", "press_reliz"),
        ("Taklifnoma yoz", "taklifnoma"),
        ("Tabriknoma yoz", "tabriknoma"),
        ("Ishonchnoma loyihasi", "ishonchnoma"),
        ("CV tayyorla", "cv"),
        ("Write a cover letter", "cover_letter"),
        ("Motivatsion xat yozib ber", "motivatsion_xat"),
        ("Tadqiqot taklifi tayyorla", "tadqiqot_taklifi"),
        ("Ўзбекистон ҳокимлигига ариза ёз", "ariza"),
    ],
)
def test_document_types_are_recognised(request_text: str, key: str) -> None:
    found = detect(normalize(request_text))
    assert found is not None and found.key == key


def test_catalog_is_complete_and_requests_need_a_writing_verb() -> None:
    assert len(CATALOG) >= 30
    assert all({"uz", "ru", "en"} <= set(d.names) and d.sections for d in CATALOG)
    assert wants_writing(normalize("Ariza haqida nima bilasan?")) is None
    assert wants_writing(normalize("Telegram post yoz")) is None
    assert wants_template(normalize("Bayonnoma shablonini ber"))
    assert wants_template(normalize("Xizmat xati blankini ber"))
    # "on our letterhead" asks for content, not an empty template.
    assert not wants_template(normalize("Tashkilotimiz blankasida xizmat xati tayyorla"))
    assert not wants_template(normalize("Подготовь служебное письмо на бланке"))


def test_options_from_words() -> None:
    short = options_from("qisqa va rasmiy uslubda, kirillda yoz")
    assert short.words == (80, 250) and short.mode == "official" and short.language == "uz-cyrl"
    assert options_from("300 so'zli ilmiy ma'ruza").words == (255, 345)
    assert options_from("2 betlik hisobot").words == (600, 900)
    assert options_from("ruscha yoz").language == "ru"
    assert options_from("qoralama belgisisiz").draft_label is False


def test_style_lint_finds_machine_writing() -> None:
    text = (
        "Hozirgi kunda bu masala dolzarb. Ulug'vor islohotlar amalga oshirilmoqda. "
        "Hujjat imzolandi. " + "biz bu ishni bajardik. " * 4
    )
    kinds = {issue.kind for issue in lint(text).issues}
    assert {"cliche", "inflated", "approval_claim", "repetition"} <= kinds
    assert "mixed_script" in {i.kind for i in lint("o'zbek " * 30 + "кирилл " * 30).issues}
    assert lint("Tuman hokimligiga ma'lumot uchun yuboriladi.").issues == []


def test_validation_never_keeps_invented_details() -> None:
    doc = BY_KEY["tushuntirish_xati"]
    draft = Draft(
        title="Tushuntirish xati",
        fields={
            "recipient": "Tuman hokimi A. Karimovga",
            "sender": "Bosh mutaxassis B. Aliyev",
            "date": "2026-yil 15-mart",  # nobody said this
            "event": None,
        },
        sections=[
            Section(text="12-martda ishga 30 daqiqa kechikdim. 2025-yilda ham shunday bo'lgan.")
        ],
    )
    report = validate(
        doc,
        draft,
        evidence="12-martda ishga 30 daqiqa kechikdim. Bosh mutaxassis B. Aliyev, hokim A. Karimov",
        options=options_from(""),
        language="uz",
    )
    assert draft.fields["date"] is None and report.replaced_fields == ["date"]
    assert "2025" in report.unverified_numbers and "12" not in report.unverified_numbers
    assert "[Holat (nima, qachon)]" in report.placeholders
    assert report.filled["recipient"] == "Tuman hokimi A. Karimovga"


@pytest.mark.parametrize("doc", CATALOG, ids=lambda d: d.key)
def test_every_type_renders_an_honest_template(doc) -> None:  # type: ignore[no-untyped-def]
    draft = blank(doc)
    data = render_docx(doc, draft, org=ORG)
    document = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs) + "\n".join(
        c.text for t in document.tables for r in t.rows for c in r.cells
    )
    assert doc.names["uz"].split(" (")[0].lower() in text.lower()
    assert "[" in text  # placeholders, never invented details
    assert "imzolandi" not in text.lower()
    header = document.sections[0].header.paragraphs[0].text
    assert (header == "QORALAMA") == doc.draft
    if doc.layout in ("letter", "act", "minutes"):
        assert "Xatirchi tumani hokimligi" in text
        assert "[imzo]" in text
    assert to_markdown(doc, draft, ORG).startswith("# ")


def test_letter_layout() -> None:
    doc = BY_KEY["xizmat_xati"]
    draft = Draft(
        title="Xizmat xati",
        fields={"recipient": "Navoiy viloyati hokimligiga", "subject": "Yoshlar tadbiri"},
        sections=[
            Section(
                text="Tumanda yoshlar bilan uchrashuv o'tkazildi.\n\n- Ishtirokchilar: 120 nafar"
            )
        ],
    )
    document = Document(io.BytesIO(render_docx(doc, draft, org=ORG, draft_label=False)))
    paragraphs = document.paragraphs
    recipient = next(p for p in paragraphs if p.text == "Navoiy viloyati hokimligiga")
    assert recipient.paragraph_format.left_indent.cm > 8
    assert any(p.text == "Mavzu: Yoshlar tadbiri" for p in paragraphs)
    bullet = next(p for p in paragraphs if p.text == "Ishtirokchilar: 120 nafar")
    assert bullet.style.name == "List Bullet"
    assert document.sections[0].header.paragraphs[0].text == ""  # label switched off
    assert document.sections[0].left_margin.cm == pytest.approx(3, abs=0.01)
    assert document.styles["Normal"].font.name == "Times New Roman"


@requires_office
def test_libreoffice_opens_a_rendered_letter() -> None:
    doc = BY_KEY["bayonnoma"]
    pdf = office_to_pdf(render_docx(doc, blank(doc), org=ORG), ".docx")
    text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "BAYONNOMA" in text and "QORALAMA" in text
