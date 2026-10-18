"""Translation building blocks: Uzbek Latin ↔ Cyrillic, reading languages from requests,
glossary matching, quality checks and the Word outputs. No model involved."""

from __future__ import annotations

import io
from collections import Counter
from itertools import pairwise

import pytest
from docx import Document

from ezo_api.translation.documents import bilingual_docx, docx_pairs, transliterate_docx
from ezo_api.translation.glossary import Term, add_request, missing_in, prompt_block, relevant
from ezo_api.translation.languages import (
    default_target,
    detect,
    domain_of,
    parse_languages,
)
from ezo_api.translation.quality import check, numbers, surnames
from ezo_api.translation.requests import inline_text, wants_translation
from ezo_api.translation.service import pairs_of, parse_reply, script_change
from ezo_api.translation.translit import detect_script, spans, to_cyrillic, to_latin

LATIN_CYRILLIC = [
    ("O'zbekiston Respublikasi Prezidenti", "Ўзбекистон Республикаси Президенти"),
    ("Toshkent shahar hokimligi", "Тошкент шаҳар ҳокимлиги"),
    ("yo'l, tog', ma'no, san'at, mas'uliyat", "йўл, тоғ, маъно, санъат, масъулият"),
    ("yetti, yer, poyezd, ekologiya, aeroport, shoir", "етти, ер, поезд, экология, аэропорт, шоир"),
    ("Konstitutsiya, aksiya, stansiya, versiya", "Конституция, акция, станция, версия"),
    ("komissiya, funksiya", "комиссия, функция"),
    ("yangi yil, yulduz, yoshlar, ishlayapti", "янги йил, юлдуз, ёшлар, ишлаяпти"),
    ("SHAHAR XO'JALIGI", "ШАҲАР ХЎЖАЛИГИ"),
    ("mashhur, Is'hoq, litsey, sotsial, obyekt", "машҳур, Исҳоқ, лицей, социал, объект"),
    ("G'afur G'ulom, Erkin Vohidov", "Ғафур Ғулом, Эркин Воҳидов"),
]  # fmt: skip


@pytest.mark.parametrize(("latin", "cyrillic"), LATIN_CYRILLIC)
def test_latin_and_cyrillic_both_ways(latin: str, cyrillic: str) -> None:
    assert to_cyrillic(latin) == cyrillic
    assert to_latin(cyrillic) == latin


def test_transliteration_details() -> None:
    # Apostrophe variants are all understood; official modifier letters on request.
    assert to_cyrillic("Oʻzbekiston, togʻ, maʼno") == "Ўзбекистон, тоғ, маъно"
    assert to_latin("Ўзбекистон маъно", apostrophe="modifier") == "Oʻzbekiston maʼno"
    # Months: the soft sign only on the bare word; Cyrillic → Latin uses the Uzbek spelling.
    assert to_cyrillic("12-sentabr, 1-yanvar, aprelda") == "12-сентябрь, 1-январь, апрелда"
    assert to_latin("сентябрь, октябрь, мўъжиза, батальон") == "sentabr, oktabr, mo'jiza, batalyon"
    # ц: s at the start or after a consonant, ts after a vowel.
    assert to_latin("цирк, лицей, конституция") == "sirk, litsey, konstitutsiya"
    # Links, e-mail addresses and digits are never changed.
    assert (
        to_cyrillic("www.gov.uz va info@ezo.uz, 2024-yil") == "www.gov.uz ва info@ezo.uz, 2024-йил"
    )
    # Text already in the target script passes through.
    assert to_cyrillic("Ўзбекистон") == "Ўзбекистон" and to_latin("Salom") == "Salom"
    assert detect_script("Salom dunyo") == "latin" and detect_script("Салом") == "cyrillic"
    assert detect_script("2024") is None
    # Spans cover the input exactly once (needed for Word runs).
    text = "Shahar xo'jaligi, 12-mart."
    pieces = spans(text, "cyrillic")
    assert [(s.start, s.end) for s in pieces] == sorted((s.start, s.end) for s in pieces)
    assert pieces[0].start == 0 and pieces[-1].end == len(text)
    assert all(a.end == b.start for a, b in pairwise(pieces))


def test_languages_from_requests() -> None:
    cases = {
        "Shu ruscha hujjatni o'zbek lotiniga tarjima qil, formatini saqla": (None, "uz"),
        "Inglizchaga tarjima qil": (None, "en"),
        "Переведи с русского на узбекский": ("ru", "uz"),
        "Bu matnni kirillga o'gir": (None, "uz-cyrl"),
        "o'zbek kirilliga tarjima qil": (None, "uz-cyrl"),
        "Translate this into Russian": (None, "ru"),
        "Inglizchadan ruschaga": ("en", "ru"),
        "Lotindan kirillga o'tkaz": (None, "uz-cyrl"),
        "ruschadagi xatni tarjima qil": ("ru", None),
        "tarjima qil": (None, None),
    }
    for request, (source, target) in cases.items():
        parsed = parse_languages(request)
        assert (parsed.source, parsed.target) == (source, target), request
    assert detect("Ўзбекистон Республикаси") == "uz-cyrl"
    assert detect("Привет, как дела?") == "ru" and detect("Hello, how are you?") == "en"
    assert domain_of("rasmiy xatni tarjima qil") == "official"
    assert domain_of("ilmiy maqolani tarjima qil") == "academic"
    assert domain_of("Переведи юридический текст") == "legal"
    assert domain_of("tarjima qil") == "general"
    assert default_target("uz", "uz") == "en" and default_target("en", "uz") == "uz"
    assert default_target("ru", "ru") == "uz"
    assert script_change("uz", "uz-cyrl") and not script_change("uz", "en")


def test_requests_and_inline_text() -> None:
    assert inline_text("Inglizchaga tarjima qil: Salom dunyo") == "Salom dunyo"
    assert inline_text("Tarjima qil\n\nBirinchi xatboshi.\n\nIkkinchi.") == (
        "Birinchi xatboshi.\n\nIkkinchi."
    )
    assert inline_text('Translate "good morning" into Uzbek') == "good morning"
    assert inline_text("Kirillga o'gir: Toshkent") == "Toshkent"
    assert inline_text("Shu faylni inglizchaga tarjima qil") is None
    assert wants_translation("Переведи на английский: доброе утро")
    assert wants_translation("Shu faylni kirill yozuviga o'gir")
    assert wants_translation("Glossariyga qo'sh: «hokimlik» → «khokimiyat»")
    assert not wants_translation("Ingliz tilini o'rganmoqchiman")
    assert not wants_translation("Salom, qalaysan?")


def test_glossary_matching_and_commands() -> None:
    terms = [
        Term("hokimlik", "khokimiyat", organization=True),
        Term("tuman", "district"),
        Term("Hokimlik", "administration"),  # a general entry: the organization's one wins
        Term("vazirlik", "ministry"),
    ]
    text = "Тошкент туманидаги ҳокимликка xat yuborildi."
    found = relevant(terms, text, "uz")
    assert [t.target for t in found] == ["khokimiyat", "district"]
    missing = missing_in(
        terms, text, "A letter was sent to the district administration.", "uz", "en"
    )
    assert [t.target for t in missing] == ["khokimiyat"]
    assert relevant([Term("администрация", "ma'muriyat")], "в администрации района", "ru")
    block = prompt_block([Term("hokimlik", "ҳокимият")], "uz")
    assert "“hokimlik” → “hokimiyat”" in block  # Uzbek terms reach the model in Latin
    assert add_request("Glossariyga qo'sh: «hokimlik» → «khokimiyat»") == ("hokimlik", "khokimiyat")
    assert add_request("Add to the glossary: district = tuman") == ("district", "tuman")
    assert add_request("Запомни термин: «район» — «tuman»") == ("район", "tuman")
    assert add_request("Hokimlik so'zini tarjima qil") is None


def test_numbers_months_and_names() -> None:
    assert numbers("1 250 000 so'm, 3,5 %") == numbers("1,250,000 soums, 3.5 %")
    assert numbers("2024-yil 12-mart") == numbers("March 12, 2024") == numbers("12 марта 2024 года")
    assert numbers("05.03.2024") == numbers("5 March 2024")
    assert numbers("Tadbirda 3 marta chiqdi") == Counter({3: 1})  # "marta" = times, not March
    assert numbers("may be 5 items") == numbers("5")  # "may" counts only next to a number…
    assert numbers("5-may") == numbers("May 5") == numbers("5 мая")  # …as here
    assert surnames("Akmal Karimovga va Dilnoza Yusupovaga") == ["Karimov", "Yusupova"]
    assert surnames("Каримов Ш.") == ["Karimov"]


def test_quality_report() -> None:
    source = (
        "Toshkent shahar hokimligi 2024-yil 12-mart kuni 1 250 000 so'm ajratdi. "
        "Mas'ul: Akmal Karimov. Batafsil: www.gov.uz"
    )
    good = (
        "On March 12, 2024 the Tashkent City Khokimiyat allocated 1,250,000 soums. "
        "Responsible: Akmal Karimov. Details: www.gov.uz"
    )
    report = check([(source, good)], source_language="uz", target_language="en")
    assert report.ok and report.lines() == []
    bad = "On March 21 the Tashkent khokimiyat allocated 1,250,000 soums. Responsible: A. Karimova."
    report = check(
        [(source, bad)],
        source_language="uz",
        target_language="en",
        terms=[Term("hokimlik", "khokimiyat"), Term("mas'ul", "in charge")],
        notes=["“mas'ul” could also mean “accountable”"],
    )
    assert report.missing_numbers == ["12", "2024"] and report.extra_numbers == ["21"]
    assert report.links_changed == ["www.gov.uz"]
    assert report.names_to_check == []  # "Karimova" keeps the surname stem
    assert report.terms_missing == ["“mas'ul” → “in charge”"] and report.terms_used == 1
    assert report.ambiguities == ["“mas'ul” could also mean “accountable”"]
    assert not report.ok and len(report.lines()) == 5
    renamed = check(
        [("Akmal Karimov keldi.", "Akmal Rahimov came.")],
        source_language="uz",
        target_language="en",
    )
    assert renamed.names_to_check == ["Karimov"]
    # Russian / Cyrillic romanisation still finds the name.
    cyr = check(
        [("Пришёл Шавкат Каримов.", "Shavkat Karimov keldi.")],
        source_language="ru",
        target_language="uz",
    )
    assert cyr.names_to_check == []
    romanised = check(
        [("Bahodir Xo'jayev keldi.", "Bahodir Khodjaev arrived.")],
        source_language="uz",
        target_language="en",
    )
    assert romanised.names_to_check == []
    paragraphs = check(
        [
            ("Birinchi xatboshi juda muhim ma'lumotlarni o'z ichiga oladi.", ""),
            ("Ikkinchi xatboshi ham muhim.", "Ikkinchi xatboshi ham muhim."),
            ("Uchinchi " * 20, "Third."),
        ],
        source_language="uz",
        target_language="en",
    )
    assert paragraphs.empty_paragraphs == [0] and paragraphs.untranslated_paragraphs == [1]
    assert paragraphs.length_warnings == [2]
    script = check([("Salom", "Salom dunyo")], source_language="ru", target_language="uz-cyrl")
    assert script.script_note == "The translation is in Latin script; Cyrillic was asked"


def test_reply_parsing_and_pairs() -> None:
    text, notes = parse_reply("Hello.\n\nSecond.\n=====\nNOTES:\n- “bank” could be a river bank")
    assert text == "Hello.\n\nSecond." and notes == ["“bank” could be a river bank"]
    assert parse_reply("Hello.\n=====\nNOTES: none") == ("Hello.", [])
    assert parse_reply("Just the text") == ("Just the text", [])
    assert pairs_of("A.\n\nB.", "X.\n\nY.") == [("A.", "X."), ("B.", "Y.")]
    assert pairs_of("A.\n\nB.", "X.") == [("A.\n\nB.", "X.")]


def _word() -> bytes:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("O'zbekiston ").bold = True
    paragraph.add_run("Respublikasi qarori")
    split = document.add_paragraph("S")  # "Sh" split across two runs
    split.add_run("hahar hokimligi www.gov.uz").italic = True
    document.add_table(rows=1, cols=1).rows[0].cells[0].text = "Yig'ilish bayonnomasi"
    document.sections[0].header.paragraphs[0].text = "Toshkent shahri"
    document.core_properties.title = "Qaror"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_word_script_conversion_keeps_runs_tables_and_headers() -> None:
    original = _word()
    data, changed = transliterate_docx(original, "cyrillic")
    document = Document(io.BytesIO(data))
    assert changed == 4
    assert [(r.text, r.bold) for r in document.paragraphs[0].runs] == [
        ("Ўзбекистон ", True),
        ("Республикаси қарори", None),
    ]
    assert [(r.text, r.italic) for r in document.paragraphs[1].runs] == [
        ("Ш", None),
        ("аҳар ҳокимлиги www.gov.uz", True),
    ]
    assert document.tables[0].cell(0, 0).text == "Йиғилиш баённомаси"
    assert document.sections[0].header.paragraphs[0].text == "Тошкент шаҳри"
    assert document.core_properties.title == "Қарор"
    back, _ = transliterate_docx(data, "latin")
    assert [p.text for p in Document(io.BytesIO(back)).paragraphs] == [
        "O'zbekiston Respublikasi qarori",
        "Shahar hokimligi www.gov.uz",
    ]
    assert docx_pairs(original, data)[0] == (
        "O'zbekiston Respublikasi qarori",
        "Ўзбекистон Республикаси қарори",
    )


def test_bilingual_word_file() -> None:
    data = bilingual_docx(
        [("Salom.", "Hello."), ("Rahmat.", "Thank you.")],
        title="Tarjima",
        source_label="O'zbekcha",
        target_label="English",
    )
    document = Document(io.BytesIO(data))
    table = document.tables[0]
    assert [[c.text for c in row.cells] for row in table.rows] == [
        ["O'zbekcha", "English"],
        ["Salom.", "Hello."],
        ["Rahmat.", "Thank you."],
    ]
    assert document.sections[0].page_width > document.sections[0].page_height  # landscape
