"""Conversions through chat with no AI provider configured: scan → Word / searchable PDF,
tables → Excel, spreadsheet operations."""

from __future__ import annotations

import io
from collections.abc import AsyncIterator

import pytest
from docx import Document
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from pypdf import PdfReader

from ezo_api.documents.pdf import engine
from ezo_api.main import create_app
from tests.conftest import make_settings
from tests.files import make_xlsx
from tests.integration.helpers import stream_chat
from tests.integration.test_document_agent import created, upload
from tests.ocr_fixtures import ENGLISH, UZ_LATIN, needs_tesseract, page_image, png, recall
from tests.pdf_fixtures import make_pdf
from tests.table_fixtures import table_pdf


@pytest.fixture
async def offline(database_url: str) -> AsyncIterator[AsyncClient]:
    app: FastAPI = create_app(make_settings(database_url, EZO_OCR_PROVIDER="auto"), providers={})
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as client:
        yield client
    await app.state.services.db.dispose()


async def meta_of(client: AsyncClient, artifact: dict[str, object]) -> dict[str, object]:
    detail = (await client.get(f"/api/artifacts/{artifact['id']}")).json()
    return detail["current_version"]  # type: ignore[no-any-return]


async def download(client: AsyncClient, artifact: dict[str, object]) -> bytes:
    return (await client.get(f"/api/artifacts/{artifact['id']}/versions/1/download")).content


@needs_tesseract
async def test_scan_to_editable_and_layout_word(offline: AsyncClient) -> None:
    scan = engine.images_to_pdf([png(page_image(UZ_LATIN)), png(page_image([]))]).data
    document = await upload(offline, "skan.pdf", scan)
    events = await stream_chat(
        offline, message="Shu skanni Word ga aylantir", attachments=[document["id"]]
    )
    [editable] = created(events)
    version = await meta_of(offline, editable)
    assert version["filename"] == "skan (editable).docx"
    meta = version["generation_metadata"]
    assert meta["operation"] == "ocr_docx_editable" and meta["ocr"]["empty_pages"] == [2]
    paragraphs = [
        p.text for p in Document(io.BytesIO(await download(offline, editable))).paragraphs
    ]
    assert recall(UZ_LATIN, " ".join(paragraphs)) >= 0.9
    assert "[page 2: no text could be read]" in paragraphs

    events = await stream_chat(
        offline,
        message="Shu skanni asl ko'rinishini saqlab Word ga o'tkaz",
        attachments=[document["id"]],
    )
    [layout] = created(events)
    doc = Document(io.BytesIO(await download(offline, layout)))
    lines = [p.text for p in doc.paragraphs if p.text]
    assert lines[:3] == UZ_LATIN  # one paragraph per recognised line
    source_width = float(PdfReader(io.BytesIO(scan)).pages[0].mediabox.width) / 72
    assert len(doc.sections) == 2  # one section per page, sized like the original page
    assert abs(doc.sections[0].page_width.inches - source_width) < 0.05


@needs_tesseract
async def test_searchable_pdf_keeps_text_pages(offline: AsyncClient) -> None:
    mixed = engine.merge(
        [make_pdf(["Kirish sahifasi matni"]), engine.images_to_pdf([png(page_image(ENGLISH))]).data]
    ).data
    document = await upload(offline, "hisobot.pdf", mixed)
    events = await stream_chat(
        offline, message="Bu PDFni qidiriladigan PDF qil", attachments=[document["id"]]
    )
    [artifact] = created(events)
    version = await meta_of(offline, artifact)
    assert version["generation_metadata"]["ocr_pages"] == [2]
    reader = PdfReader(io.BytesIO(await download(offline, artifact)))
    assert len(reader.pages) == 2
    assert "Kirish sahifasi matni" in reader.pages[0].extract_text()
    assert recall(ENGLISH, reader.pages[1].extract_text()) >= 0.9


async def test_pdf_tables_to_excel(offline: AsyncClient) -> None:
    document = await upload(offline, "aholi.pdf", table_pdf())
    events = await stream_chat(
        offline, message="Shu PDFdagi jadvalni Excelga chiqar", attachments=[document["id"]]
    )
    [artifact] = created(events)
    assert artifact["type"] == "spreadsheet"
    sheet = load_workbook(io.BytesIO(await download(offline, artifact))).worksheets[0]
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    assert rows[0] == ["Viloyat", "Aholi", "O'sish %"]
    assert rows[1] == ["Toshkent shahri", 3040000, 2.1]
    meta = (await meta_of(offline, artifact))["generation_metadata"]
    assert meta["tables"][0]["page"] == 1 and meta["method"] == "PDF text layout"


async def test_spreadsheet_operations_without_a_model(offline: AsyncClient) -> None:
    data = make_xlsx(
        [
            ["Viloyat", "Aholi"],
            ["Toshkent", 350000],
            ["Samarqand", 450000],
            ["Toshkent", 400000],
            ["Samarqand", 450000],
        ]
    )
    document = await upload(offline, "aholi.xlsx", data)
    events = await stream_chat(
        offline,
        message="Dublikatlarni o'chir va Aholi bo'yicha kamayish tartibida sarala",
        attachments=[document["id"]],
    )
    [artifact] = created(events)
    rows = list(
        load_workbook(io.BytesIO(await download(offline, artifact)))
        .worksheets[0]
        .iter_rows(values_only=True)
    )
    assert [r[1] for r in rows[1:]] == [450000, 400000, 350000]
    meta = (await meta_of(offline, artifact))["generation_metadata"]
    assert meta["rows"] == {"before": 4, "after": 3}
    assert "computed by EZO" in meta["note"]
    grouped = await stream_chat(
        offline,
        message="Viloyat bo'yicha guruhlab Aholi yig'indisini hisobla",
        attachments=[document["id"]],
    )
    [summary] = created(grouped)
    rows = list(
        load_workbook(io.BytesIO(await download(offline, summary)))
        .worksheets[0]
        .iter_rows(values_only=True)
    )
    assert rows == [("Viloyat", "Aholi (sum)"), ("Toshkent", 750000), ("Samarqand", 900000)]


async def test_old_scans_are_not_converted_by_unrelated_requests(offline: AsyncClient) -> None:
    await upload(offline, "skan.png", png(page_image(UZ_LATIN)))
    events = await stream_chat(offline, message="Hokimga ariza yozib ber, word fayl qilib")
    assert created(events) == []  # offline: writing needs a model; the scan is left alone
