"""VisualAgent end to end through the chat API (fake models, fake image provider, real DB)."""

from __future__ import annotations

import io
import json
import uuid
from collections.abc import AsyncIterator
from typing import Any

import pytest
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ezo_api.main import create_app
from ezo_api.providers.base import ModelProvider
from ezo_api.visual.service import VisualService
from tests.conftest import make_settings
from tests.fakes import FakeImages, FakeProvider, png_bytes
from tests.integration.helpers import stream_chat
from tests.integration.test_presentation_agent import deck, upload


@pytest.fixture
def images() -> FakeImages:
    return FakeImages()


@pytest.fixture
async def visual_client(
    database_url: str, providers: dict[str, ModelProvider], images: FakeImages
) -> AsyncIterator[AsyncClient]:
    app: FastAPI = create_app(
        make_settings(database_url),
        providers=providers,
        visual=VisualService([images], [images], max_per_request=4, max_per_deck=2),
    )
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as client:
            yield client
    finally:
        await app.state.services.db.dispose()


def items(*objects: dict[str, Any]) -> str:
    return json.dumps({"items": list(objects)}, ensure_ascii=False)


def artifacts_of(events: list[tuple[str, dict[str, Any]]]) -> list[dict[str, Any]]:
    return events[-1][1]["message"]["metadata"].get("artifacts") or []


def system_of(fake: FakeProvider) -> str:
    return fake.requests[-1].system or ""


async def test_generate_then_edit_an_image(
    visual_client: AsyncClient, fake_anthropic: FakeProvider, images: FakeImages
) -> None:
    fake_anthropic.replies = [
        items(
            {
                "title": "Kosmosdagi mushuk",
                "prompt": "A ginger cat floating in space, soft rim light",
                "alt_text": "Kosmosda suzayotgan mushuk",
            }
        ),
        "Tayyor!",
    ]
    events = await stream_chat(visual_client, message="Kosmosdagi mushuk rasmini chiz")
    assert events[0][1]["agent"] == "visual" and events[0][1]["requested_agent_available"]
    assert fake_anthropic.requests[-2].model == "claude-haiku-5-5"  # planning: cheap route
    assert "NEVER ask for text" in (fake_anthropic.requests[-2].system or "")
    [card] = artifacts_of(events)
    assert card["type"] == "image" and card["version"] == 1
    assert images.requests[0].prompt == "A ginger cat floating in space, soft rim light"
    assert "generated with fake-images/fake-image-1" in system_of(fake_anthropic)

    detail = (await visual_client.get(f"/api/artifacts/{card['id']}")).json()
    assert detail["metadata"]["kind"] == "image" and detail["provider"] == "fake-images"
    meta = detail["current_version"]["generation_metadata"]
    assert meta["prompt"].startswith("A ginger cat") and meta["model"] == "fake-image-1"
    assert (meta["width"], meta["height"]) == (128, 128) and meta["placeholder"] is False
    assert meta["revised_prompt"].startswith("revised:")
    preview = await visual_client.get(f"/api/artifacts/{card['id']}/versions/1/exports/preview-1")
    assert preview.headers["content-type"] == "image/png"
    assert preview.headers["content-disposition"].startswith("inline")

    # A follow-up without the word "image" still edits it, as a new version.
    fake_anthropic.replies = [
        json.dumps({"instruction": "Make the background deep blue", "change_summary": "Fon ko'k"}),
        "O'zgartirildi.",
    ]
    conversation_id = events[0][1]["conversation_id"]
    edited = await stream_chat(
        visual_client, message="Fonini to'q ko'k qil", conversation_id=conversation_id
    )
    assert edited[0][1]["agent"] == "visual"
    assert artifacts_of(edited) == [{**card, "version": 2}]
    original, request = images.edits[0]
    assert Image.open(io.BytesIO(original)).size == (128, 128)
    assert request.prompt == "Make the background deep blue"
    versions = (await visual_client.get(f"/api/artifacts/{card['id']}")).json()["versions"]
    assert [v["version_number"] for v in versions] == [2, 1]
    assert versions[0]["change_summary"] == "Fon ko'k"


async def test_no_image_provider_is_honest_and_costs_nothing(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    events = await stream_chat(client, message="Kosmosdagi mushuk rasmini chiz")
    assert events[0][1]["agent"] == "visual"
    assert artifacts_of(events) == []
    assert len(fake_anthropic.requests) == 1  # only the reply, no planning call
    assert "NO image was generated" in system_of(fake_anthropic)
    assert (await client.get("/api/artifacts")).json() == []


async def test_diagram_needs_no_image_provider(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    fake_anthropic.replies = [
        items(
            {
                "title": "EZO arxitekturasi",
                "kind": "flow",
                "nodes": [
                    {"id": "u", "label": "Foydalanuvchi"},
                    {"id": "c", "label": "EzoCore"},
                    {"id": "a", "label": "Agentlar"},
                ],
                "edges": [["u", "c"], ["c", "a"]],
            }
        ),
        "Diagramma tayyor.",
    ]
    events = await stream_chat(client, message="EZO arxitekturasi diagrammasini chiz")
    [card] = artifacts_of(events)
    detail = (await client.get(f"/api/artifacts/{card['id']}")).json()
    meta = detail["current_version"]["generation_metadata"]
    assert meta["kind"] == "diagram" and meta["provider"] == "ezo" and not meta["placeholder"]
    svg = await client.get(f"/api/artifacts/{card['id']}/versions/1/exports/svg")
    assert b"EzoCore" in svg.content
    assert svg.headers["content-disposition"].startswith("attachment")  # never inline SVG
    assert "drawn by EZO's layout engine" in system_of(fake_anthropic)


async def test_social_graphic_without_provider_is_a_labelled_placeholder(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    fake_anthropic.replies = [
        items(
            {
                "title": "Kurs e'loni",
                "headline": "AI kursi boshlanadi",
                "subline": "1-noyabr, Toshkent",
                "image_prompt": "students in a bright modern classroom",
            }
        ),
        "Mana.",
    ]
    events = await stream_chat(client, message="AI kursi uchun Instagram story tayyorla")
    [card] = artifacts_of(events)
    detail = (await client.get(f"/api/artifacts/{card['id']}")).json()
    assert detail["metadata"]["placeholder"] is True
    meta = detail["current_version"]["generation_metadata"]
    assert (meta["width"], meta["height"]) == (1080, 1920)
    assert meta["placeholder_reason"] == "no image provider is configured"
    assert "ABSTRACT PLACEHOLDER" in system_of(fake_anthropic)


async def test_logo_variants_and_text_edit_keep_the_symbol(
    visual_client: AsyncClient, fake_anthropic: FakeProvider, images: FakeImages
) -> None:
    variants = [
        {"title": f"Nur Academy {n}", "brand": "Nur Academy", "symbol_prompt": f"symbol {n}"}
        for n in (1, 2, 3)
    ]
    fake_anthropic.replies = [items(*variants), "3 ta variant tayyor."]
    events = await stream_chat(visual_client, message="Nur Academy uchun logo yarat, 3 ta variant")
    cards = artifacts_of(events)
    assert len(cards) == 3 and len(images.requests) == 3
    assert [r.prompt for r in images.requests] == ["symbol 1", "symbol 2", "symbol 3"]
    assert "All words were rendered by EZO" in system_of(fake_anthropic)

    fake_anthropic.replies = [
        json.dumps(
            {
                "spec": {"brand": "Nur Academy", "tagline": "Bilim nuri"},
                "picture_instruction": "",
                "change_summary": "Shior qo'shildi",
            },
            ensure_ascii=False,
        ),
        "Yangilandi.",
    ]
    conversation_id = events[0][1]["conversation_id"]
    edited = await stream_chat(
        visual_client,
        message="Bu logoga 'Bilim nuri' shiorini qo'sh",
        conversation_id=conversation_id,
    )
    [card] = artifacts_of(edited)
    assert card["id"] == cards[-1]["id"] and card["version"] == 2
    assert len(images.requests) == 3 and images.edits == []  # the symbol was reused
    spec = await visual_client.get(f"/api/artifacts/{card['id']}/versions/2/exports/spec")
    assert json.loads(spec.content)["tagline"] == "Bilim nuri"


async def test_refusal_is_final_and_reported(
    database_url: str, providers: dict[str, ModelProvider], fake_anthropic: FakeProvider
) -> None:
    strict, other = FakeImages(name="strict", fail_with="refused"), FakeImages(name="other")
    app = create_app(
        make_settings(database_url),
        providers=providers,
        visual=VisualService([strict, other], [strict, other]),
    )
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as client:
            fake_anthropic.replies = [
                items(
                    {"title": "x", "prompt": "a cat on a sofa"},
                    {"title": "y", "prompt": "a cat in a garden"},
                ),
                "Afsus.",
            ]
            events = await stream_chat(client, message="2 ta mushuk rasmini chiz")
            assert artifacts_of(events) == []
            assert len(strict.requests) == 1 and other.requests == []
            assert "declined this request" in system_of(fake_anthropic)
    finally:
        await app.state.services.db.dispose()


async def test_uploaded_image_is_edited_into_a_new_artifact(
    visual_client: AsyncClient, fake_anthropic: FakeProvider, images: FakeImages
) -> None:
    document = await upload(visual_client, "mahsulot.png", png_bytes(300, 200, "#ffffff"))
    fake_anthropic.replies = [
        json.dumps({"instruction": "Remove the background", "change_summary": "Fon olindi"}),
        "Tayyor.",
    ]
    events = await stream_chat(visual_client, message="Bu rasmning fonini olib tashla")
    [card] = artifacts_of(events)
    assert card["version"] == 1
    assert Image.open(io.BytesIO(images.edits[0][0])).size == (300, 200)
    detail = (await visual_client.get(f"/api/artifacts/{card['id']}")).json()
    assert detail["source_references"][0]["document_id"] == document["id"]
    meta = detail["current_version"]["generation_metadata"]
    assert meta["aspect_ratio"] == "3:2" and meta["edit"] is True


def hero_deck() -> dict[str, Any]:
    data = deck()
    data["slides"].insert(
        1,
        {
            "layout": "hero_image",
            "title": "Kelajak xizmati",
            "body": "Hammasi bir oynada.",
            "image_prompt": "a bright modern public service centre",
            "image_artifact_id": str(uuid.uuid4()),  # anything the model writes is ignored
        },
    )
    return data


async def test_presentation_gets_generated_slide_images_and_keeps_them_on_edit(
    visual_client: AsyncClient, fake_anthropic: FakeProvider, images: FakeImages
) -> None:
    fake_anthropic.replies = [json.dumps(hero_deck(), ensure_ascii=False), "Tayyor."]
    first = await stream_chat(visual_client, message="Raqamli xizmatlar haqida taqdimot tayyorla")
    [card] = artifacts_of(first)
    assert card["type"] == "presentation"
    assert len(images.requests) == 1
    assert images.requests[0].aspect_ratio == "3:4"
    assert "No text, letters or logos" in images.requests[0].prompt
    assert "1 slide image(s) were generated with fake-images/fake-image-1" in system_of(
        fake_anthropic
    )
    pptx = await visual_client.get(f"/api/artifacts/{card['id']}/versions/1/download")
    slide = Presentation(io.BytesIO(pptx.content)).slides[1]
    assert any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in slide.shapes)
    spec = json.loads(
        (await visual_client.get(f"/api/artifacts/{card['id']}/versions/1/exports/spec")).content
    )
    image_id = spec["slides"][1]["image_artifact_id"]
    image = (await visual_client.get(f"/api/artifacts/{image_id}")).json()
    assert image["type"] == "image"
    assert image["current_version"]["generation_metadata"]["used_in"] == "Raqamli xizmatlar"
    detail = (await visual_client.get(f"/api/artifacts/{card['id']}")).json()
    story = (await visual_client.get(f"/api/stories/{detail['metadata']['story_id']}")).json()
    assert story["scenes"][1]["metadata"]["image_artifact_id"] == image_id

    edited = hero_deck()
    edited["slides"][2]["title"] = "Asosiy muammo"
    fake_anthropic.replies = [json.dumps(edited, ensure_ascii=False), "Yangilandi."]
    await stream_chat(
        visual_client,
        message="3-slayd sarlavhasini o'zgartir",
        conversation_id=first[0][1]["conversation_id"],
    )
    assert len(images.requests) == 1  # the slide image was reused, not paid for again
    spec2 = json.loads(
        (await visual_client.get(f"/api/artifacts/{card['id']}/versions/2/exports/spec")).content
    )
    assert spec2["slides"][1]["image_artifact_id"] == image_id
    assert "1 existing slide image(s) were kept" in system_of(fake_anthropic)


async def test_presentation_without_provider_says_visuals_are_abstract(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    fake_anthropic.replies = [json.dumps(hero_deck(), ensure_ascii=False), "Tayyor."]
    events = await stream_chat(client, message="Raqamli xizmatlar haqida taqdimot tayyorla")
    [card] = artifacts_of(events)
    assert "abstract visual (not a real image)" in system_of(fake_anthropic)
    spec = json.loads(
        (await client.get(f"/api/artifacts/{card['id']}/versions/1/exports/spec")).content
    )
    assert spec["slides"][1]["image_artifact_id"] == ""  # the model's made-up id was dropped
    assert len((await client.get("/api/artifacts")).json()) == 1  # no image artifacts


async def test_status_lists_image_providers(
    visual_client: AsyncClient, client: AsyncClient
) -> None:
    status = (await visual_client.get("/api/status")).json()
    assert status["image_providers"] == [
        {"name": "fake-images", "model": "fake-image-1", "edit": True}
    ]
    features = {f["key"]: f for f in status["features"]}
    assert features["visual"]["maturity"] == "alpha"
    assert (await client.get("/api/status")).json()["image_providers"] == []
