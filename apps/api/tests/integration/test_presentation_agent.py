"""PresentationAgent end to end: request → deck JSON (fake model) → QA → PPTX artifact + story."""

from __future__ import annotations

import io
import json
from typing import Any

from fastapi import FastAPI
from httpx import AsyncClient
from pptx import Presentation

from ezo_api.main import create_app
from tests.conftest import make_settings
from tests.fakes import FakeProvider
from tests.integration.helpers import stream_chat

FACT = "Xatirchi tumanida 2025-yilda 12 400 ta ariza qabul qilingan, 2026-yilda 18 900 ta."


def deck(**overrides: Any) -> dict[str, Any]:
    data: dict[str, Any] = {
        "title": "Raqamli xizmatlar",
        "subtitle": "Tuman uchun reja",
        "theme": "civic",
        "audience": "hokimlik",
        "goal": "loyihani tasdiqlash",
        "slides": [
            {"layout": "cover", "title": "Raqamli xizmatlar", "subtitle": "Tuman uchun reja"},
            {
                "layout": "problem_solution",
                "title": "Muammo va yechim",
                "columns": [
                    {"title": "Muammo", "bullets": ["Navbatlar uzun", "Qog‘oz ko‘p"]},
                    {"title": "Yechim", "bullets": ["Bir oyna", "Onlayn ariza"]},
                ],
                "speaker_notes": "Asosiy muammoni tushuntiramiz.",
            },
            {
                "layout": "chart",
                "title": "Arizalar soni",
                "chart": {
                    "type": "column",
                    "categories": ["2025", "2026"],
                    "series": [{"name": "Arizalar", "values": [12400, 18900]}],
                },
                "sources": [1],
            },
            {
                "layout": "process",
                "title": "Qanday ishlaydi",
                "steps": [
                    {"label": "1", "text": "Ariza"},
                    {"label": "2", "text": "Ko‘rib chiqish"},
                    {"label": "3", "text": "Javob"},
                ],
            },
            {"layout": "closing", "title": "Rahmat!", "subtitle": "Savollar"},
        ],
    }
    data.update(overrides)
    return data


async def upload(client: AsyncClient, name: str, data: bytes, **form: str) -> dict[str, Any]:
    response = await client.post(
        "/api/knowledge/documents", files={"file": (name, data)}, data=form
    )
    assert response.status_code == 201, response.text
    return response.json()["document"]  # type: ignore[no-any-return]


async def test_presentation_from_a_file_creates_checked_pptx_story_and_exports(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    await upload(client, "arizalar.txt", FACT.encode())
    fake_anthropic.replies = [
        "Mana reja:\n```json\n" + json.dumps(deck(), ensure_ascii=False) + "\n```",
        "Taqdimot tayyor: 6 slayd [1].",
    ]
    events = await stream_chat(client, message="Ezo, shu fayl asosida 5 slaydli taqdimot tayyorla.")
    meta = events[0][1]
    assert meta["agent"] == "presentation" and meta["requested_agent_available"] is True
    stages = [data["stage"] for name, data in events if name == "activity"]
    assert "reading" in stages and "creating" in stages

    # The deck prompt carried the file as a numbered source, marked as data.
    deck_request = fake_anthropic.requests[0]
    assert deck_request.model == "claude-opus-5-5"  # the deck is written on the reasoning route
    assert FACT in deck_request.messages[0].content
    assert '<source n="1">' in deck_request.messages[0].content
    assert "Slides: 5" in deck_request.messages[0].content
    # The final answer only reports the build and gets the facts as notes.
    assert fake_anthropic.requests[-1].model == "claude-haiku-5-5"  # the report: cheap chat route
    final_system = fake_anthropic.requests[-1].system or ""
    assert "Created presentation “Raqamli xizmatlar”" in final_system
    assert "6 slides" in final_system and "Layout check: passed" in final_system

    done = events[-1][1]
    artifacts = done["message"]["metadata"]["artifacts"]
    assert len(artifacts) == 1 and artifacts[0]["type"] == "presentation"
    artifact = (await client.get(f"/api/artifacts/{artifacts[0]['id']}")).json()
    version = artifact["current_version"]
    assert version["filename"] == "Raqamli xizmatlar.pptx"
    assert version["generation_metadata"]["slides"] == 6  # + the sources slide
    assert version["generation_metadata"]["qa"]["passed"] is True
    assert "exports" not in version["generation_metadata"]  # storage keys stay internal
    names = {e["name"] for e in version["exports"]}
    assert {"pdf", "html", "spec", "preview-1", "preview-6"} <= names
    assert artifact["source_references"][0]["title"] == "arizalar"

    pptx = await client.get(f"/api/artifacts/{artifact['id']}/versions/1/download")
    presentation = Presentation(io.BytesIO(pptx.content))
    assert len(presentation.slides) == 6
    texts = [
        shape.text_frame.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    assert "Manbalar" in texts and any("[1] arizalar" in t for t in texts)
    charts = [s for slide in presentation.slides for s in slide.shapes if s.has_chart]
    assert list(charts[0].chart.plots[0].series[0].values) == [12400, 18900]

    pdf = await client.get(f"/api/artifacts/{artifact['id']}/versions/1/exports/pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert pdf.headers["content-disposition"].startswith("attachment")
    preview = await client.get(f"/api/artifacts/{artifact['id']}/versions/1/exports/preview-1")
    assert preview.content.startswith(b"\x89PNG")
    assert preview.headers["content-disposition"].startswith("inline")
    assert preview.headers["x-content-type-options"] == "nosniff"
    html = await client.get(f"/api/artifacts/{artifact['id']}/versions/1/exports/html")
    assert html.headers["content-disposition"].startswith("attachment")
    assert html.headers["content-security-policy"] == "sandbox"
    missing = await client.get(f"/api/artifacts/{artifact['id']}/versions/1/exports/zip")
    assert missing.status_code == 404

    story_id = artifact["metadata"]["story_id"]
    story = (await client.get(f"/api/stories/{story_id}")).json()
    assert [s["title"] for s in story["scenes"]][:3] == [
        "Raqamli xizmatlar",
        "Muammo va yechim",
        "Arizalar soni",
    ]
    assert story["scenes"][1]["narration"]["text"] == "Asosiy muammoni tushuntiramiz."


async def test_follow_up_edit_creates_a_new_version(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    fake_anthropic.replies = [json.dumps(deck(), ensure_ascii=False), "Tayyor."]
    first = await stream_chat(client, message="Raqamli xizmatlar haqida taqdimot tayyorla")
    conversation_id = first[0][1]["conversation_id"]
    artifact_id = first[-1][1]["message"]["metadata"]["artifacts"][0]["id"]

    edited = deck()
    edited["slides"][1]["title"] = "Asosiy muammo"
    edited["change_summary"] = "2-slayd sarlavhasi qisqartirildi."
    fake_anthropic.replies = [json.dumps(edited, ensure_ascii=False), "Yangilandi."]
    second = await stream_chat(
        client, message="2-slayd sarlavhasini o'zgartir", conversation_id=conversation_id
    )
    assert second[0][1]["agent"] == "presentation"
    edit_request = fake_anthropic.requests[-2]
    assert "Current deck JSON" in edit_request.messages[0].content
    assert "Muammo va yechim" in edit_request.messages[0].content

    reported = second[-1][1]["message"]["metadata"]["artifacts"][0]
    assert reported == {
        "id": artifact_id,
        "title": "Raqamli xizmatlar",
        "type": "presentation",
        "version": 2,
    }
    detail = (await client.get(f"/api/artifacts/{artifact_id}")).json()
    assert [v["version_number"] for v in detail["versions"]] == [2, 1]
    assert detail["versions"][0]["change_summary"] == "2-slayd sarlavhasi qisqartirildi."
    assert "What changed: 2-slayd" in (fake_anthropic.requests[-1].system or "")
    story = (await client.get(f"/api/stories/{detail['metadata']['story_id']}")).json()
    assert story["scenes"][1]["title"] == "Asosiy muammo"
    # Version 1 and its exports are untouched.
    old = await client.get(f"/api/artifacts/{artifact_id}/versions/1/exports/spec")
    assert json.loads(old.content)["slides"][1]["title"] == "Muammo va yechim"

    deleted = await client.delete(f"/api/artifacts/{artifact_id}", params={"confirm": "true"})
    assert deleted.status_code == 204
    gone = await client.get(f"/api/artifacts/{artifact_id}/versions/1/exports/pdf")
    assert gone.status_code == 404


async def test_layout_problems_get_one_repair_round(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    broken = deck()
    broken["slides"][2] = {"layout": "chart", "title": "Arizalar soni"}  # chart without data
    fake_anthropic.replies = [
        json.dumps(broken, ensure_ascii=False),
        json.dumps(deck(), ensure_ascii=False),
        "Tayyor.",
    ]
    events = await stream_chat(client, message="Arizalar haqida taqdimot tayyorla")
    repair = fake_anthropic.requests[1].messages[-1].content
    assert "layout check found problems" in repair and "slide 3" in repair
    artifact_id = events[-1][1]["message"]["metadata"]["artifacts"][0]["id"]
    artifact = (await client.get(f"/api/artifacts/{artifact_id}")).json()
    assert artifact["current_version"]["generation_metadata"]["qa"]["passed"] is True
    run = (await client.get(f"/api/agent-runs/{events[0][1]['run_id']}")).json()
    tools = [c["tool"] for c in run["tool_calls"]]
    assert tools[:3] == ["plan_deck", "layout_qa", "repair_layout"]
    assert "create_presentation" in tools


async def test_invalid_model_output_is_reported_honestly(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    fake_anthropic.replies = ["Kechirasiz, bu JSON emas.", "Hali ham emas.", "Uzr."]
    events = await stream_chat(client, message="AI haqida 10 slaydli taqdimot qil")
    assert events[-1][0] == "done"
    assert events[-1][1]["message"]["metadata"]["artifacts"] == []
    system = fake_anthropic.requests[-1].system or ""
    assert "could not be built" in system and "Nothing was created" in system
    assert (await client.get("/api/artifacts")).json() == []


async def test_offline_mode_creates_nothing(database_url: str) -> None:
    app: FastAPI = create_app(make_settings(database_url), providers={})
    try:
        from httpx import ASGITransport

        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as client:
            events = await stream_chat(client, message="15 slaydli taqdimot tayyorla")
            assert events[0][1]["agent"] == "presentation"
            assert events[-1][1]["message"]["metadata"]["artifacts"] == []
            assert (await client.get("/api/artifacts")).json() == []
    finally:
        await app.state.services.db.dispose()
