"""Education materials through chat: a lesson package (Word + Excel + slides), a quiz from
an attached lecture, a template without a model, Uzbek Cyrillic output."""

from __future__ import annotations

import io
import json
from typing import Any

from docx import Document
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from pptx import Presentation

from ezo_api.main import create_app
from tests.conftest import make_settings
from tests.fakes import FakeProvider
from tests.integration.helpers import stream_chat
from tests.integration.test_document_agent import created, upload

LESSON = {
    "title": "Fotosintez",
    "subject": "Biologiya",
    "level": "7-sinf",
    "duration_minutes": 45,
    "objectives": ["Fotosintez jarayonini tushuntirish"],
    "stages": [
        {"name": "Kirish", "minutes": 5, "teacher": "Savol beradi", "students": "Javob beradi"},
        {"name": "Yangi mavzu", "minutes": 20, "teacher": "Tushuntiradi", "students": "Yozadi"},
        {"name": "Mashq", "minutes": 15, "teacher": "Kuzatadi", "students": "Bajaradi"},
        {"name": "Xulosa", "minutes": 5, "teacher": "Umumlashtiradi", "students": "Savol beradi"},
    ],
    "lecture_notes": [
        {"heading": "Fotosintez nima?", "text": "O'simliklar yorug'likdan foydalanadi."}
    ],
    "questions": [
        {
            "question": "Fotosintez qayerda boradi?",
            "options": ["Xloroplast", "Mitoxondriya", "Yadro", "Ribosoma"],
            "correct": 0,
            "explanation": "Xlorofill xloroplastda.",
        }
    ],
    "homework": ["Darslikdagi 3-savolga javob yozing"],
    "criteria": [
        {
            "criterion": "Tushunish",
            "levels": [{"label": "A'lo", "points": 5}, {"label": "Yaxshi", "points": 4}],
        }
    ],
    "slides": [
        {"title": "Fotosintez nima?", "bullets": ["Yorug'lik energiyasi", "Xlorofill"]},
        {"title": "Natija", "bullets": ["Glyukoza", "Kislorod"]},
    ],
    "sources_needed": ["7-sinf biologiya darsligi"],
}


async def detail(client: AsyncClient, artifact: dict[str, Any]) -> tuple[dict[str, Any], bytes]:
    version = (await client.get(f"/api/artifacts/{artifact['id']}")).json()["current_version"]
    data = (await client.get(f"/api/artifacts/{artifact['id']}/versions/1/download")).content
    return version, data


async def test_lesson_package_makes_word_excel_and_slides(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    fake_anthropic.replies = [json.dumps(LESSON, ensure_ascii=False), "Dars ishlanmasi tayyor."]
    events = await stream_chat(
        client, message="7-sinf uchun fotosintez mavzusida 45 daqiqalik dars ishlanmasi tayyorla"
    )
    assert events[0][1]["agent"] == "document"
    writer = fake_anthropic.requests[0].system or ""
    assert (
        "Prepare: Lesson package" in writer and "add up to the lesson length (45 minutes)" in writer
    )
    word, sheet, deck = created(events)
    assert (word["type"], sheet["type"], deck["type"]) == (
        "document",
        "spreadsheet",
        "presentation",
    )
    version, data = await detail(client, word)
    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert "Dars bosqichlari" in text and "Javoblar kaliti (o'qituvchi uchun)" in text
    meta = version["generation_metadata"]
    assert meta["kind"] == "lesson" and meta["check"]["total_minutes"] == 45
    assert meta["check"]["issues"] == [] and meta["check"]["answer_key"][0]["answer"] == "A"
    assert {e["name"] for e in version["exports"]} >= {"pdf"}
    _, data = await detail(client, sheet)
    workbook = load_workbook(io.BytesIO(data))
    assert workbook.sheetnames == ["Test savollari", "Dars bosqichlari", "Baholash mezonlari"]
    assert workbook["Test savollari"]["G2"].value == "A"
    _, data = await detail(client, deck)
    assert len(Presentation(io.BytesIO(data)).slides) == 3
    final = fake_anthropic.requests[-1].system or ""
    assert "answer key was built by EZO" in final and "7-sinf biologiya darsligi" in final


async def test_quiz_from_an_attached_lecture_reports_problems(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    lecture = await upload(
        client, "maruza.txt", b"Fotosintez xloroplastda boradi. Natijada kislorod ajraladi."
    )
    quiz = {
        "title": "Fotosintez testi",
        "questions": [
            {
                "question": "Fotosintez qayerda boradi?",
                "options": ["Xloroplast", "Yadro"],
                "correct": 0,
            },
            {"question": "Nima ajraladi?", "options": ["Kislorod", "Azot"], "correct": 7},
        ],
    }
    fake_anthropic.replies = [json.dumps(quiz, ensure_ascii=False), "Tayyor."]
    events = await stream_chat(
        client, message="Shu ma'ruzadan 3 ta test tuz", attachments=[lecture["id"]]
    )
    word, _sheet = created(events)
    sent = fake_anthropic.requests[0]
    assert "Exactly 3 questions" in (sent.system or "")
    assert "Fotosintez xloroplastda boradi" in sent.messages[0].content  # based on the file
    version, _ = await detail(client, word)
    assert version["generation_metadata"]["check"]["issues"] == [
        "Question 2: no valid correct answer",
        "2 questions instead of 3",
    ]
    assert version["source_references"][0]["title"] == "maruza"
    final = fake_anthropic.requests[-1].system or ""
    assert "Question 2: no valid correct answer" in final


async def test_template_without_a_model_and_cyrillic(database_url: str) -> None:
    app: FastAPI = create_app(make_settings(database_url), providers={})
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as client:
            events = await stream_chat(client, message="Kirillda 4 haftalik o'qish rejasini tuz")
            word, _sheet = created(events)
            version, data = await detail(client, word)
            assert version["generation_metadata"]["template"] is True
            assert version["generation_metadata"]["language"] == "uz-cyrl"
            text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
            assert "Ҳафта 4" in text and "Hafta" not in text
            events = await stream_chat(
                client, message="45 daqiqalik dars ishlanmasi shablonini ber"
            )
            word, _sheet = created(events)  # no slides from a template
            version, _ = await detail(client, word)
            assert version["generation_metadata"]["check"]["total_minutes"] == 45
    finally:
        await app.state.services.db.dispose()
