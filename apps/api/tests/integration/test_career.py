"""Career profile: the API, a CV built exactly from the profile (no model), application
letters that use only the profile's facts, and interview preparation."""

from __future__ import annotations

import io
import json
from typing import Any

from docx import Document
from httpx import AsyncClient

from ezo_api.career.profile import CareerProfileIn, cv_markdown, facts_text
from tests.fakes import FakeProvider
from tests.integration.helpers import stream_chat
from tests.integration.test_document_agent import created

PROFILE: dict[str, Any] = {
    "full_name": "Dilnoza Yusupova",
    "headline": "Ma'lumotlar tahlilchisi",
    "email": "dilnoza@example.uz",
    "location": "Toshkent",
    "summary": "Davlat xizmatlari ma'lumotlarini tahlil qilaman.",
    "experience": [
        {
            "role": "Tahlilchi",
            "organization": "Raqamli texnologiyalar vazirligi",
            "start": "2021",
            "end": "",
            "achievements": ["Murojaatlarni ko'rib chiqish vaqtini 30% ga qisqartirdim"],
        }
    ],
    "education": [
        {"institution": "TATU", "degree": "Bakalavr", "field": "Informatika", "end": "2020"}
    ],
    "skills": ["Python", "SQL", "Power BI"],
    "languages": [{"language": "Ingliz tili", "level": "B2"}],
    "certificates": [{"name": "Google Data Analytics", "year": "2022"}],
}


def test_cv_markdown_and_facts() -> None:
    profile = CareerProfileIn.model_validate(PROFILE)
    markdown = cv_markdown(profile, "uz")
    assert markdown.startswith("# Dilnoza Yusupova\n\n**Ma'lumotlar tahlilchisi**")
    assert "**Tahlilchi** — Raqamli texnologiyalar vazirligi (2021 – hozirgacha)" in markdown
    assert "- Murojaatlarni ko'rib chiqish vaqtini 30% ga qisqartirdim" in markdown
    assert "## Ko'nikmalar\n\nPython, SQL, Power BI" in markdown
    assert "## Опыт работы" in cv_markdown(profile, "ru")
    assert "# Дилноза Юсупова" in cv_markdown(profile, "uz-cyrl")
    facts = facts_text(profile)
    assert "Experience: Tahlilchi, Raqamli texnologiyalar vazirligi (2021 – present)" in facts
    assert "Languages: Ingliz tili B2" in facts


async def test_profile_api_and_cv_without_a_model(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    assert (await client.get("/api/career/profile")).json() is None
    saved = await client.put("/api/career/profile", json=PROFILE)
    assert saved.status_code == 200 and saved.json()["skills"] == ["Python", "SQL", "Power BI"]
    updated = await client.put("/api/career/profile", json={**PROFILE, "skills": ["Python"]})
    assert updated.json()["skills"] == ["Python"]  # one profile per user, replaced
    bad = await client.put("/api/career/profile", json={**PROFILE, "full_name": ""})
    assert bad.status_code == 422

    fake_anthropic.replies = ["CV tayyor."]
    events = await stream_chat(client, message="Menga CV tayyorla")
    [artifact] = created(events)
    assert len(fake_anthropic.requests) == 1  # only the reply: the CV used no model
    version = (await client.get(f"/api/artifacts/{artifact['id']}")).json()["current_version"]
    assert version["generation_metadata"]["source"] == "career_profile"
    data = (await client.get(f"/api/artifacts/{artifact['id']}/versions/1/download")).content
    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert "Dilnoza Yusupova" in text and "Murojaatlarni ko'rib chiqish vaqtini 30%" in text
    assert "EXACTLY from the user's career profile" in (fake_anthropic.requests[-1].system or "")

    assert (await client.delete("/api/career/profile")).status_code == 204
    assert (await client.delete("/api/career/profile")).status_code == 404


async def test_letters_and_interview_use_only_profile_facts(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    await client.put("/api/career/profile", json=PROFILE)
    letter = {
        "title": "Cover letter",
        "fields": {"name": "Dilnoza Yusupova", "position": "Data analyst", "company": "Bank"},
        "sections": [
            {
                "heading": None,
                "text": "Vazirlikda murojaatlarni ko'rib chiqish vaqtini 30% ga, xarajatlarni 45% "
                "ga qisqartirdim.",
            }
        ],
        "missing": [],
    }
    fake_anthropic.replies = [json.dumps(letter, ensure_ascii=False), "Tayyor."]
    await stream_chat(client, message="Bankdagi data analyst lavozimi uchun cover letter yoz")
    writer = fake_anthropic.requests[0].messages[0].content
    assert "Career profile (entered by the user; the ONLY facts" in writer
    assert "Raqamli texnologiyalar vazirligi" in writer
    final = fake_anthropic.requests[-1].system or ""
    # 30 comes from the profile; 45 was invented and is flagged.
    line = next(ln for ln in final.splitlines() if "ask them to check" in ln)
    assert "45" in line and "30" not in line
    assert "uses the user's career profile as its only facts" in final

    interview = {
        "title": "Intervyuga tayyorgarlik",
        "questions": [
            {
                "question": "Eng katta yutug'ingiz?",
                "answer": "Vazirlikdagi 30% qisqartirishni ayting",
                "explanation": "Natijaga yo'naltirilganlik",
            }
        ],
    }
    fake_anthropic.replies = [json.dumps(interview, ensure_ascii=False), "Tayyor."]
    events = await stream_chat(
        client, message="Data analyst lavozimiga intervyu savollarini tayyorla"
    )
    [artifact] = created(events)
    sent = fake_anthropic.requests[-2]
    assert "Prepare: Interview preparation" in (sent.system or "")
    assert "Skills: Python, SQL, Power BI" in sent.messages[0].content
    data = (await client.get(f"/api/artifacts/{artifact['id']}/versions/1/download")).content
    text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    assert "Javob rejasi: Vazirlikdagi 30% qisqartirishni ayting" in text
