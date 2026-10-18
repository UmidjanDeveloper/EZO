"""Word files edited in place through chat: translation keeps the formatting, missing
paragraphs are retried, find-and-replace works offline, tracked changes on request."""

from __future__ import annotations

import io

from docx import Document
from docx.oxml.ns import qn
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient

from ezo_api.main import create_app
from tests.conftest import make_settings
from tests.docx_fixtures import rich_docx
from tests.fakes import FakeProvider
from tests.integration.helpers import stream_chat
from tests.integration.test_document_agent import created, upload

DATA = rich_docx()
REPLY = "\n".join(
    [
        "[[0]] Lease agreement",
        "[[1]] The rent is ⟦1⟧4 500 000 soums⟦/1⟧ per month, payable by the 5th of each month.",
        "[[2]] More information is ⟦1⟧on the company website⟦/1⟧.",
        "[[3]] ⟦1⟧Important: ⟦/1⟧the contract is made in two copies.",
        "[[4]] The tenant keeps the premises clean.",
        "[[5]] The tenant pays utility bills on time.",
        "[[6]] Party",
        "[[7]] Signature",
        "[[8]] Tenant",
        "[[11]] The contract enters into force on the day it is signed.",
        "[[12]] LLC «Ezgu Kelajak»",
        "[[13]] Page ",
    ]
)


async def download(client: AsyncClient, artifact_id: str, version: int = 1) -> bytes:
    response = await client.get(f"/api/artifacts/{artifact_id}/versions/{version}/download")
    assert response.status_code == 200
    return response.content


async def test_translation_keeps_the_original_formatting(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    document = await upload(client, "shartnoma.docx", DATA)
    fake_anthropic.replies = [REPLY, "Tarjima tayyor."]
    events = await stream_chat(
        client, message="Shu faylni ingliz tiliga tarjima qil", attachments=[document["id"]]
    )
    [artifact] = created(events)
    sent = fake_anthropic.requests[0]
    assert "[[1]] Ijara narxi oyiga ⟦1⟧4 500 000 so'm⟦/1⟧" in sent.messages[0].content
    assert "[[9]]" not in sent.messages[0].content  # "________" has no words to translate
    assert "EVERY input paragraph" in (sent.system or "")
    data = await download(client, artifact["id"])
    doc = Document(io.BytesIO(data))
    assert [(r.text, r.bold) for r in doc.paragraphs[1].runs][1] == ("4 500 000 soums", True)
    assert doc.paragraphs[0].style.name == "Heading 1"
    assert doc.paragraphs[4].style.name == "List Number"
    assert doc.tables[0].cell(1, 0).text == "Tenant"
    assert doc.sections[0].header.paragraphs[0].text == "LLC «Ezgu Kelajak»"
    assert doc.element.body.find(".//" + qn("w:hyperlink")) is not None

    detail = (await client.get(f"/api/artifacts/{artifact['id']}")).json()
    version = detail["current_version"]
    assert version["filename"] == "shartnoma (en).docx"
    meta = version["generation_metadata"]
    assert meta["fidelity"] == "in_place" and meta["action"] == "translate"
    assert meta["changed_properties"] == []
    assert {"tables", "images", "headers", "footers", "hyperlinks"} <= set(meta["preserved"])
    assert meta["inputs"][0]["id"] == document["id"] and len(meta["inputs"][0]["sha256"]) == 64
    assert meta["edit"]["changed_paragraphs"] == 12 and meta["model_calls"] == 1
    system = fake_anthropic.requests[-1].system or ""
    assert "IN THE ORIGINAL WORD FILE" in system and "the original upload is unchanged" in system

    original = await client.get(f"/api/knowledge/documents/{document['id']}")
    assert original.json()["size_bytes"] == len(DATA)


async def test_paragraphs_the_model_skips_are_retried_then_reported(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    document = await upload(client, "shartnoma.docx", DATA)
    first = "\n".join(line for line in REPLY.splitlines() if not line.startswith("[[5]]"))
    fake_anthropic.replies = [first, "", "Tayyor."]  # the retry also returns nothing
    events = await stream_chat(
        client, message="Shu faylni ingliz tiliga tarjima qil", attachments=[document["id"]]
    )
    [artifact] = created(events)
    retry = fake_anthropic.requests[1].messages[0].content
    assert "[[5]] Ijarachi kommunal" in retry and "[[1]]" not in retry
    meta = (await client.get(f"/api/artifacts/{artifact['id']}")).json()["current_version"][
        "generation_metadata"
    ]
    assert any("without a translation" in w for w in meta["warnings"])
    doc = Document(io.BytesIO(await download(client, artifact["id"])))
    assert doc.paragraphs[5].text == "Ijarachi kommunal to'lovlarni o'z vaqtida to'laydi."


async def test_tracked_changes_on_request(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    document = await upload(client, "shartnoma.docx", DATA)
    fake_anthropic.replies = [
        "[[5]] Ijarachi kommunal to'lovlarni vaqtida to'laydi.",
        "Tayyor.",
    ]
    events = await stream_chat(
        client,
        message="Shu hujjatni imlo xatolaridan tuzat, o'zgarishlar ko'rinsin (track changes)",
        attachments=[document["id"]],
    )
    [artifact] = created(events)
    body = Document(io.BytesIO(await download(client, artifact["id"]))).element.body
    deleted = body.findall(".//" + qn("w:del"))
    assert len(deleted) == 1 and deleted[0].get(qn("w:author")) == "EZO"
    meta = (await client.get(f"/api/artifacts/{artifact['id']}")).json()["current_version"][
        "generation_metadata"
    ]
    assert meta["tracked_changes"] is True and meta["edit"]["changed_paragraphs"] == 1


async def test_find_and_replace_needs_no_model(database_url: str) -> None:
    app: FastAPI = create_app(make_settings(database_url), providers={})
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as client:
            document = await upload(client, "shartnoma.docx", DATA)
            events = await stream_chat(
                client,
                message="«Ijarachi» ni «Arendator» ga almashtir",
                attachments=[document["id"]],
            )
            [artifact] = created(events)
            doc = Document(io.BytesIO(await download(client, artifact["id"])))
            assert doc.paragraphs[4].text == "Arendator binoni toza saqlaydi."
            assert doc.paragraphs[4].style.name == "List Number"
            assert doc.tables[0].cell(1, 0).text == "Arendator"
            meta = (await client.get(f"/api/artifacts/{artifact['id']}")).json()["current_version"][
                "generation_metadata"
            ]
            assert meta["replacements"] == 3 and meta["engine"].endswith("(no AI model)")
            missing = await stream_chat(
                client, message="«Yo'q so'z» ni «X» ga almashtir", attachments=[document["id"]]
            )
            assert created(missing) == []
    finally:
        await app.state.services.db.dispose()
