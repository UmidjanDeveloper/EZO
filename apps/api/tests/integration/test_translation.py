"""Translation through chat and the API: text in the message, Word files translated in
place (also into Uzbek Cyrillic, with a bilingual export), the glossary, Latin ↔ Cyrillic
without a model, and turn-based voice interpretation."""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient

from ezo_api.db.models import GlossaryTerm, User
from ezo_api.main import create_app
from ezo_api.voice.service import VoiceService
from tests.conftest import make_settings
from tests.fakes import FakeProvider, FakeVoice
from tests.integration.helpers import stream_chat
from tests.integration.test_document_agent import created, upload

WEBM = b"\x1a\x45\xdf\xa3" + b"\x00" * 64


def russian_docx() -> bytes:
    document = Document()
    document.add_heading("Протокол", level=1)
    document.add_paragraph("Совещание состоялось 12 марта 2024 года в хокимияте района.")
    paragraph = document.add_paragraph()
    paragraph.add_run("Ответственный: ").bold = True
    paragraph.add_run("Каримов А.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def download(client: AsyncClient, artifact: dict[str, Any]) -> tuple[dict[str, Any], bytes]:
    detail = (await client.get(f"/api/artifacts/{artifact['id']}")).json()["current_version"]
    number = detail["version_number"]
    response = await client.get(f"/api/artifacts/{artifact['id']}/versions/{number}/download")
    assert response.status_code == 200
    return detail, response.content


async def test_text_in_the_message_with_glossary_and_checks(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    term = await client.post(
        "/api/translation/glossary",
        json={
            "source_language": "uz",
            "target_language": "en",
            "source_term": "hokimlik",
            "target_term": "khokimiyat",
        },
    )
    assert term.status_code == 201
    fake_anthropic.replies = [
        "On March 12, 2024 the Tashkent district khokimiyat held a meeting.\n=====\nNOTES: none",
        "Here is the translation.",
    ]
    events = await stream_chat(
        client,
        message="Inglizchaga tarjima qil: Toshkent tuman hokimligi 2024-yil 12-mart kuni "
        "yig'ilish o'tkazdi.",
    )
    assert events[0][1]["agent"] == "document"
    writer = fake_anthropic.requests[0]
    assert "from Uzbek (Latin script) into English" in (writer.system or "")
    assert "“hokimlik” → “khokimiyat”" in (writer.system or "")
    assert "<text>\nToshkent tuman hokimligi" in writer.messages[0].content
    final = fake_anthropic.requests[-1].system or ""
    assert "<translation>\nOn March 12, 2024 the Tashkent district khokimiyat" in final
    assert "1 glossary term(s) applied" in final and "found no problem" in final


async def test_problems_are_reported_not_hidden(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    fake_anthropic.replies = [
        "Привет, Рахимов. Встреча 21 мая.\n=====\nNOTES:\n- “uchrashuv” может быть “свидание”",
        "Готово.",
    ]
    await stream_chat(
        client, message="Ruschaga tarjima qil: Salom, Karimov. Uchrashuv 12-may kuni."
    )
    final = fake_anthropic.requests[-1].system or ""
    assert "Numbers in the source but not in the translation: 12" in final
    assert "Numbers in the translation that are not in the source: 21" in final
    assert "Names to check (not found unchanged in the translation): Karimov" in final
    assert "“uchrashuv” может быть “свидание”" in final


async def test_russian_word_file_into_uzbek_cyrillic_with_bilingual_export(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    await client.post(
        "/api/translation/glossary",
        json={
            "source_language": "ru",
            "target_language": "uz",
            "source_term": "хокимият",
            "target_term": "hokimlik",
        },
    )
    document = await upload(client, "protokol.docx", russian_docx())
    fake_anthropic.replies = [
        "\n".join(
            [
                "[[0]] Bayonnoma",
                "[[1]] Yig'ilish 2024-yil 12-mart kuni tuman hokimligida bo'lib o'tdi.",
                "[[2]] Mas'ul: ⟦1⟧Karimov A.⟦/1⟧",
            ]
        ),
        "Tarjima tayyor.",
    ]
    events = await stream_chat(
        client,
        message="Shu ruscha hujjatni o'zbek kirilliga tarjima qil, ikki tilli variantini ham ber",
        attachments=[document["id"]],
    )
    [artifact] = created(events)
    editor = fake_anthropic.requests[0].system or ""
    assert "into Uzbek (Latin script)" in editor  # the model writes Latin; EZO converts
    assert "Source language: Russian" in editor and "“хокимият” → “hokimlik”" in editor
    version, data = await download(client, artifact)
    word = Document(io.BytesIO(data))
    assert [p.text for p in word.paragraphs] == [
        "Баённома",
        "Йиғилиш 2024-йил 12-март куни туман ҳокимлигида бўлиб ўтди.",
        "Масъул: Каримов А.",
    ]
    assert word.paragraphs[0].style.name == "Heading 1"
    assert [(r.text, r.bold) for r in word.paragraphs[2].runs] == [
        ("Масъул: ", True),
        ("Каримов А.", None),
    ]
    meta = version["generation_metadata"]
    assert meta["translation"]["source_language"] == "ru"
    assert meta["translation"]["target_language"] == "uz-cyrl"
    assert meta["translation_quality"]["ok"] is True
    assert meta["translation_quality"]["terms_used"] == 1
    assert {e["name"] for e in version["exports"]} >= {"markdown", "bilingual"}
    bilingual = await client.get(f"/api/artifacts/{artifact['id']}/versions/1/exports/bilingual")
    assert bilingual.status_code == 200
    rows = Document(io.BytesIO(bilingual.content)).tables[0].rows
    assert [c.text for c in rows[2].cells] == [
        "Совещание состоялось 12 марта 2024 года в хокимияте района.",
        "Йиғилиш 2024-йил 12-март куни туман ҳокимлигида бўлиб ўтди.",
    ]
    final = fake_anthropic.requests[-1].system or ""
    assert "Translated from Russian into Uzbek (Cyrillic script) with 1 glossary term" in final


async def _offline(database_url: str) -> tuple[FastAPI, AsyncClient]:
    app = create_app(make_settings(database_url), providers={})
    return app, AsyncClient(transport=ASGITransport(app=app), base_url="http://t")


async def test_script_conversion_works_without_a_model(database_url: str) -> None:
    app, client = await _offline(database_url)
    try:
        async with client:
            events = await stream_chat(
                client, message="Kirillga o'gir: O'zbekiston Respublikasi Prezidentining qarori"
            )
            [artifact] = created(events)
            version, data = await download(client, artifact)
            assert data.decode() == "Ўзбекистон Республикаси Президентининг қарори"
            assert version["generation_metadata"]["engine"].startswith("EZO transliteration")

            word = Document()
            word.add_paragraph().add_run("Toshkent shahar hokimligi").bold = True
            buffer = io.BytesIO()
            word.save(buffer)
            document = await upload(client, "xat.docx", buffer.getvalue())
            events = await stream_chat(
                client, message="Shu faylni kirill yozuviga o'gir", attachments=[document["id"]]
            )
            [artifact] = created(events)
            version, data = await download(client, artifact)
            run = Document(io.BytesIO(data)).paragraphs[0].runs[0]
            assert (run.text, run.bold) == ("Тошкент шаҳар ҳокимлиги", True)
            assert version["filename"] == "xat (Кирилл).docx"
            assert version["generation_metadata"]["fidelity"] == "in_place"

            # Translation itself needs a model: said plainly, nothing invented.
            events = await stream_chat(client, message="Inglizchaga tarjima qil: Salom dunyo")
            assert events[-1][0] == "done" and created(events) == []
            api = await client.post(
                "/api/translation/translate", json={"text": "Salom", "target": "en"}
            )
            assert api.status_code == 503 and api.json()["error"]["code"] == "offline"
            # …but an Uzbek script change through the API does not.
            api = await client.post(
                "/api/translation/translate", json={"text": "Salom dunyo", "target": "uz-cyrl"}
            )
            assert api.status_code == 200
            assert api.json()["text"] == "Салом дунё" and api.json()["method"] == "transliteration"
    finally:
        await app.state.services.db.dispose()


async def test_translate_and_transliterate_api(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    fake_anthropic.replies = ["Добрый день, коллеги.\n=====\nNOTES: none"]
    response = await client.post(
        "/api/translation/translate",
        json={"text": "Xayrli kun, hamkasblar.", "target": "ru", "domain": "official"},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["text"] == "Добрый день, коллеги." and body["source_language"] == "uz"
    assert body["warnings"] == [] and body["quality"]["ok"] is True
    assert "Official correspondence" in (fake_anthropic.requests[-1].system or "")
    same = await client.post(
        "/api/translation/translate", json={"text": "Hello there", "target": "en"}
    )
    assert same.status_code == 422
    converted = await client.post(
        "/api/translation/transliterate", json={"text": "Ўзбекистон", "to": "latin"}
    )
    assert converted.json() == {
        "text": "O'zbekiston",
        "from_script": "cyrillic",
        "to": "latin",
        "engine": "rules",
    }
    languages = (await client.get("/api/translation/languages")).json()
    assert [entry["code"] for entry in languages["languages"]] == ["uz", "uz-cyrl", "ru", "en"]
    assert languages["interpretation"] == "turn-based"


async def test_glossary_api_chat_and_privacy(
    client: AsyncClient, app: FastAPI, fake_anthropic: FakeProvider
) -> None:
    org = (await client.post("/api/writing/organizations", json={"name": "Hokimlik"})).json()
    created_term = await client.post(
        "/api/translation/glossary",
        json={
            "source_language": "uz",
            "target_language": "en",
            "source_term": " tuman ",
            "target_term": "district",
            "note": "official usage",
            "organization_id": org["id"],
        },
    )
    assert created_term.status_code == 201
    term = created_term.json()
    assert term["source_term"] == "tuman" and term["organization_id"] == org["id"]
    bad = await client.post(
        "/api/translation/glossary",
        json={
            "source_language": "uz",
            "target_language": "uz",
            "source_term": "a",
            "target_term": "b",
        },
    )
    assert bad.status_code == 422
    updated = await client.put(
        f"/api/translation/glossary/{term['id']}",
        json={**{k: term[k] for k in ("source_language", "target_language", "source_term")},
              "target_term": "district (tuman)"},
    )  # fmt: skip
    assert updated.json()["target_term"] == "district (tuman)"

    fake_anthropic.replies = ["Saqlandi."]
    await stream_chat(client, message="Glossariyga qo'sh (inglizcha): «viloyat» → «region»")
    listed = (
        await client.get("/api/translation/glossary", params={"target_language": "en"})
    ).json()
    assert {(t["source_term"], t["target_term"]) for t in listed} == {
        ("tuman", "district (tuman)"),
        ("viloyat", "region"),
    }
    assert "Saved in the user's glossary" in (fake_anthropic.requests[-1].system or "")
    fake_anthropic.replies = ["Qaysi til?"]
    await stream_chat(client, message="Glossariyga qo'sh: «viloyat» → «viloyatlar»")
    assert "cannot tell the two languages apart" in (fake_anthropic.requests[-1].system or "")

    # Another user's terms are invisible and cannot be changed.
    async with app.state.services.db.sessionmaker() as session:
        stranger = User(email="other@x", display_name="Other", preferences={})
        session.add(stranger)
        await session.flush()
        foreign = GlossaryTerm(
            user_id=stranger.id,
            source_language="uz",
            target_language="en",
            source_term="maxfiy",
            target_term="secret",
        )
        session.add(foreign)
        await session.commit()
        foreign_id = foreign.id
    assert len((await client.get("/api/translation/glossary")).json()) == 2
    assert (await client.delete(f"/api/translation/glossary/{foreign_id}")).status_code == 404
    assert (await client.delete(f"/api/translation/glossary/{term['id']}")).status_code == 204
    # Deleting an organization deletes its own terms.
    again = await client.post(
        "/api/translation/glossary",
        json={"source_language": "uz", "target_language": "en", "source_term": "hokim",
              "target_term": "khokim", "organization_id": org["id"]},
    )  # fmt: skip
    assert again.status_code == 201
    await client.delete(f"/api/writing/organizations/{org['id']}")
    assert [t["source_term"] for t in (await client.get("/api/translation/glossary")).json()] == [
        "viloyat"
    ]


async def test_turn_based_interpretation(database_url: str) -> None:
    voice = FakeVoice("openai", transcript="Bugun yig'ilish soat 10 da boshlanadi.")
    provider = FakeProvider(name="anthropic", reply="unused")
    provider.replies = ["The meeting starts at 10 today.\n=====\nNOTES: none"]
    app = create_app(
        make_settings(database_url),
        providers={"anthropic": provider},
        voice=VoiceService([voice], [voice]),
    )
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as client:
            response = await client.post(
                "/api/translation/interpret",
                files={"audio": ("turn.webm", WEBM)},
                data={"target": "en", "source": "uz"},
            )
            assert response.status_code == 200, response.text
            body = response.json()
            assert body["mode"] == "turn-based"
            assert body["transcript"] == "Bugun yig'ilish soat 10 da boshlanadi."
            assert body["translation"] == "The meeting starts at 10 today."
            assert body["source_language"] == "uz" and body["target_language"] == "en"
            assert body["audio_mime"] == "audio/mpeg" and body["audio_base64"]
            assert [c[0] for c in voice.calls] == ["stt", "tts"] and voice.calls[1][2] == "en"
    finally:
        await app.state.services.db.dispose()
