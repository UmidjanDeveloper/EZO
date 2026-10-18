"""Professional writing through chat: drafts in the official layout, honest placeholders,
organization profiles, templates without a model, and the API."""

from __future__ import annotations

import io
import json
from typing import Any

from docx import Document
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient

from ezo_api.main import create_app
from tests.conftest import make_settings
from tests.fakes import FakeProvider
from tests.integration.helpers import stream_chat
from tests.integration.test_document_agent import created

FACTS = (
    "Shu faktlar asosida professional tushuntirish xati yoz va Word qilib ber: men, bosh "
    "mutaxassis Bobur Aliyev, 12-mart kuni ishga 30 daqiqa kechikdim, sababi avtobus "
    "buzilib qoldi. Xat tuman hokimi Akmal Karimovga."
)


def reply(**overrides: Any) -> str:
    body: dict[str, Any] = {
        "title": "Tushuntirish xati",
        "fields": {
            "recipient": "Tuman hokimi Akmal Karimovga",
            "sender": "Bosh mutaxassis Bobur Aliyev",
            "date": "2026-yil 13-mart",
            "event": "12-mart kuni ishga 30 daqiqa kechikish",
        },
        "sections": [
            {
                "heading": None,
                "text": "12-mart kuni ishga 30 daqiqa kechikdim, chunki men chiqqan avtobus "
                "yo'lda buzilib qoldi.",
            },
            {
                "heading": None,
                "text": "Bundan keyin bunday holat takrorlanmasligi uchun ertaroq yo'lga chiqaman.",
            },
        ],
        "missing": ["xat sanasi"],
    }
    body.update(overrides)
    return json.dumps(body, ensure_ascii=False)


async def artifact_detail(
    client: AsyncClient, artifact: dict[str, Any]
) -> tuple[dict[str, Any], Document]:
    detail = (await client.get(f"/api/artifacts/{artifact['id']}")).json()["current_version"]
    data = (await client.get(f"/api/artifacts/{artifact['id']}/versions/1/download")).content
    return detail, Document(io.BytesIO(data))


async def test_explanatory_note_from_facts(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    fake_anthropic.replies = [reply(), "Tushuntirish xati tayyor."]
    events = await stream_chat(client, message=FACTS)
    assert events[0][1]["agent"] == "document"
    [artifact] = created(events)
    writer = fake_anthropic.requests[0]
    system = writer.system or ""
    assert "Explanatory note" in system and "Uzbek quality rules" in system
    assert "Never invent names" in system and FACTS in writer.messages[0].content
    version, document = await artifact_detail(client, artifact)
    meta = version["generation_metadata"]
    assert meta["doc_type"] == "tushuntirish_xati" and meta["draft"] is True
    # "13-mart" was never given: the date is a placeholder, not an invention.
    assert meta["validation"]["replaced_fields"] == ["date"]
    paragraphs = [p.text for p in document.paragraphs]
    assert "TUSHUNTIRISH XATI" in paragraphs and "Tuman hokimi Akmal Karimovga" in paragraphs
    assert "[Sana]" in paragraphs and not any("13-mart" in p for p in paragraphs)
    cells = [c.text for t in document.tables for r in t.rows for c in r.cells]
    assert "[imzo]   Bosh mutaxassis Bobur Aliyev" in cells
    final = fake_anthropic.requests[-1].system or ""
    assert "DRAFT" in final and "the user did not give them: date" in final


async def test_invented_numbers_and_claims_are_reported(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    fake_anthropic.replies = [
        reply(
            sections=[
                {
                    "heading": None,
                    "text": "Kechikish 45 daqiqa bo'ldi. Xat hokim tomonidan imzolandi.",
                }
            ]
        ),
        "Tayyor.",
    ]
    await stream_chat(client, message=FACTS)
    final = fake_anthropic.requests[-1].system or ""
    line = next(ln for ln in final.splitlines() if "ask them to check" in ln)
    assert "45" in line and "30" not in line  # 30 was in the user's facts
    assert "claims signing/approval (imzolandi)" in final


async def test_organization_profile_letterhead_and_terms(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    body = {
        "name": "Xatirchi tumani hokimligi",
        "org_type": "government",
        "department": "Yoshlar ishlari bo'limi",
        "address": "Yangirabod sh., Mustaqillik ko'chasi 1",
        "terminology": {"mutasaddi": "mas'ul"},
        "style_notes": "Qisqa va aniq yozamiz.",
    }
    created_org = await client.post("/api/writing/organizations", json=body)
    assert created_org.status_code == 201
    org = created_org.json()
    assert org["is_default"] is True
    fake_anthropic.replies = [
        json.dumps(
            {
                "title": "Xizmat xati",
                "fields": {
                    "recipient": "Navoiy viloyati hokimligiga",
                    "subject": "Yoshlar bilan uchrashuv",
                },
                "sections": [
                    {"heading": None, "text": "Tumanda yoshlar bilan uchrashuv o'tkazildi."}
                ],
                "missing": [],
            }
        ),
        "Tayyor.",
    ]
    events = await stream_chat(
        client,
        message="Tashkilotimiz uslubida viloyat hokimligiga xizmat xati tayyorla: yoshlar "
        "bilan uchrashuv o'tkazildi",
    )
    [artifact] = created(events)
    profile_block = fake_anthropic.requests[0].messages[0].content
    assert "Organization: Xatirchi tumani hokimligi" in profile_block
    assert "use “mas'ul” instead of “mutasaddi”" in profile_block
    version, document = await artifact_detail(client, artifact)
    assert version["generation_metadata"]["organization_profile_id"] == org["id"]
    assert document.paragraphs[0].text == "Xatirchi tumani hokimligi"
    # Profiles can be edited and removed.
    body["name"] = "Xatirchi tumani hokimligi (yangi)"
    assert (
        (await client.put(f"/api/writing/organizations/{org['id']}", json=body))
        .json()["name"]
        .endswith("(yangi)")
    )
    assert (await client.delete(f"/api/writing/organizations/{org['id']}")).status_code == 204
    assert (await client.get("/api/writing/organizations")).json() == []


async def test_template_request_uses_no_model(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    fake_anthropic.replies = ["Mana shablon."]
    events = await stream_chat(client, message="Yig'ilish bayonnomasi shablonini tayyorla")
    [artifact] = created(events)
    assert len(fake_anthropic.requests) == 1  # only the short reply, no drafting call
    version, document = await artifact_detail(client, artifact)
    assert version["generation_metadata"]["action"] == "template"
    text = "\n".join(p.text for p in document.paragraphs)
    assert "BAYONNOMA № [Hujjat raqami]" in text and "Rais: [Rais]" in text


async def test_document_types_endpoint_and_russian(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    types = (await client.get("/api/writing/document-types")).json()
    assert len(types) >= 30 and {"ariza", "bayonnoma", "cv", "press_reliz"} <= {
        t["key"] for t in types
    }
    fake_anthropic.replies = [reply(title="Объяснительная записка"), "Готово."]
    await stream_chat(
        client, message="Напиши объяснительную записку на русском: опоздал на 30 минут 12 марта"
    )
    assert "Language: Russian" in (fake_anthropic.requests[0].system or "")


async def test_offline_gets_a_labelled_template(database_url: str) -> None:
    app: FastAPI = create_app(make_settings(database_url), providers={})
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as client:
            events = await stream_chat(client, message="Hokimga ariza yozib ber")
            [artifact] = created(events)
            version, document = await artifact_detail(client, artifact)
            assert version["generation_metadata"]["action"] == "template"
            assert {e["name"] for e in version["exports"]} == {"pdf", "markdown", "text"}
            assert "ARIZA" in [p.text for p in document.paragraphs]
    finally:
        await app.state.services.db.dispose()
