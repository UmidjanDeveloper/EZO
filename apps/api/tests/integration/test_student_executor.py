"""Actual isolated parsers + durable DB queue + K2 + MOCK AI end-to-end.

Only the model is mocked: uploads, container boundaries, DB transactions and
source offsets use their real implementations. Never contacts a paid provider.
"""

import io
import json
import os
import uuid

import docx
import pytest
from reportlab.pdfgen.canvas import Canvas

from ezo_api.documents.executor.controller import IsolatedExecutor
from ezo_api.student.ingestion import StudentIngestion
from ezo_api.student.worker import StudentWorker
from tests.unit.test_isolated_executor import book

IMAGE = os.environ.get("EZO_EXECUTOR_TEST_IMAGE")
pytestmark = pytest.mark.skipif(not IMAGE, reason="requires pinned isolated executor image")


@pytest.mark.parametrize("format", ["pdf", "docx", "epub"])
async def test_upload_isolated_worker_index_answer_guide_quiz(client, app, fake_anthropic, format):
    if format == "epub":
        data = book(
            "<html><body>Fotosintez quyosh energiyasini saqlaydi. Ўзбек 🙂</body></html>".encode()
        )
    elif format == "docx":
        document = docx.Document()
        document.add_heading("Fotosintez", level=1)
        document.add_paragraph("Fotosintez quyosh energiyasini saqlaydi. Ўзбек 🙂")
        output = io.BytesIO()
        document.save(output)
        data = output.getvalue()
    else:
        output = io.BytesIO()
        canvas = Canvas(output)
        canvas.drawString(20, 800, "Fotosintez quyosh energiyasini saqlaydi.")
        canvas.save()
        data = output.getvalue()
    response = await client.post(
        "/api/student/uploads", params={"filename": f"book.{format}"}, content=data
    )
    assert response.status_code == 201, response.text
    job = response.json()
    assert await StudentWorker(
        StudentIngestion(app.state.services.db), IsolatedExecutor(IMAGE), "real-native"
    ).process_one()
    result = await client.get(f"/api/student/readings/{job['id']}")
    assert result.status_code == 200, result.text
    text = result.json()["text"]
    assert "Fotosintez" in text
    if format != "pdf":
        assert "Ўзбек 🙂" in text
    doc = job["document_id"]
    indexed = await client.post(f"/api/student/documents/{doc}/semantic-index")
    assert indexed.status_code == 200 and indexed.json()["indexed"] > 0
    fake_anthropic.reply = "Fotosintez quyosh energiyasini saqlaydi [1]."
    for kind in ("answer", "study_guide", "summary"):
        answer = await client.post(
            f"/api/student/documents/{doc}/learn", json={"kind": kind, "question": "Fotosintez"}
        )
        assert answer.status_code == 200, answer.text
        citation = answer.json()["citations"][0]
        assert citation["revision_id"] and citation["extra"]["source_sha256"]
        assert text[citation["extra"]["start"] : citation["extra"]["end"]].startswith(
            citation["excerpt"]
        )
    fake_anthropic.reply = json.dumps(
        {
            "items": [
                {
                    "question": "Nima saqlanadi?",
                    "choices": ["Energiya", "Qum"],
                    "answer": 0,
                    "explanation": "Energiya [1].",
                    "source": 1,
                }
            ]
        }
    )
    quiz = await client.post(f"/api/student/documents/{doc}/learn", json={"kind": "quiz"})
    assert quiz.status_code == 200 and len(quiz.json()["quiz"]["items"]) == 1


@pytest.mark.parametrize("action", ["cancel", "expire"])
async def test_real_worker_revocation_removes_container_and_fences_publication(client, app, action):
    import asyncio

    from tests.integration.test_document_jobs import expire
    from tests.integration.test_isolated_executor_container import PROBE, names, wait_new

    if not PROBE:
        pytest.skip("requires hostile fixture image")

    class StallExecutor(IsolatedExecutor):
        def extract(self, data, format, *, ocr_languages=(), cancelled=None):
            # Trusted adversarial fixture, not a user-selectable execution mode.
            return super().extract(b"stall", "pdf", cancelled=cancelled)

    output = io.BytesIO()
    canvas = Canvas(output)
    canvas.drawString(20, 800, "Fotosintez test")
    canvas.save()
    job = (
        await client.post("/api/student/uploads?filename=revoke.pdf", content=output.getvalue())
    ).json()
    jobs = StudentIngestion(app.state.services.db)
    previous = await asyncio.to_thread(names)
    task = asyncio.create_task(
        StudentWorker(jobs, StallExecutor(PROBE), "revocation").process_one()
    )
    await asyncio.to_thread(wait_new, previous)
    if action == "cancel":
        cancelled = await client.post(f"/api/knowledge/document-jobs/{job['id']}/cancel")
        assert cancelled.status_code == 200
    else:
        from sqlalchemy import select

        from ezo_api.db.models import DocumentJob
        from ezo_api.knowledge.jobs.service import Lease

        async with app.state.services.db.sessionmaker() as session:
            current = await session.scalar(
                select(DocumentJob).where(DocumentJob.id == uuid.UUID(job["id"]))
            )
            lease = Lease(
                current.id,
                current.user_id,
                current.project_id,
                current.lease_token,
                current.attempts,
                current.input_revision_id,
                current.parser_version,
            )
        await expire(app.state.services.db, lease)
    assert await asyncio.wait_for(task, timeout=40)
    assert await asyncio.to_thread(names) == previous
    assert (await client.get(f"/api/student/readings/{job['id']}")).status_code == 409
    if action == "expire":
        assert await StudentWorker(jobs, IsolatedExecutor(IMAGE), "recovery").process_one()
        assert (await client.get(f"/api/student/readings/{job['id']}")).status_code == 200


async def test_scanned_pdf_upload_is_ocrd_indexed_and_cited(client, app, fake_anthropic):
    """Real isolated OCR → fenced publication → K2 retrieval → cited answer (mock model)."""
    from ezo_api.documents.pdf import engine
    from tests.ocr_fixtures import UZ_CYRILLIC, page_image, png, recall

    scan = engine.images_to_pdf([png(page_image(UZ_CYRILLIC))]).data
    response = await client.post(
        "/api/student/uploads",
        params={"filename": "qaror.pdf", "ocr": ["uzb_cyrl", "rus"]},
        content=scan,
    )
    assert response.status_code == 201, response.text
    job = response.json()
    assert await StudentWorker(
        StudentIngestion(app.state.services.db), IsolatedExecutor(IMAGE), "real-ocr"
    ).process_one()
    reading = await client.get(f"/api/student/readings/{job['id']}")
    assert reading.status_code == 200, reading.text
    text = reading.json()["text"]
    assert recall(UZ_CYRILLIC, text) >= 0.9, text
    structure = await client.get(f"/api/student/documents/{job['document_id']}/structure")
    assert structure.json()["sections"][0]["page"] == 1
    fake_anthropic.reply = "Таълим тизимини ривожлантириш [1]."
    answer = await client.post(
        f"/api/student/documents/{job['document_id']}/learn",
        json={"kind": "answer", "question": "Таълим тизимини ривожлантириш"},
    )
    assert answer.status_code == 200, answer.text
    citation = answer.json()["citations"][0]
    assert text[citation["extra"]["start"] : citation["extra"]["end"]].startswith(
        citation["excerpt"]
    )


async def test_scan_uploaded_without_ocr_fails_then_retry_with_ocr_succeeds(client, app):
    from ezo_api.documents.pdf import engine
    from tests.ocr_fixtures import ENGLISH, page_image, png

    scan = engine.images_to_pdf([png(page_image(ENGLISH))]).data
    job = (
        await client.post(
            "/api/student/uploads",
            params={"filename": "scan.pdf", "no_ocr": "true"},
            content=scan,
        )
    ).json()
    worker = StudentWorker(
        StudentIngestion(app.state.services.db), IsolatedExecutor(IMAGE), "real-ocr"
    )
    assert await worker.process_one()
    failed = (await client.get("/api/student/library")).json()[0]["job"]
    assert failed["id"] == job["id"] and failed["error_code"] == "ocr_required"
    retry = await client.post(f"/api/student/readings/{job['id']}/retry", params={"ocr": ["eng"]})
    assert retry.status_code == 201, retry.text
    assert await worker.process_one()
    reading = await client.get(f"/api/student/readings/{retry.json()['id']}")
    assert reading.status_code == 200 and "quick brown fox" in reading.json()["text"].lower()
