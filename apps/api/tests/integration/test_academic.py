"""Academic assistance through chat and the API: a research plan citing only works EZO
found, a literature list without a model, reference formatting that drops invented
fields, and the deterministic /api/academic endpoints."""

from __future__ import annotations

import io
import json
from dataclasses import dataclass, field
from typing import Any

from docx import Document
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from ezo_api.academic.citations import Reference
from ezo_api.academic.scholar import Work
from ezo_api.main import create_app
from tests.conftest import make_settings
from tests.fakes import FakeProvider
from tests.integration.helpers import stream_chat
from tests.integration.test_document_agent import created

WORKS = [
    Work(
        reference=Reference(
            title="E-government adoption in Central Asia",
            authors=["Aziza Nazarova", "John Smith"],
            year=2021,
            container="Government Information Quarterly",
            volume="38",
            issue="2",
            pages="101-115",
            doi="10.5555/egov.2021.7",
        ),
        openalex_id="https://openalex.org/W1",
        cited_by=42,
        open_access_url="",
        abstract="This study reviews e-government adoption in Central Asia.",
        language="en",
    ),
    Work(
        reference=Reference(title="Digital services and trust", year=2019, container="Policy"),
        openalex_id="https://openalex.org/W2",
        cited_by=3,
        open_access_url="https://example.org/trust.pdf",
        abstract="",
        language="en",
    ),
]


@dataclass
class FakeScholar:
    name: str = "openalex"
    works: list[Work] = field(default_factory=lambda: list(WORKS))
    queries: list[str] = field(default_factory=list)

    async def search(self, query: str, limit: int = 10, from_year: int | None = None) -> list[Work]:
        self.queries.append(query)
        return self.works[:limit]


def _app(database_url: str, scholar: FakeScholar | None, provider: FakeProvider | None) -> FastAPI:
    providers = {"anthropic": provider} if provider else {}
    return create_app(make_settings(database_url), providers=providers, scholar=scholar)


async def detail(client: AsyncClient, artifact: dict[str, Any]) -> tuple[dict[str, Any], bytes]:
    version = (await client.get(f"/api/artifacts/{artifact['id']}")).json()["current_version"]
    data = (await client.get(f"/api/artifacts/{artifact['id']}/versions/1/download")).content
    return version, data


PLAN = {
    "title": "Raqamli hukumat: tadqiqot rejasi",
    "research_questions": ["Fuqarolar e-hukumatni qanday qabul qiladi? [1]"],
    "hypotheses": ["Ishonch qabul qilishni oshiradi."],
    "methodology": {"design": "Aralash usul", "methods": ["So'rovnoma", "Intervyu"]},
    "literature_review": [
        {
            "theme": "Qabul qilish",
            "text": "Markaziy Osiyoda qabul qilish o'smoqda [1, 7]. "
            "Natijalar shuni ko'rsatdiki, ishonch muhim.",
        }
    ],
    "timeline": [{"period": "1-oy", "task": "Adabiyotlar tahlili"}],
    "defense_questions": [{"question": "Ilmiy yangilik nimada?", "hint": "Yangilikni aniq ayting"}],
}


async def test_research_plan_cites_only_found_works(database_url: str) -> None:
    scholar = FakeScholar()
    provider = FakeProvider(name="anthropic", reply="unused")
    provider.replies = [
        json.dumps(
            {
                "topic": "Raqamli hukumat",
                "queries": ["e-government adoption Central Asia"],
                "from_year": 2015,
            }
        ),
        json.dumps(PLAN, ensure_ascii=False),
        "Reja tayyor.",
    ]
    app = _app(database_url, scholar, provider)
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as client:
            events = await stream_chat(
                client,
                message="Raqamli hukumat mavzusi bo'yicha ilmiy manbalarni top, tadqiqot rejasini "
                "tuz va himoyaga tayyorlanishimga yordam ber",
            )
            assert events[0][1]["agent"] == "document"
            word, sheet = created(events)
            assert scholar.queries == ["e-government adoption Central Asia"]
            plan_call = provider.requests[1]
            assert (
                '<work n="1">E-government adoption in Central Asia' in plan_call.messages[0].content
            )
            assert "Never invent a work" in (plan_call.system or "")
            version, data = await detail(client, word)
            text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
            # In-text citations in APA, built by EZO; the bad [7] is removed.
            assert "Markaziy Osiyoda qabul qilish o'smoqda (Nazarova & Smith, 2021)." in text
            assert "Foydalanilgan adabiyotlar (APA 7)" in text
            assert (
                "Nazarova, A., & Smith, J. (2021). E-government adoption in Central Asia. "
                "Journal" not in text
            )
            assert "https://doi.org/10.5555/egov.2021.7" in text
            check = version["generation_metadata"]["check"]
            assert check["bad_citations"] == ["[7]"] and check["cited"] == [1]
            assert check["unsupported_results"] == ["Natijalar shuni ko'rsatdiki, ishonch muhim."]
            _, data = await detail(client, sheet)
            rows = list(load_workbook(io.BytesIO(data)).active.iter_rows(values_only=True))
            assert rows[1][3] == "E-government adoption in Central Asia" and rows[1][7] == 42
            assert rows[2][9] == "authors, volume, pages"  # missing, not invented
            final = provider.requests[-1].system or ""
            assert "Citations to works that do not exist were removed: [7]" in final
            assert "rewrite them as hypotheses" in final
    finally:
        await app.state.services.db.dispose()


async def test_literature_list_without_a_model_and_search_off(database_url: str) -> None:
    scholar = FakeScholar()
    app = _app(database_url, scholar, None)
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as client:
            events = await stream_chat(
                client, message="Find research papers on e-government adoption, IEEE style"
            )
            word, _sheet = created(events)
            assert scholar.queries == ["e-government adoption ieee style"]
            version, data = await detail(client, word)
            assert version["generation_metadata"]["action"] == "literature_list"
            text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
            assert "[1] A. Nazarova and J. Smith, “E-government adoption in Central Asia,”" in text
    finally:
        await app.state.services.db.dispose()
    app = _app(database_url, None, None)
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as client:
            events = await stream_chat(client, message="Ilmiy manbalarni top: suv tejash")
            assert created(events) == []
            search = await client.get("/api/academic/search", params={"q": "water"})
            assert search.status_code == 503
    finally:
        await app.state.services.db.dispose()


async def test_reference_formatting_drops_invented_fields(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    extracted = {
        "references": [
            {
                "kind": "book",
                "authors": ["Karimov A."],
                "year": 2020,
                "title": "Raqamli hukumat",
                "publisher": "Fan",
                "place": "Toshkent",
                "doi": "10.9999/invented",
            }
        ]
    }
    fake_anthropic.replies = [json.dumps(extracted, ensure_ascii=False), "Tayyor."]
    events = await stream_chat(
        client,
        message="Shu adabiyotlarni APA formatiga o'tkaz:\n"
        "Karimov A. Raqamli hukumat. Toshkent: Fan, 2020.",
    )
    [word] = created(events)
    final = fake_anthropic.requests[-1].system or ""
    assert "Karimov, A. (2020). *Raqamli hukumat.* Fan." in final
    assert "removed values not in the text (doi)" in final
    version, _ = await detail(client, word)
    assert version["generation_metadata"]["style"] == "apa"


async def test_cite_and_search_api(database_url: str) -> None:
    app = _app(database_url, FakeScholar(), None)
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as client:
            styles = (await client.get("/api/academic/styles")).json()
            assert [s["key"] for s in styles] == [
                "apa",
                "mla",
                "chicago",
                "harvard",
                "ieee",
                "gost",
            ]
            body = {
                "style": "harvard",
                "references": [
                    {"title": "Open data portal", "kind": "web", "url": "https://data.gov.uz"}
                ],
            }
            cited = (await client.post("/api/academic/cite", json=body)).json()
            assert cited["entries"] == [
                "Open data portal (n.d.). Available at: https://data.gov.uz."
            ]
            assert cited["missing"] == [
                {"title": "Open data portal", "fields": ["authors", "year"]}
            ]
            found = (await client.get("/api/academic/search", params={"q": "e-government"})).json()
            assert found["works"][0]["apa"].startswith("Nazarova, A., & Smith, J. (2021).")
    finally:
        await app.state.services.db.dispose()
