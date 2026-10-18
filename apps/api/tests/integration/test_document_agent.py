"""DocumentAgent end to end: create, transform (new artifact / new version), convert, compare,
spreadsheet analysis and creation, project scope and offline honesty."""

from __future__ import annotations

import io
import json
from typing import Any

from docx import Document
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from ezo_api.main import create_app
from tests.conftest import make_settings
from tests.fakes import FakeProvider
from tests.files import make_xlsx
from tests.integration.helpers import stream_chat

ORIGINAL = (
    "Shartnoma shartlari\n\nIjara narxi oyiga 4 500 000 so'm.\n\n"
    "To'lov har oyning 5-sanasigacha amalga oshiriladi."
)


async def upload(client: AsyncClient, name: str, data: bytes, **form: str) -> dict[str, Any]:
    response = await client.post(
        "/api/knowledge/documents", files={"file": (name, data)}, data=form
    )
    assert response.status_code == 201, response.text
    return response.json()["document"]  # type: ignore[no-any-return]


def created(events: list[tuple[str, dict[str, Any]]]) -> list[dict[str, Any]]:
    assert events[-1][0] == "done", events[-1]
    return events[-1][1]["message"]["metadata"]["artifacts"]  # type: ignore[no-any-return]


async def test_create_official_document_as_docx_with_pdf_export(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    fake_anthropic.replies = [
        "TITLE: Ariza\nCHANGES: new document\n=====\n# Ariza\n\nTuman hokimi [F.I.Sh.] ga\n\n"
        "## Mazmuni\n\nMen, [F.I.Sh.], uy-joy masalasida yordam so'rayman.\n\n"
        "- Manzil: [manzil]\n\n"
        "Sana: [sana]",
        "Ariza tayyor.",
    ]
    events = await stream_chat(client, message="Hokimga rasmiy ariza yozib ber, word fayl qilib")
    assert events[0][1]["agent"] == "document"
    [artifact] = created(events)
    assert artifact["type"] == "document" and artifact["title"] == "Ariza"
    writer = fake_anthropic.requests[0]
    assert writer.model == "claude-opus-5-5" and "Uzbekistan" in (writer.system or "")
    system = fake_anthropic.requests[-1].system or ""
    assert "Placeholders the user must fill in" in system and "[F.I.Sh.]" in system

    detail = (await client.get(f"/api/artifacts/{artifact['id']}")).json()
    version = detail["current_version"]
    assert version["filename"] == "Ariza.docx"
    assert {e["name"] for e in version["exports"]} == {"pdf", "markdown"}
    assert version["generation_metadata"]["action"] == "create"
    docx = await client.get(f"/api/artifacts/{artifact['id']}/versions/1/download")
    paragraphs = [p.text for p in Document(io.BytesIO(docx.content)).paragraphs]
    assert "Ariza" in paragraphs and "Mazmuni" in paragraphs
    pdf = await client.get(f"/api/artifacts/{artifact['id']}/versions/1/exports/pdf")
    assert pdf.content.startswith(b"%PDF")


async def test_translation_creates_a_new_artifact_and_keeps_the_original(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    document = await upload(client, "shartnoma.txt", ORIGINAL.encode())
    fake_anthropic.replies = [
        "TITLE: Contract terms\nCHANGES: Translated into English.\n=====\n# Contract terms\n\n"
        "The rent is 4,500,000 soums per month.\n\nPayment is due by the 5th of each month.",
        "Tarjima tayyor.",
    ]
    events = await stream_chat(client, message="Shu faylni ingliz tiliga tarjima qil")
    [artifact] = created(events)
    editor = fake_anthropic.requests[0]
    assert editor.model == "claude-haiku-5-5"  # translation runs on the cheap chat route
    assert "into English" in (editor.system or "") and ORIGINAL in editor.messages[0].content
    detail = (await client.get(f"/api/artifacts/{artifact['id']}")).json()
    metadata = detail["current_version"]["generation_metadata"]
    assert metadata["action"] == "translate" and metadata["target_language"] == "en"
    assert metadata["derived_from"] == {
        "kind": "knowledge_document",
        "id": document["id"],
        "title": "shartnoma",
    }
    assert "paragraph" in metadata["comparison"]
    assert detail["current_version"]["change_summary"] == "Translated into English."
    system = fake_anthropic.requests[-1].system or ""
    assert "the original file is unchanged" in system and "Change summary" in system
    # The uploaded file itself is untouched.
    doc = (await client.get(f"/api/knowledge/documents/{document['id']}")).json()
    assert doc["size_bytes"] == len(ORIGINAL.encode()) and doc["status"] == document["status"]

    # A follow-up edit of EZO's own Word document is done in place (paragraph by paragraph)
    # and becomes version 2 of the same artifact.
    fake_anthropic.replies = [
        "[[2]] The monthly rent amounts to 4,500,000 soums.\n"
        "[[3]] Payment shall be made by the 5th of each month.",
        "Yangilandi.",
    ]
    conversation_id = events[0][1]["conversation_id"]
    follow = await stream_chat(
        client, message="Hujjatni rasmiy uslubda qayta yoz", conversation_id=conversation_id
    )
    [again] = created(follow)
    assert again["id"] == artifact["id"] and again["version"] == 2
    sent = fake_anthropic.requests[-2].messages[0].content
    assert "[[0]] Contract terms" in sent and "[[2]] The rent is 4,500,000" in sent
    versions = (await client.get(f"/api/artifacts/{artifact['id']}")).json()["versions"]
    assert versions[0]["change_summary"] == "Formal style"
    assert versions[0]["parent_version_id"] == versions[1]["id"]
    v2 = (await client.get(f"/api/artifacts/{artifact['id']}/versions/2/download")).content
    paragraphs = [p.text for p in Document(io.BytesIO(v2)).paragraphs if p.text]
    assert paragraphs[-3:] == [
        "Contract terms",
        "The monthly rent amounts to 4,500,000 soums.",
        "Payment shall be made by the 5th of each month.",
    ]


async def test_compare_two_uploaded_files(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    await upload(client, "eski.txt", ORIGINAL.encode())
    await upload(client, "yangi.txt", ORIGINAL.replace("4 500 000", "5 000 000").encode())
    fake_anthropic.replies = ["Ijara narxi o'zgargan."]
    events = await stream_chat(client, message="Bu ikki hujjatni solishtir")
    assert events[0][1]["agent"] == "document"
    system = fake_anthropic.requests[-1].system or ""
    assert "Comparison of “eski” (original) and “yangi” (newer)" in system
    assert "5 000 000" in system and "4 500 000" in system
    assert created(events) == []  # no file was asked for


async def test_spreadsheet_analysis_uses_computed_facts(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    rows: list[list[object]] = [["Oy", "Savdo"], ["Yanvar", 120], ["Fevral", 180], ["Mart", 150]]
    await upload(client, "savdo.xlsx", make_xlsx(rows))
    fake_anthropic.replies = ["Savdo jami 450."]
    await stream_chat(client, message="Shu excel jadvalni tahlil qil")
    system = fake_anthropic.requests[-1].system or ""
    assert "Spreadsheet facts computed by EZO" in system
    assert "sum 450" in system and "mean 150" in system


async def test_spreadsheet_creation(client: AsyncClient, fake_anthropic: FakeProvider) -> None:
    table = {
        "title": "Oylik xarajatlar",
        "sheets": [
            {
                "name": "Xarajatlar",
                "columns": ["Modda", "Summa (so'm)"],
                "rows": [["Ijara", 4500000], ["Kommunal", None], ["=HYPERLINK(1)", 1]],
            }
        ],
    }
    fake_anthropic.replies = [json.dumps(table, ensure_ascii=False), "Jadval tayyor."]
    events = await stream_chat(client, message="Oylik xarajatlar uchun excel jadval yarat")
    [artifact] = created(events)
    assert artifact["type"] == "spreadsheet"
    data = await client.get(f"/api/artifacts/{artifact['id']}/versions/1/download")
    sheet = load_workbook(io.BytesIO(data.content)).active
    assert sheet["A2"].value == "Ijara" and sheet["B2"].value == 4500000
    assert sheet["B3"].value is None  # unknown values stay empty, never invented
    assert not str(sheet["A4"].value).startswith("=")  # formula injection is neutralised
    detail = (await client.get(f"/api/artifacts/{artifact['id']}")).json()
    assert {e["name"] for e in detail["current_version"]["exports"]} == {"csv"}


async def test_files_outside_the_project_scope_are_not_used(
    client: AsyncClient, fake_anthropic: FakeProvider
) -> None:
    project = (await client.post("/api/projects", json={"name": "Boshqa loyiha"})).json()
    await upload(client, "maxfiy.txt", b"Maxfiy reja matni.", project_id=project["id"])
    fake_anthropic.replies = ["Fayl topilmadi."]
    await stream_chat(client, message="Shu faylni ingliz tiliga tarjima qil")
    system = fake_anthropic.requests[-1].system or ""
    assert "no matching uploaded file was found" in system
    assert "Maxfiy reja" not in str(fake_anthropic.requests[-1].messages)


async def test_conversion_works_offline(database_url: str) -> None:
    app: FastAPI = create_app(make_settings(database_url), providers={})
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as client:
            await upload(client, "shartnoma.txt", ORIGINAL.encode())
            events = await stream_chat(client, message="Shu faylni PDF ga aylantir")
            [artifact] = created(events)
            detail = (await client.get(f"/api/artifacts/{artifact['id']}")).json()
            assert detail["current_version"]["filename"] == "shartnoma.pdf"
            assert detail["current_version"]["generation_metadata"]["action"] == "convert"
            # Writing needs a model: offline it is refused honestly, nothing is created.
            refused = await stream_chat(client, message="Hokimga rasmiy ariza yozib ber")
            assert created(refused) == []
    finally:
        await app.state.services.db.dispose()
